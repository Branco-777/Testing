"""Define the Excel-independent contract before Bonds Lite implementation."""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import sys
import threading
import time
from copy import deepcopy
from contextlib import ExitStack, contextmanager
from datetime import datetime
from itertools import groupby
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterator

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from bonds_lite_input import (
    STRESS_BLOCKS,
    VBA_LAUNCHER,
    BondsLiteError,
    choose_output_path,
    column_name,
    resolve_paths,
    validate_headers,
    validate_identifiers,
    validate_stress_values,
    workbook_mutex_name,
)

CASES = json.loads(
    (Path(__file__).parent / "fixtures" / "bonds_lite_cases.json").read_text(
        encoding="utf-8"
    )
)
BLOCK_CASES = CASES["stress_blocks"]
RAW_FIRST_COLUMN = 2
RAW_LAST_COLUMN = 1216
RAW_COLUMN_COUNT = RAW_LAST_COLUMN - RAW_FIRST_COLUMN + 1
EXPECTED_INPUT_FILENAMES = {
    "Bonds": "Bonds Lite input_Bonds.xlsx",
    "MAP": "Bonds Lite input_MAP.xlsx",
}


def _complete_run(folder: Path) -> Path:
    """Create distinct nonempty members for a filesystem-only handoff."""
    folder.mkdir(parents=True)
    for dataset, filename in EXPECTED_INPUT_FILENAMES.items():
        (folder / filename).write_bytes(f"Published {dataset} sentinel".encode())
    return folder


def _pair_bytes(folder: Path) -> dict[str, bytes]:
    """Capture both members to detect accidental deletion or overwriting."""
    return {
        name: (folder / name).read_bytes() for name in EXPECTED_INPUT_FILENAMES.values()
    }


def test_stage_field_names_preserve_initial_install_anchors() -> None:
    """Separate permanent field identities from first-install coordinates."""
    import bonds_lite_input as workflow

    expected = {
        "last_run": "SST_Stage1_LastRun",
        "start": "SST_Stage1_Start",
        "finish": "SST_Stage1_Finish",
        "elapsed": "SST_Stage1_Elapsed",
        "user": "SST_Stage1_User",
        "status": "SST_Stage1_Status",
        "last_output": "SST_Stage1_LastSuccessfulOutput",
        "stage_two_input": "SST_Stage2_Input",
    }
    assert workflow.STAGE_ONE_FIELD_NAMES == expected
    assert workflow.STAGE_ONE_FIELDS == {
        "last_run": "L74", "start": "L75", "finish": "L76", "elapsed": "L77",
        "user": "L78", "status": "L79", "last_output": "L80",
        "stage_two_input": "L82",
    }
    assert workflow.STAGE_ONE_NAMES == {
        "SST_Stage1_Panel": "H73:O83",
        **{name: workflow.STAGE_ONE_FIELDS[field] for field, name in expected.items()},
    }
    assert callable(workflow._stage_field_cells)


@pytest.fixture
def named_stage_cells(monkeypatch: pytest.MonkeyPatch) -> Any:
    """Keep audit I/O tests behind the owned named-cell resolution boundary."""
    import bonds_lite_input as workflow

    cells = {
        field: SimpleNamespace(value=f"Moved {field}")
        for field in workflow.STAGE_ONE_FIELDS
    }
    events: list[str] = []

    def forbidden_range(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Runtime audit I/O must not use a fixed worksheet address")

    caller = SimpleNamespace(
        sheets={"I. Control": SimpleNamespace(range=forbidden_range)},
        app=SimpleNamespace(api=SimpleNamespace(EnableEvents=True)),
        api=SimpleNamespace(Saved=True),
    )

    def validate(book: Any) -> None:
        assert book is caller
        events.append("validate")

    def resolve(book: Any) -> dict[str, Any]:
        assert book is caller
        events.append("resolve")
        return cells

    def save() -> None:
        assert caller.app.api.EnableEvents is False
        events.append("save")

    caller.save = save
    monkeypatch.setattr(workflow, "require_stage_panel", validate)
    monkeypatch.setattr(workflow, "_stage_field_cells", resolve, raising=False)
    return SimpleNamespace(caller=caller, cells=cells, events=events)


def test_read_stage_fields_uses_named_cells_without_fixed_address_fallback(
    named_stage_cells: Any,
) -> None:
    from bonds_lite_input import _read_stage_fields

    state = named_stage_cells
    assert _read_stage_fields(state.caller) == {
        field: cell.value for field, cell in state.cells.items()
    }
    assert "resolve" in state.events


def test_write_stage_audit_uses_named_cells_and_saves_after_validation(
    named_stage_cells: Any,
) -> None:
    from bonds_lite_input import write_stage_audit

    state = named_stage_cells
    unchanged = {field: cell.value for field, cell in state.cells.items()}
    assert write_stage_audit(state.caller, {"status": "Running", "elapsed": 0.25}) is None
    assert state.cells["status"].value in ("Running", "'Running")
    assert state.cells["elapsed"].value == 0.25
    for field in unchanged.keys() - {"status", "elapsed"}:
        assert state.cells[field].value == unchanged[field]
    assert state.events.index("validate") < state.events.index("save")
    assert state.events.index("resolve") < state.events.index("save")
    assert state.events.count("save") == 1
    assert state.caller.app.api.EnableEvents is True


@pytest.mark.parametrize("field", ["last_output", "stage_two_input"])
@pytest.mark.parametrize("value", ["Bonds Lite output/20261028_090000", ""])
def test_write_stage_audit_updates_named_hyperlink_cell(
    named_stage_cells: Any, field: str, value: str,
) -> None:
    from bonds_lite_input import write_stage_audit

    state = named_stage_cells
    links: list[tuple[str, str]] = []
    deleted: list[bool] = []
    cell = state.cells[field]
    cell.api = SimpleNamespace(
        Hyperlinks=SimpleNamespace(Count=1, Delete=lambda: deleted.append(True))
    )
    cell.add_hyperlink = lambda address, text_to_display: links.append(
        (address, text_to_display)
    )
    write_stage_audit(state.caller, {field: value})
    assert deleted == [True]
    assert links == ([(value, value)] if value else [])
    if not value:
        assert cell.value in (None, "")
    assert state.events.count("save") == 1


def test_stage_field_cells_resolves_names_and_current_merge_top_left(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A name may point inside its merged value region after a user edit."""
    import bonds_lite_input as workflow

    names = workflow.STAGE_ONE_FIELD_NAMES
    anchors: dict[tuple[int, int], Any] = {}
    sheet = SimpleNamespace(name="I. Control")

    def range_at(first: Any, *args: Any) -> Any:
        if isinstance(first, str):
            assert "L7" not in first, "Do not fall back to installation anchors"
            match = re.fullmatch(r"\$?U\$?(\d+)(?::\$?X\$?\d+)?", first)
            assert match, f"Unexpected range lookup: {first}"
            first = (int(match.group(1)), 21)
        return anchors[first]

    sheet.range = range_at
    definitions = {}
    for offset, (field, name) in enumerate(names.items()):
        row = 101 + offset * 3
        anchor = SimpleNamespace(value=field, row=row, column=21, sheet=sheet)
        anchor.cells = {(0, 0): anchor}
        anchor.api = SimpleNamespace(Row=row, Column=21)
        anchors[(row, 21)] = anchor
        merged = SimpleNamespace(
            Row=row, Column=21, Address=f"$U${row}:$X${row}",
            Cells=lambda *args, cell=anchor: cell.api,
        )
        region = SimpleNamespace(
            row=row, column=23, sheet=sheet,
            api=SimpleNamespace(
                MergeCells=True, MergeArea=merged,
                Cells=lambda *args, area=merged: SimpleNamespace(MergeArea=area),
            ),
        )
        definitions[name] = SimpleNamespace(refers_to_range=region)
    caller = SimpleNamespace(names=definitions, sheets={"I. Control": sheet})
    monkeypatch.setattr(workflow, "require_stage_panel", lambda book: None)
    resolved = workflow._stage_field_cells(caller)
    assert set(resolved) == set(names)
    assert {field: cell.value for field, cell in resolved.items()} == {
        field: field for field in names
    }
    assert all(cell.column == 21 for cell in resolved.values())


def test_stage_one_audit_interfaces_exist() -> None:
    """Expose the approved panel and handoff contract before implementation."""
    from bonds_lite_input import (
        STAGE_ONE_FIELDS,
        _generate_bonds_lite_input,
        _read_stage_fields,
        initialise_stage_panel,
        portable_output_path,
        require_stage_panel,
        resolve_stage_input,
        run_stage_one,
        write_stage_audit,
    )

    assert STAGE_ONE_FIELDS == {
        "last_run": "L74", "start": "L75", "finish": "L76",
        "elapsed": "L77", "user": "L78", "status": "L79",
        "last_output": "L80", "stage_two_input": "L82",
    }
    assert all(callable(helper) for helper in (
        _generate_bonds_lite_input, _read_stage_fields, initialise_stage_panel,
        portable_output_path, require_stage_panel, resolve_stage_input,
        run_stage_one, write_stage_audit,
    ))


@pytest.fixture
def stage_attempt(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """Observe owned audit boundaries and saves, never external Excel objects."""
    import getpass
    import subprocess
    import bonds_lite_input as workflow
    from bonds_lite_input import run_stage_one

    caller_path = tmp_path / "Draft.xlsm"
    caller_path.write_bytes(b"Caller sentinel")
    folder = tmp_path / "Bonds Lite output"
    folder.mkdir()
    previous = _complete_run(folder / "20261027_090000")
    output = folder / "20261028_090000"
    old_relative = "Bonds Lite output/20261027_090000"
    state = SimpleNamespace(
        fields={"last_output": old_relative, "stage_two_input": old_relative,
                "status": "Succeeded", "finish": datetime(2026, 10, 27),
                "elapsed": 0.5},
        saves=[], writes=[], events=[], locked=False, guard_calls=[],
        clock=100.0, fail_saves=set(), backend_error=None, panel_error=None,
        output=output, previous=previous, caller_path=caller_path,
        old_relative=old_relative, run=run_stage_one,
    )
    caller = SimpleNamespace(fullname=str(caller_path), api=SimpleNamespace(Saved=True))
    state.caller = caller

    def save() -> None:
        assert state.locked, "Every caller save belongs inside the run guard"
        state.events.append("save")
        state.saves.append(deepcopy(state.fields))
        if len(state.saves) in state.fail_saves:
            caller.api.Saved = False
            raise BondsLiteError("Audit save sentinel")
        caller.api.Saved = True

    caller.save = save

    @contextmanager
    def guard(path: Path) -> Iterator[None]:
        state.guard_calls.append(path)
        assert not state.locked, "Do not acquire a nested workbook guard"
        assert Path(path) == caller_path
        state.locked = True
        state.events.append("lock")
        try:
            yield
        finally:
            state.events.append("release")
            state.locked = False

    def require(book: Any) -> None:
        assert book is caller and state.locked
        state.events.append("require")
        if state.panel_error:
            raise state.panel_error

    def read(book: Any) -> dict[str, object]:
        assert book is caller and state.locked
        state.events.append("read")
        return deepcopy(state.fields)

    def write(book: Any, fields: dict[str, object]) -> None:
        assert book is caller and state.locked
        assert set(fields) <= {
            "last_run", "start", "finish", "elapsed", "user", "status",
            "last_output", "stage_two_input",
        }
        state.events.append("write")
        state.writes.append(deepcopy(fields))
        state.fields.update(fields)
        caller.save()
        if not caller.api.Saved:
            raise BondsLiteError("Audit save sentinel: Saved is false")

    def backend(book: Any) -> Path:
        assert book is caller and state.locked
        assert state.saves[-1]["status"] == "Running"
        assert not state.saves[-1]["stage_two_input"]
        assert not state.saves[-1]["finish"]
        assert not state.saves[-1]["elapsed"]
        assert state.saves[-1]["last_output"] == old_relative
        state.events.append("backend")
        try:
            state.clock = 105.5
            if state.backend_error:
                raise state.backend_error
            _complete_run(output)
            return output
        finally:
            state.events.append("cleanup")

    class AuditDatetime(datetime):
        @classmethod
        def now(cls, tz: Any = None) -> datetime:
            return datetime(2026, 10, 28, 9, 0 if state.clock == 100.0 else 3)

    monkeypatch.setattr(getpass, "getuser", lambda: "audit-user")
    monkeypatch.setattr(workflow, "datetime", AuditDatetime)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: state.clock)
    monkeypatch.setattr(workflow, "workbook_run_guard", guard)
    monkeypatch.setattr(workflow, "require_stage_panel", require)
    monkeypatch.setattr(workflow, "_read_stage_fields", read)
    monkeypatch.setattr(workflow, "write_stage_audit", write)
    monkeypatch.setattr(workflow, "_generate_bonds_lite_input", backend)
    monkeypatch.setattr(
        subprocess, "Popen",
        lambda *args, **kwargs: pytest.fail("Stage 1 must not execute Stage 2"),
    )
    monkeypatch.setattr(
        workflow, "generate_bonds_lite_input",
        lambda *args: pytest.fail("Wrapper must call the private backend"),
    )
    return state


def test_run_stage_one_saves_running_then_success_after_cleanup(stage_attempt: Any) -> None:
    state = stage_attempt
    assert state.run(state.caller) == state.output
    assert state.guard_calls == [state.caller_path]
    assert state.events == [
        "lock", "require", "read", "write", "save", "backend", "cleanup",
        "write", "save", "release",
    ]
    running, success = state.saves
    assert running["status"] == "Running"
    assert running["user"] == "audit-user"
    assert running["start"] == datetime(2026, 10, 28, 9, 0)
    assert isinstance(running["last_run"], datetime)
    assert success["start"] == running["start"]
    assert success["finish"] == datetime(2026, 10, 28, 9, 3)
    assert success["elapsed"] == pytest.approx(5.5 / 86400)
    assert success["status"] == "Succeeded"
    assert success["last_output"] == "Bonds Lite output/20261028_090000"
    assert success["stage_two_input"] == success["last_output"]
    assert _pair_bytes(state.output) == _pair_bytes(state.previous)
    assert state.caller_path.read_bytes() == b"Caller sentinel"
    assert not state.locked


@pytest.mark.parametrize("seconds", [2.0, 90061.0])
def test_run_stage_one_elapsed_handles_midnight_and_more_than_one_day(
    stage_attempt: Any, monkeypatch: pytest.MonkeyPatch, seconds: float,
) -> None:
    """Store full timestamps and duration days without time-of-day subtraction."""
    from datetime import timedelta

    import bonds_lite_input as workflow

    state = stage_attempt
    started = datetime(2026, 10, 28, 23, 59, 59)
    backend = workflow._generate_bonds_lite_input

    class RunDatetime(datetime):
        @classmethod
        def now(cls, tz: Any = None) -> datetime:
            return started if state.clock == 100.0 else started + timedelta(seconds=seconds)

    def generate(caller: Any) -> Path:
        output = backend(caller)
        state.clock = 100.0 + seconds
        return output

    monkeypatch.setattr(workflow, "datetime", RunDatetime)
    monkeypatch.setattr(workflow, "_generate_bonds_lite_input", generate)
    assert state.run(state.caller) == state.output
    assert state.fields["start"] == started
    assert state.fields["finish"] == started + timedelta(seconds=seconds)
    assert state.fields["elapsed"] == pytest.approx(seconds / 86400)


@pytest.mark.parametrize("error", [BondsLiteError("Backend sentinel"),
                                  RuntimeError("Backend sentinel"), KeyboardInterrupt()])
@pytest.mark.parametrize("audit_failure", [False, True])
def test_run_stage_one_failure_clears_handoff_preserves_success_and_original_error(
    stage_attempt: Any, error: BaseException, audit_failure: bool,
) -> None:
    state = stage_attempt
    state.backend_error = error
    if audit_failure:
        state.fail_saves.add(2)
    with pytest.raises(type(error)) as caught:
        state.run(state.caller)
    assert caught.value is error
    terminal = state.saves[-1]
    assert terminal["status"] == (
        "Interrupted" if isinstance(error, KeyboardInterrupt) else "Failed"
    )
    assert terminal["finish"] == datetime(2026, 10, 28, 9, 3)
    assert terminal["elapsed"] == pytest.approx(5.5 / 86400)
    assert not terminal["stage_two_input"]
    assert terminal["last_output"] == state.old_relative
    assert set(_pair_bytes(state.previous)) == set(EXPECTED_INPUT_FILENAMES.values())
    assert not state.output.exists()
    assert state.events.index("cleanup") < len(state.events) - 2
    assert state.events[-1] == "release"
    assert len(state.guard_calls) == 1


def test_run_stage_one_initial_audit_save_failure_does_not_generate(stage_attempt: Any) -> None:
    state = stage_attempt
    state.fail_saves.add(1)
    with pytest.raises(BondsLiteError, match="Audit save sentinel"):
        state.run(state.caller)
    assert "backend" not in state.events
    assert not state.fields["stage_two_input"]
    assert state.fields["last_output"] == state.old_relative
    assert not state.output.exists()
    assert not state.locked


@pytest.mark.parametrize("recovery_save_fails", [False, True])
def test_run_stage_one_final_save_failure_retains_published_output_without_handoff(
    stage_attempt: Any, recovery_save_fails: bool, caplog: pytest.LogCaptureFixture,
) -> None:
    state = stage_attempt
    state.fail_saves = {2, 3} if recovery_save_fails else {2}
    with caplog.at_level(logging.INFO, logger="bonds_lite_input"):
        with pytest.raises(BondsLiteError) as caught:
            state.run(state.caller)
    message = str(caught.value)
    assert "audit" in message.lower() and "retain" in message.lower()
    assert state.output.name in message
    assert "Audit save sentinel" in message
    assert _pair_bytes(state.output) == _pair_bytes(state.previous)
    assert state.fields["status"] == "Failed"
    assert not state.fields["stage_two_input"]
    assert len(state.writes) >= 3
    assert state.previous.is_dir()
    assert not state.locked
    assert "Saved Bonds Lite input" not in caplog.text


def test_run_stage_one_requires_panel_under_lock_before_any_write(stage_attempt: Any) -> None:
    state = stage_attempt
    state.panel_error = BondsLiteError("Panel missing sentinel")
    with pytest.raises(BondsLiteError, match="Panel missing"):
        state.run(state.caller)
    assert state.events == ["lock", "require", "release"]
    assert not state.saves and not state.writes


def test_run_stage_one_duplicate_guard_rejection_never_writes(
    stage_attempt: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow
    state = stage_attempt

    @contextmanager
    def reject(path: Path) -> Iterator[None]:
        assert path == state.caller_path
        raise BondsLiteError("Already in progress sentinel")
        yield

    monkeypatch.setattr(workflow, "workbook_run_guard", reject)
    with pytest.raises(BondsLiteError, match="Already in progress"):
        state.run(state.caller)
    assert state.events == []
    assert not state.writes and not state.saves


def test_generate_bonds_lite_input_guards_private_backend_without_audit_or_save(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow
    from bonds_lite_input import _generate_bonds_lite_input

    events: list[str] = []
    path = tmp_path / "Draft.xlsm"
    caller = SimpleNamespace(fullname=str(path))
    output = _complete_run(tmp_path / "Bonds Lite output" / "20261028_090000")

    @contextmanager
    def guard(identity: Path) -> Iterator[None]:
        assert identity == path
        events.append("lock")
        try:
            yield
        finally:
            events.append("release")

    def backend(book: Any) -> Path:
        assert book is caller
        assert events == ["lock"]
        events.append("backend")
        return output

    def forbidden(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Direct generation must not read/write the audit or save caller")

    caller.save = forbidden
    monkeypatch.setattr(workflow, "workbook_run_guard", guard)
    monkeypatch.setattr(workflow, "_generate_bonds_lite_input", backend)
    for name in ("require_stage_panel", "_read_stage_fields", "write_stage_audit"):
        monkeypatch.setattr(workflow, name, forbidden)
    assert callable(_generate_bonds_lite_input)
    assert workflow.generate_bonds_lite_input(caller) == output
    assert events == ["lock", "backend", "release"]


@pytest.mark.parametrize("username", ["Joao Branco", "Sarah Walker"])
@pytest.mark.parametrize("suffix", ["", "_001", "_1001"])
def test_portable_output_path_round_trips_for_different_users(
    tmp_path: Path,
    username: str,
    suffix: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow
    from bonds_lite_input import portable_output_path

    folder = tmp_path / "Users" / username / "OneDrive - Just Group" / "Testing"
    output = _complete_run(folder / "Bonds Lite output" / f"20261028_090000{suffix}")
    relative = portable_output_path(folder / "Draft.xlsm", output)
    assert relative == f"Bonds Lite output/20261028_090000{suffix}"
    assert not Path(relative).is_absolute()
    assert ".." not in Path(relative).parts and "\\" not in relative
    assert (folder / relative).resolve() == output.resolve()
    caller = SimpleNamespace(
        fullname=str(folder / "Draft.xlsm"), api=SimpleNamespace(Saved=True)
    )
    monkeypatch.setattr(workflow, "require_stage_panel", lambda book: None)
    monkeypatch.setattr(
        workflow,
        "_read_stage_fields",
        lambda book: {
            "status": "Succeeded",
            "stage_two_input": relative,
            "last_output": relative,
        },
    )
    assert workflow.resolve_stage_input(caller) == output.resolve()
    assert workflow.resolve_stage_inputs(caller) == {
        dataset: output.resolve() / filename
        for dataset, filename in EXPECTED_INPUT_FILENAMES.items()
    }


@pytest.mark.parametrize(
    "scenario",
    [
        "outside",
        "sibling_prefix",
        "missing",
        "empty",
        "legacy_file",
        "wrong_suffix",
        "temporary",
        "traversal",
        "nested",
        "missing_bonds",
        "missing_map",
        "empty_bonds",
        "empty_map",
        "directory_member",
        "short_suffix",
        "foreign_symlink",
        "bad_timestamp",
        "alpha_suffix",
    ],
)
def test_portable_output_path_rejects_unpublished_or_outside_files(
    tmp_path: Path,
    scenario: str,
) -> None:
    from bonds_lite_input import portable_output_path

    folder = tmp_path / "Bonds Lite output"
    folder.mkdir()
    output = folder / "20261028_090000"
    if scenario == "outside":
        output = tmp_path / "20261028_090000"
    elif scenario == "sibling_prefix":
        output = tmp_path / "Bonds Lite output other" / "20261028_090000"
    elif scenario == "wrong_suffix":
        output = folder / "20261028_090000.xlsx"
    elif scenario == "temporary":
        output = folder / ".bonds-lite-pending"
    elif scenario == "traversal":
        _complete_run(folder / "20261028_090000")
        output = folder / "nested" / ".." / "20261028_090000"
        (folder / "nested").mkdir()
    elif scenario == "nested":
        output = folder / "nested" / "20261028_090000"
    elif scenario == "short_suffix":
        output = folder / "20261028_090000_01"
    elif scenario == "bad_timestamp":
        output = folder / "20261028-090000"
    elif scenario == "alpha_suffix":
        output = folder / "20261028_090000_abc"
    if scenario == "legacy_file":
        output = folder / "input.xlsx"
        output.write_bytes(b"Legacy single input")
    elif scenario == "empty":
        output.mkdir()
    elif scenario not in ("missing", "traversal"):
        _complete_run(output)
    if scenario in (
        "missing_bonds",
        "missing_map",
        "empty_bonds",
        "empty_map",
        "directory_member",
        "foreign_symlink",
    ):
        dataset = "Bonds" if scenario.endswith("bonds") else "MAP"
        member = output / EXPECTED_INPUT_FILENAMES[dataset]
        member.unlink()
        if scenario.startswith("empty"):
            member.touch()
        elif scenario == "directory_member":
            member.mkdir()
        elif scenario == "foreign_symlink":
            foreign = tmp_path / "foreign.xlsx"
            foreign.write_bytes(b"Foreign nonempty input")
            try:
                member.symlink_to(foreign)
            except OSError as error:
                pytest.skip(
                    f"File symlinks unavailable on this Windows profile: {error}"
                )
    with pytest.raises(BondsLiteError):
        portable_output_path(tmp_path / "Draft.xlsm", output)


@pytest.mark.parametrize(
    "scenario",
    [
        "success",
        "not_installed",
        "Running",
        "Failed",
        "Interrupted",
        "blank_status",
        "blank_input",
        "mismatch",
        "absolute",
        "outside",
        "traversal",
        "missing",
        "empty",
        "wrong_suffix",
        "legacy_file",
        "temporary",
        "traversal_inside",
        "unsaved",
        "missing_bonds",
        "missing_map",
        "empty_bonds",
        "empty_map",
    ],
)
@pytest.mark.parametrize("resolver", ["resolve_stage_input", "resolve_stage_inputs"])
def test_resolve_stage_input_gates_handoff_without_executing_stage_two(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    scenario: str,
    resolver: str,
) -> None:
    import bonds_lite_input as workflow

    resolve = getattr(workflow, resolver)

    caller = SimpleNamespace(
        fullname=str(tmp_path / "Draft.xlsm"), api=SimpleNamespace(Saved=True)
    )
    folder = tmp_path / "Bonds Lite output"
    folder.mkdir()
    output = _complete_run(folder / "20261028_090000")
    relative = "Bonds Lite output/20261028_090000"
    fields = {
        "status": "Succeeded",
        "stage_two_input": relative,
        "last_output": relative,
    }
    events: list[str] = []

    def require(book: Any) -> None:
        assert book is caller
        events.append("require")
        if scenario == "not_installed":
            raise BondsLiteError("Panel not installed")

    def read(book: Any) -> dict[str, object]:
        assert book is caller and events == ["require"]
        events.append("read")
        return fields.copy()

    if scenario == "unsaved":
        caller.api.Saved = False
    elif scenario in ("Running", "Failed", "Interrupted"):
        fields["status"] = scenario
    elif scenario == "blank_status":
        fields["status"] = ""
    elif scenario == "blank_input":
        fields["stage_two_input"] = ""
    elif scenario == "mismatch":
        fields["stage_two_input"] = "Bonds Lite output/20261027_090000"
    elif scenario in (
        "absolute",
        "outside",
        "traversal",
        "wrong_suffix",
        "temporary",
        "traversal_inside",
    ):
        invalid = {
            "absolute": str(output),
            "outside": "20261028_090000",
            "traversal": "Bonds Lite output/../20261028_090000",
            "wrong_suffix": "Bonds Lite output/20261028_090000.xlsx",
            "temporary": "Bonds Lite output/.bonds-lite-pending/20261028_090000",
            "traversal_inside": "Bonds Lite output/nested/../20261028_090000",
        }[scenario]
        if scenario != "absolute":
            invalid_file = tmp_path / invalid
            invalid_file.parent.mkdir(parents=True, exist_ok=True)
            if not invalid_file.exists():
                _complete_run(invalid_file)
        fields.update(stage_two_input=invalid, last_output=invalid)
    elif scenario == "missing":
        shutil.rmtree(output)
    elif scenario == "empty":
        for member in output.iterdir():
            member.unlink()
    elif scenario == "legacy_file":
        legacy = folder / "input.xlsx"
        legacy.write_bytes(b"Legacy output")
        fields.update(
            stage_two_input="Bonds Lite output/input.xlsx",
            last_output="Bonds Lite output/input.xlsx",
        )
    elif scenario.startswith(("missing_", "empty_")):
        dataset = "MAP" if scenario.endswith("map") else "Bonds"
        member = output / EXPECTED_INPUT_FILENAMES[dataset]
        member.unlink()
        if scenario.startswith("empty_"):
            member.touch()

    monkeypatch.setattr(workflow, "require_stage_panel", require)
    monkeypatch.setattr(workflow, "_read_stage_fields", read)
    monkeypatch.setattr(
        workflow,
        "write_stage_audit",
        lambda *args: pytest.fail("Resolving Stage 2 input must not write or save"),
    )
    import subprocess

    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *args, **kwargs: pytest.fail("Stage 2 must never execute"),
    )
    if scenario == "success":
        expected = (
            output.resolve()
            if resolver == "resolve_stage_input"
            else {
                dataset: output.resolve() / filename
                for dataset, filename in EXPECTED_INPUT_FILENAMES.items()
            }
        )
        assert resolve(caller) == expected
        assert events == ["require", "read"]
    else:
        with pytest.raises(
            BondsLiteError,
            match=r"(?i)rerun|run.*stage.*1" if scenario == "legacy_file" else None,
        ):
            resolve(caller)


def test_validate_bond_headers_accepts_raw_schema_and_destination_tail() -> None:
    """Accept only the raw prefix while retaining destination formula headers."""
    from bonds_lite_input import validate_bond_headers

    source = [f"Bond header {index}" for index in range(RAW_COLUMN_COUNT)]
    destination = source + [f"Formula {index}" for index in range(20)]
    originals = deepcopy((source, destination))
    assert validate_bond_headers(source, destination) is None
    assert (source, destination) == originals


@pytest.mark.parametrize("tail_count", [0, 20, 21])
def test_validate_bond_headers_accepts_exact_raw_prefix(tail_count: int) -> None:
    """Allow destination-only columns without weakening raw header matching."""
    from bonds_lite_input import validate_bond_headers

    source = [f"Bond header {index}" for index in range(RAW_COLUMN_COUNT)]
    destination = source + [f"Formula {index}" for index in range(tail_count)]
    assert validate_bond_headers(source, destination) is None


@pytest.mark.parametrize(
    "scenario",
    [
        "missing", "short_source", "empty_source", "extra", "legacy_aun",
        "short_target", "empty_target", "reordered", "renamed", "case",
        "whitespace",
    ],
)
def test_validate_bond_headers_rejects_schema_drift(scenario: str) -> None:
    """Reject wrong widths and strict-prefix mismatches without mutating inputs."""
    from bonds_lite_input import validate_bond_headers

    source = [f"Bond header {index}" for index in range(RAW_COLUMN_COUNT)]
    destination = source + [f"Formula {index}" for index in range(20)]
    if scenario == "missing":
        del source[600]
    elif scenario == "short_source":
        source = source[:-1]
        destination = source.copy()
    elif scenario == "empty_source":
        source = []
    elif scenario == "extra":
        source.append("Unexpected")
    elif scenario == "legacy_aun":
        source = destination.copy()
    elif scenario == "short_target":
        destination = source[:-1]
    elif scenario == "empty_target":
        destination = []
    elif scenario == "reordered":
        source[0], source[1] = source[1], source[0]
    elif scenario == "renamed":
        source[-1] = "Renamed final raw header"
    elif scenario == "case":
        source[0] = source[0].lower()
    elif scenario == "whitespace":
        source[0] += " "
    originals = deepcopy((source, destination))
    with pytest.raises(BondsLiteError, match="(?i)header|schema|column"):
        validate_bond_headers(source, destination)
    assert (source, destination) == originals


@pytest.mark.parametrize("invalid", [None, "", " \t ", 42, False, ["header"]])
@pytest.mark.parametrize("index", [0, RAW_COLUMN_COUNT - 1])
def test_validate_bond_headers_rejects_matching_malformed_headers(
    invalid: object, index: int,
) -> None:
    """Reject equal but malformed raw headers on both sides of the comparison."""
    from bonds_lite_input import validate_bond_headers

    source = [f"Bond header {number}" for number in range(RAW_COLUMN_COUNT)]
    source[index] = invalid
    destination = deepcopy(source) + ["Formula tail"]
    originals = deepcopy((source, destination))
    with pytest.raises(BondsLiteError, match="(?i)header|schema|column"):
        validate_bond_headers(source, destination)
    assert (source, destination) == originals


@pytest.fixture
def stress_values() -> dict[str, list[list[object]]]:
    """Build complete stress inputs with distinctive boundary values."""
    blocks: dict[str, list[list[object]]] = {}
    for case in BLOCK_CASES:
        grid: list[list[object]] = [
            [0.0] * case["columns"] for _ in range(case["rows"])
        ]
        grid[0][0] = case["first_value"]
        grid[-1][-1] = case["last_value"]
        blocks[case["source"]] = grid
    return blocks


@pytest.fixture
def local_paths(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    """Create local stand-ins without opening supplied workbooks."""
    paths = (
        tmp_path / "Draft caller.xlsm",
        tmp_path / "Bond data.xlsx",
        tmp_path / "MAP data.xlsx",
        tmp_path / "Input template.xlsx",
    )
    for path in paths:
        path.write_bytes(b"Read-only test stand-in")
    return paths


def test_bonds_lite_error_is_a_value_error() -> None:
    """Expose a contextual validation exception to callers."""
    assert issubclass(BondsLiteError, ValueError)


@pytest.fixture
def console_logger(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str],
) -> Iterator[logging.Logger]:
    """Isolate the module's handlers without changing the root logger."""
    class CaptureHandler(logging.StreamHandler):
        def emit(self, record: logging.LogRecord) -> None:
            self.stream = sys.stdout
            super().emit(record)

    logger = logging.getLogger("bonds_lite_input")
    monkeypatch.setattr(logger, "handlers", [])
    monkeypatch.setattr(logger, "level", logging.INFO)
    monkeypatch.setattr(logger, "propagate", False)
    monkeypatch.setattr(logger, "disabled", False)
    handler = CaptureHandler(sys.stdout)
    logger.addHandler(handler)
    try:
        yield logger
    finally:
        for attached in logger.handlers[:]:
            logger.removeHandler(attached)
            attached.close()


@pytest.fixture
def reporting_threads(monkeypatch: pytest.MonkeyPatch) -> list[Any]:
    """Record reporter lifecycle without starting threads or waiting."""
    threads: list[Any] = []

    class FakeThread:
        """Capture the reporting target and require shutdown before joining."""

        def __init__(
            self, target: Any = None, args: tuple = (),
            kwargs: dict[str, Any] | None = None, **options: Any,
        ) -> None:
            self.target = target
            self.args = args
            self.kwargs = kwargs or {}
            self.started = False
            self.joined = False
            threads.append(self)

        def start(self) -> None:
            self.started = True

        def join(self, timeout: float | None = None) -> None:
            values = (*self.args, *self.kwargs.values())
            events = [value for value in values if hasattr(value, "is_set")]
            assert len(events) == 1
            assert events[0].is_set(), "Stop the reporter before joining it"
            self.joined = True

        def is_alive(self) -> bool:
            return self.started and not self.joined

    monkeypatch.setattr(threading, "Thread", FakeThread)
    return threads


def _assert_logged_durations(message: str, *expected: float) -> None:
    """Check elapsed values without fixing whitespace, units or precision."""
    numbers = [float(value) for value in re.findall(r"-?\d+(?:\.\d+)?", message)]
    assert "total" in message.lower()
    for duration in expected:
        assert any(value == pytest.approx(duration) for value in numbers), message


def test_configure_console_logging_is_stdout_only_and_idempotent(
    console_logger: logging.Logger, capsys: pytest.CaptureFixture[str],
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Keep one owned INFO handler and preserve unrelated/root configuration."""
    import bonds_lite_input as workflow

    monkeypatch.chdir(tmp_path)
    unrelated = console_logger.handlers[0]
    root = logging.getLogger()
    root_handlers = root.handlers[:]
    root_level = root.level
    console_logger.setLevel(logging.WARNING)
    workflow.configure_console_logging()
    workflow.configure_console_logging()
    owned = [handler for handler in console_logger.handlers if handler is not unrelated]
    assert len(owned) == 1
    assert isinstance(owned[0], logging.StreamHandler)
    assert not isinstance(owned[0], logging.FileHandler)
    assert owned[0].stream is sys.stdout
    assert owned[0].level == logging.INFO
    assert console_logger.level == logging.INFO
    assert unrelated in console_logger.handlers
    assert root.handlers == root_handlers
    assert root.level == root_level
    console_logger.removeHandler(unrelated)
    flushes: list[bool] = []
    flush = owned[0].flush

    def record_flush() -> None:
        flushes.append(True)
        flush()

    monkeypatch.setattr(owned[0], "flush", record_flush)
    console_logger.info("Console sentinel")
    assert flushes, "Progress output must be flushed immediately"
    captured = capsys.readouterr()
    assert captured.out.count("Console sentinel") == 1
    assert captured.err == ""
    assert not list(tmp_path.iterdir())


def test_timed_stage_logs_nested_order_and_stage_total_durations(
    monkeypatch: pytest.MonkeyPatch, console_logger: logging.Logger,
    reporting_threads: list[Any], capsys: pytest.CaptureFixture[str],
) -> None:
    """Use one processing clock while keeping nested stage times separate."""
    import bonds_lite_input as workflow

    clock = SimpleNamespace(now=100.0)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: clock.now)
    with workflow.run_progress() as state:
        assert state.started == 100.0
        clock.now = 102.0
        with workflow.timed_stage("Outer"):
            entry = capsys.readouterr().out
            assert "outer" in entry.lower() and "start" in entry.lower()
            _assert_logged_durations(entry, 2.0)
            clock.now = 105.0
            with workflow.timed_stage("Inner"):
                entry = capsys.readouterr().out
                assert "inner" in entry.lower() and "start" in entry.lower()
                _assert_logged_durations(entry, 5.0)
                clock.now = 109.0
            completed = capsys.readouterr().out
            assert "inner" in completed.lower() and "done" in completed.lower()
            _assert_logged_durations(completed, 4.0, 9.0)
            clock.now = 112.0
        completed = capsys.readouterr().out
        assert "outer" in completed.lower() and "done" in completed.lower()
        _assert_logged_durations(completed, 10.0, 12.0)


@pytest.mark.parametrize("error_type", [RuntimeError, KeyboardInterrupt])
def test_timed_stage_logs_failure_and_rethrows_original(
    error_type: type[BaseException], monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger, reporting_threads: list[Any],
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Keep original failures authoritative and remove the failed active stage."""
    import bonds_lite_input as workflow

    clock = SimpleNamespace(now=20.0)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: clock.now)
    original = error_type("Save sentinel")
    with workflow.run_progress() as state:
        clock.now = 22.0
        with pytest.raises(error_type) as caught:
            with workflow.timed_stage("Save"):
                capsys.readouterr()
                clock.now = 27.0
                raise original
        assert caught.value is original
        message = capsys.readouterr().out
        assert "save" in message.lower() and "failed" in message.lower()
        assert "Save sentinel" in message
        assert "done" not in message.lower()
        _assert_logged_durations(message, 5.0, 7.0)
        with state.lock:
            assert not state.stages


@pytest.mark.parametrize("interval", [5.0, 1.25])
@pytest.mark.parametrize("error_type", [None, RuntimeError, KeyboardInterrupt])
def test_run_progress_stops_and_joins_one_reporter_on_every_exit(
    error_type: type[BaseException] | None, interval: float,
    monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger, reporting_threads: list[Any],
) -> None:
    """Bound the reporter to a run and pass only plain progress state to it."""
    import bonds_lite_input as workflow

    monkeypatch.setattr(workflow.time, "perf_counter", lambda: 40.0)
    original = error_type("Run sentinel") if error_type else None
    options = {} if interval == 5.0 else {"heartbeat_interval": interval}
    try:
        with workflow.run_progress(**options) as state:
            assert state.started == 40.0
            assert len(reporting_threads) == 1
            reporter = reporting_threads[0]
            assert reporter.started and not reporter.joined
            values = (*reporter.args, *reporter.kwargs.values())
            assert any(value is state for value in values)
            assert interval in values
            assert len(values) == 3
            with state.lock:
                assert not state.stages
            if original is not None:
                raise original
        if original is not None:
            pytest.fail("The progress context suppressed the original failure")
    except BaseException as caught:
        if original is None or caught is not original:
            raise
    assert len(reporting_threads) == 1
    assert reporting_threads[0].joined
    assert not reporting_threads[0].is_alive()


@pytest.mark.parametrize("active_stage", [False, True])
def test_report_waiting_uses_current_stage_clock_and_stops_without_sleeping(
    active_stage: bool, monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger, capsys: pytest.CaptureFixture[str],
) -> None:
    """Report waiting, not completion, and resume the parent after a stage exits."""
    import bonds_lite_input as workflow

    state = SimpleNamespace(
        started=100.0, lock=threading.Lock(),
        stages=[("Outer", 101.0), ("Inner", 104.0)] if active_stage else [],
    )
    clock = SimpleNamespace(now=109.0)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: clock.now)

    class StopEvent:
        """Advance two reporting ticks, then request immediate termination."""

        def __init__(self) -> None:
            self.waits: list[float] = []

        def wait(self, interval: float) -> bool:
            self.waits.append(interval)
            if len(self.waits) == 2:
                clock.now = 112.0
                with state.lock:
                    if state.stages:
                        state.stages.pop()
            assert len(self.waits) <= 3, "Reporter ignored its stop event"
            return len(self.waits) == 3

        def is_set(self) -> bool:
            return len(self.waits) >= 3

    stop = StopEvent()
    workflow._report_waiting(stop, state, 5.0)
    messages = [
        line for line in capsys.readouterr().out.splitlines()
        if "waiting" in line.lower()
    ]
    assert stop.waits == [5.0, 5.0, 5.0]
    assert len(messages) == 2
    _assert_logged_durations(messages[0], 9.0)
    _assert_logged_durations(messages[1], 12.0)
    if active_stage:
        assert "Inner" in messages[0]
        assert "Outer" in messages[1] and "Inner" not in messages[1]
        _assert_logged_durations(messages[0], 5.0)
        _assert_logged_durations(messages[1], 11.0)
    assert all(
        "%" not in message and "done" not in message.lower()
        for message in messages
    )


def test_launcher_uses_configured_python_without_an_addin() -> None:
    assert "BONDS_LITE_PYTHON" in VBA_LAUNCHER
    assert 'Environment("User")' in VBA_LAUNCHER
    assert "--excel-hwnd" in VBA_LAUNCHER
    assert "ThisWorkbook.FullName" in VBA_LAUNCHER
    assert "RunPython" not in VBA_LAUNCHER
    assert "shell.Run command, 1, False" in VBA_LAUNCHER
    assert "-u -m bonds_lite_input" in VBA_LAUNCHER
    assert "--pause" in VBA_LAUNCHER


@pytest.mark.parametrize("username", ["Joao Branco", "Sarah Walker"])
@pytest.mark.parametrize("cloud_path", [True, False])
@pytest.mark.parametrize("matching", [True, False])
def test_main_matches_native_caller_identity(
    monkeypatch: pytest.MonkeyPatch,
    username: str,
    cloud_path: bool,
    matching: bool,
    console_logger: logging.Logger,
    reporting_threads: list[Any],
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Match VBA's native identity without replacing the user's local path."""
    import bonds_lite_input as workflow

    local_path = (
        f"C:\\Users\\{username}\\OneDrive - Just Group\\Testing\\Draft.xlsm"
    )
    native_path = (
        "https://example.sharepoint.com/Shared Documents/Testing/Draft.xlsm"
        if cloud_path else local_path
    )
    caller_argument = native_path if matching else native_path.replace("Testing", "Other")
    book = SimpleNamespace(
        fullname=local_path, api=SimpleNamespace(FullName=native_path)
    )
    application = SimpleNamespace(api=SimpleNamespace(Hwnd=1234), books=[book])
    monkeypatch.setitem(sys.modules, "xlwings", SimpleNamespace(apps=[application]))
    monkeypatch.setattr(
        sys, "argv",
        ["bonds_lite_input", "--caller", caller_argument, "--excel-hwnd", "1234"],
    )
    received = []

    def generate(caller: Any) -> Path:
        received.append(caller)
        return Path("Bonds Lite output/20261028_090000")

    def forbidden(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Ordinary CLI calls must not pause or show a message box")

    monkeypatch.setattr(workflow, "run_stage_one", generate, raising=False)
    monkeypatch.setattr(workflow, "generate_bonds_lite_input", forbidden)
    monkeypatch.setattr(workflow, "notify_user", forbidden, raising=False)
    monkeypatch.setattr("builtins.input", forbidden)

    assert workflow.main() == (0 if matching else 1)
    assert received == ([book] if matching else [])
    assert book.fullname == local_path
    captured = capsys.readouterr()
    console = captured.out + captured.err
    assert "total" in console.lower()
    if matching:
        assert "saved" in console.lower()
        assert "20261028_090000" in console
    else:
        assert "failed" in console.lower()
        assert "exact calling workbook" in console
    assert all(reporter.joined for reporter in reporting_threads)


@pytest.mark.parametrize("failure", [False, True], ids=["success", "failure"])
@pytest.mark.parametrize("pause_mode", ["default", "enter", "eof"])
def test_main_reports_total_and_cleans_up_before_optional_pause(
    failure: bool, pause_mode: str, monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger, reporting_threads: list[Any],
    capsys: pytest.CaptureFixture[str], tmp_path: Path,
) -> None:
    """Keep summaries readable without pausing noninteractive calls or holding work."""
    import bonds_lite_input as workflow

    monkeypatch.chdir(tmp_path)
    clock = SimpleNamespace(now=50.0)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: clock.now)
    caller = SimpleNamespace(
        fullname="Draft.xlsm", api=SimpleNamespace(FullName="Draft.xlsm")
    )
    application = SimpleNamespace(api=SimpleNamespace(Hwnd=1234), books=[caller])
    monkeypatch.setitem(sys.modules, "xlwings", SimpleNamespace(apps=[application]))
    arguments = [
        "bonds_lite_input", "--caller", "Draft.xlsm", "--excel-hwnd", "1234"
    ]
    if pause_mode != "default":
        arguments.append("--pause")
    monkeypatch.setattr(sys, "argv", arguments)
    lifecycle: list[str] = []
    prompts: list[str] = []
    guard_active = False

    @contextmanager
    def cli_guard(path: Path) -> Iterator[None]:
        nonlocal guard_active
        assert path == Path(caller.fullname)
        assert not guard_active
        guard_active = True
        try:
            yield
        finally:
            guard_active = False
            lifecycle.append("release")

    def generate(book: Any) -> Path:
        assert book is caller
        with workflow.workbook_run_guard(Path(book.fullname)):
            lifecycle.append("generate")
            try:
                clock.now = 57.0
                if failure:
                    raise BondsLiteError("Save failed: disk full sentinel")
                return Path("Bonds Lite output/20261028_090000")
            finally:
                lifecycle.append("cleanup")

    def assert_summary() -> None:
        captured = capsys.readouterr()
        console = captured.out + captured.err
        _assert_logged_durations(console, 7.0)
        if failure:
            assert "failed" in console.lower()
            assert "disk full sentinel" in console
        else:
            assert "saved" in console.lower()
            assert "20261028_090000" in console
        assert lifecycle == ["generate", "cleanup", "release"]
        assert not guard_active
        assert len(reporting_threads) == 1
        assert reporting_threads[0].joined

    def read_enter(prompt: str = "") -> str:
        assert pause_mode != "default", "Pause must be explicitly requested"
        assert_summary()
        prompts.append(prompt)
        clock.now = 999.0
        if pause_mode == "eof":
            raise EOFError("No interactive console")
        return ""

    def forbidden_notification(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Console completion must not also display a modal message box")

    monkeypatch.setattr(workflow, "run_stage_one", generate, raising=False)
    monkeypatch.setattr(workflow, "workbook_run_guard", cli_guard)
    monkeypatch.setattr(
        workflow, "generate_bonds_lite_input", forbidden_notification
    )
    monkeypatch.setattr(workflow, "notify_user", forbidden_notification, raising=False)
    monkeypatch.setattr("builtins.input", read_enter)
    assert workflow.main() == (1 if failure else 0)
    if pause_mode == "default":
        assert_summary()
        assert not prompts
    else:
        assert len(prompts) == 1
    assert not list(tmp_path.iterdir())


def test_main_install_button_is_noninteractive_by_default(
    monkeypatch: pytest.MonkeyPatch, console_logger: logging.Logger,
    reporting_threads: list[Any], capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
) -> None:
    """Keep installer CLI calls noninteractive without invoking Excel."""
    import bonds_lite_input as workflow

    caller = tmp_path / "Draft.xlsm"
    backup = tmp_path / "backup.xlsm"
    received: list[Path] = []

    def install(path: Path) -> Path:
        received.append(path)
        return backup

    def forbidden(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Installation must not prompt, notify, or generate output")

    monkeypatch.setattr(
        sys, "argv", ["bonds_lite_input", "--install-button", str(caller)]
    )
    monkeypatch.setattr(workflow, "install_button", install)
    monkeypatch.setattr(workflow, "run_stage_one", forbidden, raising=False)
    monkeypatch.setattr(workflow, "generate_bonds_lite_input", forbidden)
    monkeypatch.setattr(workflow, "notify_user", forbidden, raising=False)
    monkeypatch.setattr("builtins.input", forbidden)
    assert workflow.main() == 0
    assert received == [caller]
    assert str(backup) in capsys.readouterr().out
    assert all(reporter.joined for reporter in reporting_threads)
    assert not list(tmp_path.iterdir())


def test_workbook_mutex_is_stable_and_does_not_expose_paths(tmp_path: Path) -> None:
    caller = tmp_path / "Caller.xlsm"
    name = workbook_mutex_name(caller)
    assert name.startswith("Local\\BondsLite_")
    assert "Caller.xlsm" not in name
    assert name == workbook_mutex_name(tmp_path / "." / "Caller.xlsm")
    assert name != workbook_mutex_name(tmp_path / "Another.xlsm")


def test_stress_blocks_match_all_twelve_expected_ranges() -> None:
    """Keep the complete ordered mapping explicit and independently sourced."""
    assert isinstance(STRESS_BLOCKS, tuple)
    assert all(isinstance(block, str) for block in STRESS_BLOCKS)
    assert STRESS_BLOCKS == tuple(case["source"] for case in BLOCK_CASES)
    assert len(set(STRESS_BLOCKS)) == 12
    assert sum(case["rows"] * case["columns"] for case in BLOCK_CASES) == 5880


@pytest.mark.parametrize(
    ("number", "expected"),
    [(1, "A"), (3, "C"), (10, "J"), (13, "M"), (26, "Z"),
     (27, "AA"), (52, "AZ"), (53, "BA"), (113, "DI"),
     (116, "DL"), (216, "HH"), (1236, "AUN"), (16384, "XFD")],
)
def test_column_name_converts_one_based_columns(number: int, expected: str) -> None:
    assert column_name(number) == expected


@pytest.mark.parametrize("number", [0, -1])
def test_column_name_rejects_nonpositive_columns(number: int) -> None:
    with pytest.raises(ValueError):
        column_name(number)


@pytest.mark.parametrize("case", BLOCK_CASES, ids=lambda case: case["source"])
def test_validate_stress_values_maps_boundaries_without_scaling(
    case: dict[str, Any], stress_values: dict[str, list[list[object]]]
) -> None:
    """Retain first/last coordinates and raw decimal units in every block."""
    result = validate_stress_values(stress_values)
    grid = result[case["destination"]]
    assert len(grid) == case["rows"]
    assert all(len(row) == case["columns"] for row in grid)
    assert grid[0][0] == case["first_value"]
    assert grid[-1][-1] == case["last_value"]
    first_column, first_row = re.fullmatch(
        r"([A-Z]+)(\d+)", case["first_cell"]
    ).groups()
    column_number = 0
    for character in first_column:
        column_number = column_number * 26 + ord(character) - ord("A") + 1
    assert case["source"] == f'{case["first_cell"]}:{case["last_cell"]}'
    assert case["last_cell"] == (
        f'{column_name(column_number + case["columns"] - 1)}'
        f'{int(first_row) + case["rows"] - 1}'
    )


def test_validate_stress_values_preserves_every_cell_and_inputs(
    stress_values: dict[str, list[list[object]]],
) -> None:
    """Accept zero without imposing probability sums or transforming spreads."""
    stress_values["C9:J15"][1][1] = 0
    stress_values["M9:DI15"][2][2] = -2
    original = deepcopy(stress_values)
    result = validate_stress_values(stress_values)
    assert result == original
    assert stress_values == original
    assert all(
        type(value) is float
        for grid in result.values()
        for row in grid
        for value in row
    )


@pytest.mark.parametrize(
    ("invalid", "reason"),
    [(None, r"blank|missing|empty"), ("", r"blank|missing|empty"),
     (" \t", r"blank|missing|empty"), (True, r"bool|numeric|number"),
     (False, r"bool|numeric|number"), ("0.01", r"numeric|number|text|string"),
     ("bad", r"numeric|number|text|string"), ("#N/A", r"error|numeric|number"),
     ("#VALUE!", r"error|numeric|number"), ("#REF!", r"error|numeric|number"),
     (float("nan"), r"finite|nan"), (float("inf"), r"finite|inf"),
     (-float("inf"), r"finite|inf"), (1 + 2j, r"real|numeric|number|complex")],
)
def test_validate_stress_values_rejects_invalid_numeric_cells(
    invalid: object, reason: str, stress_values: dict[str, list[list[object]]]
) -> None:
    stress_values["M9:DI15"][0][0] = invalid
    with pytest.raises(ValueError) as caught:
        validate_stress_values(stress_values)
    message = str(caught.value)
    assert "M9" in message
    assert re.search(reason, message, re.IGNORECASE)


def test_validate_stress_values_aggregates_locations_across_all_blocks(
    stress_values: dict[str, list[list[object]]],
) -> None:
    for case in BLOCK_CASES:
        stress_values[case["source"]][0][0] = None
    stress_values["DL42:HH48"][-1][-1] = True
    with pytest.raises(ValueError) as caught:
        validate_stress_values(stress_values)
    message = str(caught.value)
    for case in BLOCK_CASES:
        assert re.search(rf'\b{case["first_cell"]}\b', message)
    assert "HH48" in message
    assert re.search(r"blank|missing|empty", message, re.IGNORECASE)
    assert re.search(r"bool|numeric|number", message, re.IGNORECASE)


@pytest.mark.parametrize("case", BLOCK_CASES, ids=lambda case: case["source"])
def test_validate_stress_values_rejects_each_missing_block(
    case: dict[str, Any], stress_values: dict[str, list[list[object]]]
) -> None:
    del stress_values[case["source"]]
    with pytest.raises(ValueError) as caught:
        validate_stress_values(stress_values)
    assert case["source"] in str(caught.value)


def test_validate_stress_values_aggregates_all_missing_blocks() -> None:
    with pytest.raises(ValueError) as caught:
        validate_stress_values({})
    for case in BLOCK_CASES:
        assert case["source"] in str(caught.value)


@pytest.mark.parametrize("case", BLOCK_CASES, ids=lambda case: case["source"])
@pytest.mark.parametrize("defect", ["fewer_rows", "extra_rows", "short_row", "long_row"])
def test_validate_stress_values_rejects_wrong_block_shapes(
    case: dict[str, Any], defect: str,
    stress_values: dict[str, list[list[object]]],
) -> None:
    grid = stress_values[case["source"]]
    if defect == "fewer_rows":
        grid.pop()
    elif defect == "extra_rows":
        grid.append(list(grid[0]))
    elif defect == "short_row":
        grid[3].pop()
    else:
        grid[3].append(0.0)
    with pytest.raises(ValueError) as caught:
        validate_stress_values(stress_values)
    assert case["source"] in str(caught.value)
    assert re.search(r"shape|row|column|size|dimension", str(caught.value), re.I)


@pytest.mark.parametrize("relative", [True, False])
@pytest.mark.parametrize("caller_as_string", [True, False])
def test_resolve_paths_uses_saved_caller_folder_not_working_directory(
    local_paths: tuple[Path, Path, Path, Path], tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch, relative: bool, caller_as_string: bool,
) -> None:
    caller, data, map_data, template = local_paths
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)
    caller_argument = str(caller) if caller_as_string else caller
    result = resolve_paths(
        caller_argument,
        data.name if relative else str(data),
        map_data.name if relative else str(map_data),
        template.name if relative else str(template),
    )
    assert result == local_paths
    assert isinstance(result, tuple)
    assert all(isinstance(path, Path) and path.is_absolute() for path in result)


@pytest.mark.parametrize("position", [0, 1, 2, 3])
def test_resolve_paths_rejects_missing_files(
    local_paths: tuple[Path, Path, Path, Path], position: int,
) -> None:
    arguments = [str(path) for path in local_paths]
    missing = local_paths[0].parent / "Missing workbook.xlsx"
    arguments[position] = str(missing)
    with pytest.raises(ValueError) as caught:
        resolve_paths(*arguments)
    assert missing.name in str(caught.value)


@pytest.mark.parametrize("position", [0, 1, 2, 3])
@pytest.mark.parametrize("url", ["https://example.invalid/book.xlsx", "http://example.invalid/book.xlsx"])
def test_resolve_paths_rejects_urls(
    local_paths: tuple[Path, Path, Path, Path], position: int, url: str,
) -> None:
    arguments = [str(path) for path in local_paths]
    arguments[position] = url
    with pytest.raises(ValueError) as caught:
        resolve_paths(*arguments)
    assert re.search(r"url|local|sync|https?", str(caught.value), re.I)


@pytest.mark.parametrize("position", [0, 1, 2, 3])
@pytest.mark.parametrize("blank", ["", "   "])
def test_resolve_paths_rejects_unsaved_or_blank_paths(
    local_paths: tuple[Path, Path, Path, Path], position: int, blank: str,
) -> None:
    arguments = [str(path) for path in local_paths]
    arguments[position] = blank
    with pytest.raises(ValueError):
        resolve_paths(*arguments)


@pytest.mark.parametrize("position", [0, 1, 2, 3])
def test_resolve_paths_rejects_directories_as_workbooks(
    local_paths: tuple[Path, Path, Path, Path], position: int,
) -> None:
    arguments = [str(path) for path in local_paths]
    arguments[position] = str(local_paths[0].parent)
    with pytest.raises(ValueError):
        resolve_paths(*arguments)


@pytest.mark.parametrize(
    ("left", "right"), [(0, 1), (0, 2), (0, 3), (1, 2), (1, 3), (2, 3)]
)
def test_resolve_paths_rejects_every_pairwise_alias(
    local_paths: tuple[Path, Path, Path, Path], left: int, right: int,
) -> None:
    arguments = [str(path) for path in local_paths]
    arguments[right] = str(local_paths[left].parent / "." / local_paths[left].name)
    with pytest.raises(ValueError) as caught:
        resolve_paths(*arguments)
    assert re.search(r"same|alias|distinct|different", str(caught.value), re.I)


def test_resolve_paths_rejects_relative_alias_after_normalisation(
    local_paths: tuple[Path, Path, Path, Path],
) -> None:
    caller, data, _, template = local_paths
    with pytest.raises(ValueError):
        resolve_paths(
            caller, data.name, f"../{caller.parent.name}/{data.name}", template.name
        )


@pytest.mark.parametrize("alias_kind", ["hardlink", "symlink", "case"])
@pytest.mark.parametrize(
    ("left", "right"), [(0, 1), (0, 2), (0, 3), (1, 2), (1, 3), (2, 3)]
)
def test_resolve_paths_rejects_existing_file_identity_aliases(
    local_paths: tuple[Path, Path, Path, Path],
    alias_kind: str,
    left: int,
    right: int,
) -> None:
    arguments = list(local_paths)
    if alias_kind == "case":
        arguments[right] = Path(str(arguments[left]).swapcase())
        if not arguments[right].is_file():
            pytest.skip("Case-insensitive filesystem required")
    else:
        alias = local_paths[0].parent / f"{alias_kind} alias.xlsx"
        try:
            if alias_kind == "hardlink":
                alias.hardlink_to(arguments[left])
            else:
                alias.symlink_to(arguments[left])
        except OSError as error:
            pytest.skip(f"{alias_kind} creation unavailable: {error}")
        arguments[right] = alias
    with pytest.raises(BondsLiteError, match=r"(?i)same|alias|distinct|different"):
        resolve_paths(*arguments)


def test_validate_headers_accepts_exact_schema_without_mutation() -> None:
    source = CASES["schema"]["valid"].copy()
    destination = CASES["schema"]["destination"].copy()
    assert validate_headers(source, destination) is None
    assert source == CASES["schema"]["valid"]
    assert destination == CASES["schema"]["destination"]


@pytest.mark.parametrize(
    "case", CASES["schema"]["invalid"], ids=lambda case: case["name"]
)
def test_validate_headers_rejects_schema_drift(case: dict[str, Any]) -> None:
    with pytest.raises(ValueError) as caught:
        validate_headers(case["headers"], CASES["schema"]["destination"])
    assert re.search(r"header|schema|column", str(caught.value), re.I)


def test_validate_identifiers_accepts_nonblank_values_without_new_constraints() -> None:
    values = CASES["identifiers"]["valid"].copy()
    assert validate_identifiers(values) is None
    assert values == CASES["identifiers"]["valid"]


@pytest.mark.parametrize(
    "case", CASES["identifiers"]["invalid"], ids=lambda case: case["name"]
)
def test_validate_identifiers_rejects_empty_data_or_blank_records(
    case: dict[str, Any],
) -> None:
    with pytest.raises(ValueError) as caught:
        validate_identifiers(case["values"])
    assert re.search(r"identifier|record|row|empty|bond", str(caught.value), re.I)


@pytest.mark.parametrize("collisions", [0, 1, 2])
def test_choose_output_path_uses_timestamp_and_preserves_existing_files(
    tmp_path: Path, collisions: int,
) -> None:
    case = CASES["output"]
    names = [case["base"], case["first_collision"], case["second_collision"]]
    for name in names[:collisions]:
        (tmp_path / name).write_bytes(b"Earlier successful output")
    before = {path.name: path.read_bytes() for path in tmp_path.iterdir()}
    result = choose_output_path(tmp_path, datetime.fromisoformat(case["timestamp"]))
    assert result == tmp_path / names[collisions]
    assert not result.exists()
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


def test_choose_output_path_continues_after_many_collisions(tmp_path: Path) -> None:
    case = CASES["output"]
    (tmp_path / case["base"]).touch()
    for suffix in range(1, 12):
        (tmp_path / f"Bonds Lite input_20261027_140506_{suffix:03d}.xlsx").touch()
    result = choose_output_path(tmp_path, datetime.fromisoformat(case["timestamp"]))
    assert result.name == "Bonds Lite input_20261027_140506_012.xlsx"
    assert not result.exists()


def test_choose_output_path_treats_existing_directory_as_collision(tmp_path: Path) -> None:
    case = CASES["output"]
    (tmp_path / case["base"]).mkdir()
    result = choose_output_path(tmp_path, datetime.fromisoformat(case["timestamp"]))
    assert result == tmp_path / case["first_collision"]


@pytest.fixture
def publication_case(tmp_path: Path) -> tuple[Path, Path, datetime]:
    """Use disposable file-operation stand-ins, not workbook data."""
    folder = tmp_path / "Bonds Lite output"
    folder.mkdir()
    working_folder = folder / ".working"
    working_folder.mkdir()
    working = working_folder / "working.xlsx"
    working.write_bytes(b"Temporary working workbook stand-in")
    return working, folder, datetime.fromisoformat(CASES["output"]["timestamp"])


@pytest.mark.parametrize("collisions", [0, 1, 2])
def test_publish_output_moves_file_without_overwriting(
    publication_case: tuple[Path, Path, datetime], collisions: int,
) -> None:
    """Publish the same file under the first unused timestamped name."""
    from bonds_lite_input import publish_output

    working, folder, timestamp = publication_case
    case = CASES["output"]
    names = [case["base"], case["first_collision"], case["second_collision"]]
    earlier = {folder / name: f"Earlier output {name}".encode() for name in names[:collisions]}
    for path, content in earlier.items():
        path.write_bytes(content)
    content = working.read_bytes()
    original = working.stat()

    result = publish_output(working, folder, timestamp)

    assert isinstance(result, Path)
    assert result == folder / names[collisions]
    assert result.read_bytes() == content
    assert (result.stat().st_dev, result.stat().st_ino) == (
        original.st_dev, original.st_ino,
    )
    assert not working.exists()
    assert {path: path.read_bytes() for path in earlier} == earlier
    assert set(folder.iterdir()) == {working.parent, result, *earlier}


@pytest.mark.parametrize("races", [1, 2])
def test_publish_output_retries_rename_collisions_without_deleting_outputs(
    publication_case: tuple[Path, Path, datetime],
    monkeypatch: pytest.MonkeyPatch, races: int,
) -> None:
    """A competing publisher can claim a name after it was selected."""
    from bonds_lite_input import publish_output

    working, folder, timestamp = publication_case
    previous = folder / CASES["output"]["base"]
    previous.write_bytes(b"Earlier successful output")
    original_rename = Path.rename
    attempts: list[Path] = []
    competing: dict[Path, bytes] = {previous: previous.read_bytes()}
    content = working.read_bytes()

    def racing_rename(source: Path, target: Path) -> Path:
        assert source == working
        assert target.parent == folder
        assert not target.exists()
        attempts.append(target)
        if len(attempts) <= races:
            competing[target] = f"Competing output {len(attempts)}".encode()
            with target.open("xb") as stream:
                stream.write(competing[target])
            raise FileExistsError(f"Destination already exists: {target}")
        return original_rename(source, target)

    monkeypatch.setattr(Path, "rename", racing_rename)

    result = publish_output(working, folder, timestamp)

    assert attempts == [
        folder / f"Bonds Lite input_20261027_140506_{suffix:03d}.xlsx"
        for suffix in range(1, races + 2)
    ]
    assert result == attempts[-1]
    assert result.read_bytes() == content
    assert not working.exists()
    assert {path: path.read_bytes() for path in competing} == competing
    assert set(folder.iterdir()) == {working.parent, result, *competing}


@pytest.mark.parametrize("race_first", [False, True])
def test_publish_output_propagates_rename_failure_and_preserves_files(
    publication_case: tuple[Path, Path, datetime],
    monkeypatch: pytest.MonkeyPatch, race_first: bool,
) -> None:
    """Failed publication must not clean up an earlier or competing output."""
    from bonds_lite_input import publish_output

    working, folder, timestamp = publication_case
    previous = folder / CASES["output"]["base"]
    previous.write_bytes(b"Earlier successful output")
    preserved = {previous: previous.read_bytes(), working: working.read_bytes()}
    attempts: list[Path] = []
    denied = PermissionError("Publication denied")

    def failing_rename(source: Path, target: Path) -> Path:
        assert source == working
        assert target.parent == folder
        assert not target.exists()
        attempts.append(target)
        if race_first and len(attempts) == 1:
            preserved[target] = b"Competing successful output"
            with target.open("xb") as stream:
                stream.write(preserved[target])
            raise FileExistsError(f"Destination already exists: {target}")
        raise denied

    monkeypatch.setattr(Path, "rename", failing_rename)

    with pytest.raises(PermissionError, match="Publication denied") as caught:
        publish_output(working, folder, timestamp)

    assert caught.value is denied
    assert attempts == [
        folder / CASES["output"]["first_collision"],
        *([folder / CASES["output"]["second_collision"]] if race_first else []),
    ]
    assert {path: path.read_bytes() for path in preserved} == preserved
    assert set(folder.iterdir()) == {working.parent, *preserved.keys() - {working}}


@pytest.mark.parametrize("kind", ["missing", "directory"])
def test_publish_output_rejects_invalid_working_path_without_side_effects(
    publication_case: tuple[Path, Path, datetime], kind: str,
) -> None:
    """Reject non-files before publishing or disturbing existing outputs."""
    from bonds_lite_input import publish_output

    working, folder, timestamp = publication_case
    working.unlink()
    if kind == "directory":
        working.mkdir()
        (working / "keep.txt").write_bytes(b"Directory contents")
    previous = folder / CASES["output"]["base"]
    previous.write_bytes(b"Earlier successful output")
    before = {
        path.relative_to(folder): path.read_bytes() if path.is_file() else None
        for path in folder.rglob("*")
    }

    with pytest.raises(
        (BondsLiteError, FileNotFoundError, IsADirectoryError),
        match=re.escape(working.name),
    ):
        publish_output(working, folder, timestamp)

    assert {
        path.relative_to(folder): path.read_bytes() if path.is_file() else None
        for path in folder.rglob("*")
    } == before


def test_pure_workflow_contract_validates_before_reserving_paired_run(
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict[str, list[list[object]]],
) -> None:
    """Validate inputs before reserving a new, empty paired-run directory."""
    from bonds_lite_input import allocate_run_folder

    caller, data, map_data, template = local_paths
    before = {path: path.read_bytes() for path in local_paths}
    metadata = {path: path.stat().st_mtime_ns for path in local_paths}
    resolved = resolve_paths(caller, data.name, map_data.name, template.name)
    validate_headers(CASES["schema"]["valid"], CASES["schema"]["destination"])
    validate_identifiers(CASES["identifiers"]["valid"])
    validated = validate_stress_values(stress_values)
    output_folder = resolved[0].parent / "Bonds Lite output"
    output = allocate_run_folder(
        output_folder, datetime.fromisoformat(CASES["output"]["timestamp"])
    )
    assert validated == stress_values
    assert output == output_folder / "20261027_140506"
    assert output not in resolved
    assert output.is_dir() and not list(output.iterdir())
    assert {path: path.read_bytes() for path in local_paths} == before
    assert {path: path.stat().st_mtime_ns for path in local_paths} == metadata


@pytest.mark.parametrize("calculation_state", [0, 2])
def test_refresh_does_not_wait_for_pending_manual_calculation(
    monkeypatch: pytest.MonkeyPatch, calculation_state: int,
) -> None:
    """Pending manual work is not active asynchronous Excel calculation."""
    import bonds_lite_input as workflow

    calculated: list[bool] = []

    def forbidden_wait() -> None:
        pytest.fail("Do not poll when Excel is not actively calculating")

    monkeypatch.setitem(
        sys.modules, "pythoncom", SimpleNamespace(PumpWaitingMessages=forbidden_wait)
    )
    monkeypatch.setitem(
        sys.modules, "pywintypes", SimpleNamespace(com_error=RuntimeError)
    )
    monkeypatch.setattr(workflow.time, "monotonic", lambda: 0.0)
    monkeypatch.setattr(workflow, "required_links", lambda *args: [])
    table = SimpleNamespace(Range=SimpleNamespace(Row=8), ListRows=SimpleNamespace(Count=1))
    sheet = SimpleNamespace(
        name="I. Bonds Data",
        api=SimpleNamespace(
            Calculate=lambda: calculated.append(True),
            ListObjects=lambda name: table,
            Evaluate=lambda expression: 0,
        ),
    )

    class Sheets(dict):
        def __iter__(self):
            return iter(self.values())

    working = SimpleNamespace(
        sheets=Sheets({
            "I. Bonds Data": sheet,
            "P. Base Curves": _fake_model_base_curve(),
        }),
        app=SimpleNamespace(api=SimpleNamespace(CalculationState=calculation_state)),
    )

    workflow.refresh_and_validate(working, [2])

    assert calculated == [True]


def _fake_model_base_curve(error_count: int = 0, numeric_count: int = 3) -> Any:
    curve_range = SimpleNamespace(Address="$C$6:$C$8")
    table = SimpleNamespace(
        ListRows=SimpleNamespace(Count=3),
        ListColumns=lambda name: SimpleNamespace(DataBodyRange=curve_range),
    )

    def evaluate(expression: str) -> int:
        return error_count if "ISERROR" in expression else numeric_count

    return SimpleNamespace(
        name="P. Base Curves",
        api=SimpleNamespace(
            Calculate=lambda: None,
            ListObjects=lambda name: table,
            Evaluate=evaluate,
        ),
    )


def test_refresh_rejects_non_numeric_model_base_curve_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    monkeypatch.setitem(
        sys.modules, "pythoncom", SimpleNamespace(PumpWaitingMessages=lambda: None)
    )
    monkeypatch.setitem(
        sys.modules, "pywintypes", SimpleNamespace(com_error=RuntimeError)
    )
    monkeypatch.setattr(workflow.time, "monotonic", lambda: 0.0)
    monkeypatch.setattr(workflow, "required_links", lambda *args: [])
    bond_table = SimpleNamespace(
        Range=SimpleNamespace(Row=8), ListRows=SimpleNamespace(Count=1)
    )
    bond_sheet = SimpleNamespace(
        name="I. Bonds Data",
        api=SimpleNamespace(
            Calculate=lambda: None,
            ListObjects=lambda name: bond_table,
            Evaluate=lambda expression: 0,
        ),
    )

    class Sheets(dict):
        def __iter__(self):
            return iter(self.values())

    working = SimpleNamespace(
        sheets=Sheets({
            "I. Bonds Data": bond_sheet,
            "P. Base Curves": _fake_model_base_curve(error_count=1, numeric_count=2),
        }),
        app=SimpleNamespace(api=SimpleNamespace(CalculationState=0)),
    )

    with pytest.raises(BondsLiteError, match=r"(?i)GBP SONIA.*numeric|numeric.*GBP SONIA"):
        workflow.refresh_and_validate(working, [2])


@pytest.mark.parametrize("status", [0, 3, 1, 2, 4, 5, 6, 7, 8, 9, 10])
@pytest.mark.parametrize("formula_errors", [0, 1])
def test_refresh_warns_only_for_old_status_and_still_rejects_formula_errors(
    status: int, formula_errors: int, tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch, console_logger: logging.Logger,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """An old-status warning cannot bypass refresh or formula validation."""
    import bonds_lite_input as workflow

    dependency = tmp_path / "Dependency.xlsx"
    dependency.write_bytes(b"Saved dependency stand-in")
    updated: list[str] = []
    monkeypatch.setitem(
        sys.modules, "pythoncom", SimpleNamespace(PumpWaitingMessages=lambda: None)
    )
    monkeypatch.setitem(
        sys.modules, "pywintypes", SimpleNamespace(com_error=RuntimeError)
    )
    monkeypatch.setattr(workflow.time, "monotonic", lambda: 0.0)
    monkeypatch.setattr(workflow, "required_links", lambda *args: [str(dependency)])
    table = SimpleNamespace(Range=SimpleNamespace(Row=8), ListRows=SimpleNamespace(Count=1))
    sheet = SimpleNamespace(
        name="I. Bonds Data",
        api=SimpleNamespace(
            Calculate=lambda: None, ListObjects=lambda name: table,
            Evaluate=lambda expression: formula_errors,
        ),
    )

    class Sheets(dict):
        def __iter__(self):
            return iter(self.values())

    working = SimpleNamespace(
        sheets=Sheets({
            "I. Bonds Data": sheet,
            "P. Base Curves": _fake_model_base_curve(),
        }),
        app=SimpleNamespace(api=SimpleNamespace(CalculationState=2), books=[]),
        api=SimpleNamespace(
            UpdateLink=lambda Name, Type: updated.append(Name),
            LinkInfo=lambda *args: status,
        ),
    )
    if status in (0, 3) and not formula_errors:
        workflow.refresh_and_validate(working, [2])
    else:
        with pytest.raises(BondsLiteError) as caught:
            workflow.refresh_and_validate(working, [2])
        if status in (0, 3):
            assert "formula errors" in str(caught.value)
        else:
            assert f"status {status}" in str(caught.value)
    assert updated == [str(dependency)]
    assert ("Warning:" in capsys.readouterr().out) is (status == 3)


@pytest.mark.parametrize("entity", ["JRL", "Just Retirement Limited", "JRL-MAP_2"])
def test_validate_entity_preserves_valid_names(entity: str) -> None:
    from bonds_lite_input import validate_entity

    assert validate_entity(entity) == entity


@pytest.mark.parametrize(
    "entity",
    [
        None,
        42,
        True,
        "",
        " ",
        " JRL",
        "JRL ",
        "JRL\n",
        "J\x00RL",
        "J\x1fRL",
        "J\x7fRL",
        *[f"J{character}RL" for character in '<>:"/\\|?*'],
    ],
)
def test_validate_entity_rejects_unsafe_names_without_renaming(entity: Any) -> None:
    from bonds_lite_input import validate_entity

    with pytest.raises(BondsLiteError, match=r"(?i)entity|D9"):
        validate_entity(entity)


@pytest.fixture
def input_controls(local_paths: tuple[Path, Path, Path, Path]) -> Any:
    """Expose counted live values at the owned validation boundary."""
    caller_path, bonds, map_data, template = local_paths
    values = {
        "D39": bonds.name,
        "D40": map_data.name,
        "D42": template.name,
        "D41": "",
        "D11": datetime(2026, 9, 30),
        "D9": "JRL",
    }
    reads: list[str] = []

    def control_range(address: str) -> Any:
        assert address in values, f"Unexpected control read: {address}"
        reads.append(address)
        return SimpleNamespace(value=values[address])

    caller = SimpleNamespace(
        fullname=str(caller_path),
        sheets={"I. Control": SimpleNamespace(range=control_range)},
        app=object(),
    )
    return SimpleNamespace(caller=caller, values=values, reads=reads, paths=local_paths)


def test_validate_inputs_captures_frozen_snapshot_once_from_d39_d40_d42(
    input_controls: Any,
    stress_values: dict,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from dataclasses import FrozenInstanceError, fields, is_dataclass

    import bonds_lite_input as workflow

    state = input_controls
    stress_reads: list[Any] = []
    block_reads: list[str] = []
    original_read = workflow.read_credit_stresses
    ratings = ["GOV", "AAA", "AA", "A", "BBB", "BB", "B"]
    header_values = {}
    for row in (9, 20, 31, 42):
        for column in ("B", "L", "DK"):
            header_values[f"{column}{row}:{column}{row + 6}"] = [[rating] for rating in ratings]
        header_values[f"C{row - 1}:J{row - 1}"] = [ratings + ["D"]]
        for first, last in (("M", "DI"), ("DL", "HH")):
            header_values[f"{first}{row - 1}:{last}{row - 1}"] = [list(range(101))]

    def stress_range(address: str) -> Any:
        block_reads.append(address)
        return SimpleNamespace(
            options=lambda **kwargs: SimpleNamespace(
                value=stress_values[address] if address in stress_values else header_values[address]
            )
        )

    state.caller.sheets["I. Credit Stresses"] = SimpleNamespace(range=stress_range)

    def read(book: Any) -> dict:
        assert book is state.caller
        stress_reads.append(book)
        return original_read(book)

    monkeypatch.setattr(workflow, "read_credit_stresses", read)
    settings = workflow.validate_inputs(state.caller)
    assert isinstance(settings, workflow.StageOneInputs)
    assert is_dataclass(settings)
    assert {field.name for field in fields(settings)} == {
        "caller_path",
        "bonds_path",
        "map_path",
        "template_path",
        "valuation_date",
        "entity",
        "stresses",
    }
    assert (
        settings.caller_path,
        settings.bonds_path,
        settings.map_path,
        settings.template_path,
    ) == state.paths
    assert settings.valuation_date == state.values["D11"]
    assert settings.entity == "JRL" and settings.stresses == stress_values
    assert sorted(state.reads) == sorted(["D39", "D40", "D42", "D11", "D9"])
    assert stress_reads == [state.caller]
    assert sorted(block_reads) == sorted([*workflow.STRESS_BLOCKS, *header_values])
    captured_stresses = deepcopy(settings.stresses)
    state.values.update(D9="Changed live entity", D11=datetime(2027, 1, 1))
    stress_values[workflow.STRESS_BLOCKS[0]][0][0] = 999.0
    assert settings.entity == "JRL" and settings.valuation_date == datetime(2026, 9, 30)
    assert settings.stresses == captured_stresses
    with pytest.raises(FrozenInstanceError):
        settings.entity = "Changed"


@pytest.mark.parametrize("address", ["D39", "D40", "D42"])
def test_validate_inputs_missing_required_control_creates_no_outputs(
    input_controls: Any,
    address: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    state = input_controls
    state.values[address] = ""
    state.values["D41"] = str(state.paths[3])
    before = {path: path.read_bytes() for path in state.paths}
    monkeypatch.setattr(workflow, "read_credit_stresses", lambda caller: {})
    monkeypatch.setattr(
        workflow,
        "excel_state",
        lambda *args: pytest.fail("Missing inputs must fail before Excel state setup"),
    )
    with pytest.raises(BondsLiteError):
        workflow._generate_bonds_lite_input(state.caller)
    assert not (state.paths[0].parent / "Bonds Lite output").exists()
    assert {path: path.read_bytes() for path in state.paths} == before


@pytest.mark.parametrize(
    ("address", "value"),
    [
        ("D9", "JRL/unsafe"),
        ("D9", " JRL"),
        ("D9", 42),
        ("D11", None),
        ("D11", "2026-09-30"),
    ],
)
def test_validate_inputs_rejects_invalid_snapshot_controls(
    input_controls: Any,
    monkeypatch: pytest.MonkeyPatch,
    address: str,
    value: Any,
) -> None:
    import bonds_lite_input as workflow

    input_controls.values[address] = value
    input_controls.values["D41"] = str(input_controls.paths[3])
    monkeypatch.setattr(workflow, "read_credit_stresses", lambda caller: {})
    with pytest.raises(BondsLiteError, match=rf"(?i){address}|entity|date"):
        workflow.validate_inputs(input_controls.caller)
    assert not (input_controls.paths[0].parent / "Bonds Lite output").exists()


def test_synchronise_controls_uses_snapshot_and_preserves_other_controls(
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict,
) -> None:
    from bonds_lite_input import StageOneInputs, synchronise_controls

    settings = StageOneInputs(*local_paths, datetime(2026, 9, 30), "JRL", stress_values)
    cells = {
        address: SimpleNamespace(value=f"Original {address}")
        for address in ("D10", "D13", "D19", "D11", "D9")
    }
    working = SimpleNamespace(
        sheets={"I. Control": SimpleNamespace(range=cells.__getitem__)}
    )
    assert synchronise_controls(settings, working) is None
    assert cells["D10"].value == settings.valuation_date
    assert cells["D13"].value == settings.entity
    for address in ("D19", "D11", "D9"):
        assert cells[address].value == f"Original {address}"


@pytest.mark.parametrize(
    "defect", [None, "date", "entity", "rows", "formula", "stress"]
)
def test_verify_saved_copy_compares_snapshot_not_live_caller(
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict,
    defect: str | None,
) -> None:
    from bonds_lite_input import StageOneInputs, verify_saved_copy

    settings = StageOneInputs(*local_paths, datetime(2026, 9, 30), "JRL", stress_values)
    table = SimpleNamespace(
        ListRows=SimpleNamespace(Count=2 if defect == "rows" else 1)
    )
    control = {
        "D10": SimpleNamespace(
            value=datetime(2027, 1, 1) if defect == "date" else settings.valuation_date
        ),
        "D13": SimpleNamespace(
            value="Wrong" if defect == "entity" else settings.entity
        ),
    }

    def sensitivity(address: str) -> Any:
        grid = deepcopy(settings.stresses[address])
        if defect == "stress":
            grid[0][0] += 1
        return SimpleNamespace(options=lambda **kwargs: SimpleNamespace(value=grid))

    book = SimpleNamespace(
        sheets={
            "I. Control": SimpleNamespace(range=control.__getitem__),
            "I. Sensitivity": SimpleNamespace(range=sensitivity),
            "I. Bonds Data": SimpleNamespace(
                api=SimpleNamespace(ListObjects=lambda name: table),
                range=lambda *args: SimpleNamespace(
                    api=SimpleNamespace(HasFormula=defect != "formula")
                ),
            ),
        }
    )
    if defect:
        with pytest.raises(BondsLiteError):
            verify_saved_copy(book, settings, 1, [10, 11], settings.stresses)
    else:
        assert verify_saved_copy(book, settings, 1, [10, 11], settings.stresses) is None


@pytest.mark.parametrize("collisions", [0, 1, 3, 1001])
def test_allocate_run_folder_reserves_without_overwriting(
    tmp_path: Path,
    collisions: int,
) -> None:
    from bonds_lite_input import allocate_run_folder

    root = tmp_path / "new parent" / "Bonds Lite output"
    stamp = "20261028_090102"
    for index in range(collisions):
        name = stamp if index == 0 else f"{stamp}_{index:03d}"
        previous = root / name
        previous.mkdir(parents=True)
        (previous / "sentinel").write_bytes(b"Do not overwrite")
    result = allocate_run_folder(root, datetime(2026, 10, 28, 9, 1, 2))
    expected = stamp if collisions == 0 else f"{stamp}_{collisions:03d}"
    assert result == root / expected and result.is_dir()
    assert list(result.iterdir()) == []
    assert len(list(root.iterdir())) == collisions + 1
    assert all(
        path.read_bytes() == b"Do not overwrite" for path in root.glob("*/sentinel")
    )


def test_allocate_run_folder_retries_atomic_mkdir_collision_only(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from bonds_lite_input import allocate_run_folder

    root = tmp_path / "Bonds Lite output"
    root.mkdir()
    original = Path.mkdir
    attempts: list[tuple[Path, bool]] = []

    def mkdir(
        path: Path, mode: int = 0o777, parents: bool = False, exist_ok: bool = False
    ) -> None:
        if path.parent == root:
            attempts.append((path, exist_ok))
            assert not exist_ok, "Reservation must atomically reject a competing owner"
            if len(attempts) == 1:
                original(path)
                (path / "competitor").write_bytes(b"Other owner")
                raise FileExistsError(str(path))
        original(path, mode=mode, parents=parents, exist_ok=exist_ok)

    monkeypatch.setattr(Path, "mkdir", mkdir)
    result = allocate_run_folder(root, datetime(2026, 10, 28, 9, 1, 2))
    assert result.name == "20261028_090102_001"
    assert len(attempts) == 2
    assert (attempts[0][0] / "competitor").read_bytes() == b"Other owner"


def test_allocate_run_folder_propagates_permission_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from bonds_lite_input import allocate_run_folder

    root = tmp_path / "Bonds Lite output"
    root.mkdir()
    original = Path.mkdir
    error = PermissionError("Run reservation denied")

    def mkdir(path: Path, *args: Any, **kwargs: Any) -> None:
        if path.parent == root:
            raise error
        original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", mkdir)
    with pytest.raises(PermissionError) as caught:
        allocate_run_folder(root, datetime(2026, 10, 28, 9, 1, 2))
    assert caught.value is error and list(root.iterdir()) == []


@pytest.mark.parametrize(
    "failure",
    [
        None,
        "Bonds",
        "MAP",
        "interrupt_bonds",
        "interrupt_map",
        "missing_map",
        "empty_map",
    ],
)
def test_generate_bonds_lite_input_pairs_snapshot_and_rolls_back_owned_attempt(
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict,
    monkeypatch: pytest.MonkeyPatch,
    failure: str | None,
) -> None:
    import bonds_lite_input as workflow

    settings = workflow.StageOneInputs(
        *local_paths, datetime(2026, 9, 30), "JRL", deepcopy(stress_values)
    )
    assert workflow.INPUT_FILENAMES == EXPECTED_INPUT_FILENAMES
    root = local_paths[0].parent / "Bonds Lite output"
    previous = _complete_run(root / "20261027_090000")
    previous_bytes = _pair_bytes(previous)
    originals = {path: path.read_bytes() for path in local_paths}
    caller = SimpleNamespace(fullname=str(local_paths[0]), app=object())
    events: list[str] = []
    calls: list[tuple[str, Path, Any, Path]] = []
    error = (
        KeyboardInterrupt("Dataset interruption")
        if failure and failure.startswith("interrupt")
        else BondsLiteError("Dataset generation sentinel")
    )

    @contextmanager
    def guard(path: Path) -> Iterator[None]:
        assert path == local_paths[0] and events == []
        events.append("lock")
        try:
            yield
        finally:
            events.append("release")

    def validate(book: Any) -> Any:
        assert book is caller and events == ["lock"]
        events.append("validate")
        return settings

    @contextmanager
    def excel_state(application: Any) -> Iterator[None]:
        assert application is caller.app
        events.append("state")
        try:
            yield
        finally:
            events.append("restore")

    def dataset(
        application: Any, dataset: str, data_path: Path, snapshot: Any, run_folder: Path
    ) -> Path:
        assert application is caller.app and snapshot is settings
        assert events[-1] in ("state", "Bonds")
        assert run_folder.parent == root and run_folder != previous
        assert (run_folder / "JRL Bonds").is_dir()
        assert (run_folder / "JRL MAP").is_dir()
        assert not list((run_folder / "JRL Bonds").iterdir())
        assert not list((run_folder / "JRL MAP").iterdir())
        if calls:
            assert run_folder is calls[0][3]
            assert _pair_bytes(previous) == previous_bytes
        calls.append((dataset, data_path, snapshot, run_folder))
        events.append(dataset)
        private = run_folder / f".private-{dataset}"
        private.mkdir()
        (private / "partial.xlsx").write_bytes(b"Incomplete input")
        if failure in (dataset, f"interrupt_{dataset.lower()}"):
            raise error
        shutil.rmtree(private)
        output = run_folder / EXPECTED_INPUT_FILENAMES[dataset]
        if failure != "missing_map" or dataset != "MAP":
            output.write_bytes(
                b"" if failure == "empty_map" and dataset == "MAP" else dataset.encode()
            )
        return output

    monkeypatch.setattr(workflow, "validate_inputs", validate)
    monkeypatch.setattr(workflow, "workbook_run_guard", guard)
    monkeypatch.setattr(workflow, "excel_state", excel_state)
    monkeypatch.setattr(workflow, "_generate_dataset_input", dataset)
    if failure:
        with pytest.raises(type(error)) as caught:
            workflow.generate_bonds_lite_input(caller)
        if failure not in ("missing_map", "empty_map"):
            assert caught.value is error
        assert set(root.iterdir()) == {previous}
    else:
        result = workflow.generate_bonds_lite_input(caller)
        assert result == calls[0][3]
        assert _pair_bytes(result) == {
            filename: dataset.encode()
            for dataset, filename in EXPECTED_INPUT_FILENAMES.items()
        }
        assert set(path.name for path in result.iterdir()) == {
            *EXPECTED_INPUT_FILENAMES.values(),
            "JRL Bonds",
            "JRL MAP",
        }
    expected_datasets = (
        ["Bonds"] if failure in ("Bonds", "interrupt_bonds") else ["Bonds", "MAP"]
    )
    assert [(call[0], call[1]) for call in calls] == [
        (dataset, local_paths[index])
        for index, dataset in enumerate(expected_datasets, 1)
    ]
    assert events == [
        "lock",
        "validate",
        "state",
        *expected_datasets,
        "restore",
        "release",
    ]
    assert _pair_bytes(previous) == previous_bytes
    assert {path: path.read_bytes() for path in local_paths} == originals


def test_paired_cleanup_failure_preserves_generation_error_and_blocks_complete_pair(
    local_paths: tuple[Path, Path, Path, Path], monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger, capsys: pytest.CaptureFixture[str],
) -> None:
    """Filesystem cleanup failure must not conceal the original MAP failure."""
    from contextlib import nullcontext

    import bonds_lite_input as workflow

    settings = workflow.StageOneInputs(
        *local_paths, datetime(2026, 9, 30), "JRL", {}
    )
    caller = SimpleNamespace(app=object())
    original = BondsLiteError("MAP generation failed sentinel")

    def generate(application: Any, dataset: str, data: Path, snapshot: Any, folder: Path) -> Path:
        path = folder / workflow.INPUT_FILENAMES[dataset]
        if dataset == "MAP":
            raise original
        path.write_bytes(b"Completed Bonds member")
        return path

    def denied_cleanup(path: Path) -> None:
        raise PermissionError("Cleanup denied sentinel")

    monkeypatch.setattr(workflow, "validate_inputs", lambda book: settings)
    monkeypatch.setattr(workflow, "excel_state", lambda application: nullcontext())
    monkeypatch.setattr(workflow, "_generate_dataset_input", generate)
    monkeypatch.setattr(workflow.shutil, "rmtree", denied_cleanup)
    with pytest.raises(BondsLiteError) as caught:
        workflow._generate_bonds_lite_input(caller)
    assert caught.value is original
    assert "Cleanup denied sentinel" in capsys.readouterr().out
    run = next((local_paths[0].parent / "Bonds Lite output").iterdir())
    with pytest.raises(BondsLiteError):
        workflow.portable_output_path(local_paths[0], run)


@pytest.mark.parametrize("dataset", ["Bonds", "MAP"])
@pytest.mark.parametrize(
    "failure",
    [
        None,
        "refresh",
        "save",
        "verify",
        "interrupt",
        "saved_false",
        "empty_save",
    ],
)
def test_generate_dataset_input_owns_only_source_and_unique_private_copy(
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict,
    monkeypatch: pytest.MonkeyPatch,
    dataset: str,
    failure: str | None,
) -> None:
    import bonds_lite_input as workflow

    settings = workflow.StageOneInputs(
        *local_paths, datetime(2026, 9, 30), "JRL", stress_values
    )
    run = local_paths[0].parent / "Bonds Lite output" / "20261028_090000"
    run.mkdir(parents=True)
    for label in EXPECTED_INPUT_FILENAMES:
        (run / f"JRL {label}").mkdir()
    source_path = settings.bonds_path if dataset == "Bonds" else settings.map_path
    originals = {path: path.read_bytes() for path in local_paths}
    application = object()
    source = object()
    working = SimpleNamespace(api=SimpleNamespace(Saved=True))
    opened: list[tuple[Path, bool]] = []
    source_copies: list[Path] = []
    events: list[str] = []
    error = (
        KeyboardInterrupt("Helper interruption")
        if failure == "interrupt"
        else BondsLiteError("Helper sentinel")
    )

    @contextmanager
    def owned_book(app: Any, path: Path, read_only: bool = True) -> Iterator[Any]:
        assert app is application
        opened.append((path, read_only))
        if read_only:
            assert path != source_path
            assert path.name == source_path.name
            assert path.is_file() and path.read_bytes() == originals[source_path]
            assert not path.is_relative_to(run)
            source_copies.append(path)
            book = source
        else:
            assert path.is_relative_to(run) and path.parent != run
            assert path.name != "working.xlsx"
            assert path.name not in EXPECTED_INPUT_FILENAMES.values()
            assert path.read_bytes() == originals[settings.template_path]
            book = working
        events.append("open_source" if read_only else "open_copy")
        try:
            yield book
        finally:
            events.append("close_source" if read_only else "close_copy")

    def populate(data: Any, book: Any) -> tuple[int, list[int]]:
        assert data is source and book is working
        events.append("populate")
        return (1 if dataset == "Bonds" else 2), [10, 11, 1217]

    def stress(*args: Any, **kwargs: Any) -> None:
        assert any(value is settings.stresses or value is settings for value in args)
        events.append("stress")

    def controls(snapshot: Any, book: Any) -> None:
        assert snapshot is settings and book is working
        events.append("controls")

    def output_folder(book: Any, folder: Path) -> None:
        assert book is working and folder == run / f"JRL {dataset}"
        events.append("output_folder")

    def refresh(book: Any, columns: list[int]) -> None:
        assert book is working and columns == [10, 11, 1217]
        assert "output_folder" in events and "controls" in events
        events.append("refresh")
        if failure in ("refresh", "interrupt"):
            raise error

    def save() -> None:
        events.append("save")
        if failure == "save":
            raise error
        opened[-1][0].write_bytes(b"" if failure == "empty_save" else b"Saved dataset")
        working.api.Saved = failure != "saved_false"

    def verify(
        book: Any,
        snapshot: Any,
        rows: int,
        columns: list[int],
        stresses: dict,
        model_output_folder: Path | None = None,
    ) -> None:
        assert (
            book is working and snapshot is settings and stresses is settings.stresses
        )
        assert rows == (1 if dataset == "Bonds" else 2) and columns == [10, 11, 1217]
        assert model_output_folder == run / f"JRL {dataset}"
        assert events[-1] == "save"
        assert not (run / EXPECTED_INPUT_FILENAMES[dataset]).exists()
        events.append("verify")
        if failure == "verify":
            raise error

    working.save = save
    monkeypatch.setattr(workflow, "owned_book", owned_book)
    monkeypatch.setattr(workflow, "populate_bond_data", populate)
    monkeypatch.setattr(workflow, "populate_credit_stresses", stress)
    monkeypatch.setattr(workflow, "synchronise_controls", controls)
    monkeypatch.setattr(workflow, "set_model_output_folder", output_folder)
    monkeypatch.setattr(workflow, "refresh_and_validate", refresh)
    monkeypatch.setattr(workflow, "verify_saved_copy", verify)
    final = run / EXPECTED_INPUT_FILENAMES[dataset]
    if failure:
        with pytest.raises(type(error)) as caught:
            workflow._generate_dataset_input(
                application, dataset, source_path, settings, run
            )
        if failure not in ("saved_false", "empty_save"):
            assert caught.value is error
        assert not final.exists()
    else:
        assert (
            workflow._generate_dataset_input(
                application, dataset, source_path, settings, run
            )
            == final
        )
        assert final.read_bytes() == b"Saved dataset"
        assert (
            events.index("save") < events.index("verify") < events.index("close_copy")
        )
    assert (
        len(opened) == 2
        and opened[0][0] != source_path
        and opened[0][0].name == source_path.name
        and opened[0][1] is True
        and opened[1][1] is False
    )
    assert len(source_copies) == 1 and not source_copies[0].exists()
    assert events.index("close_copy") < events.index("close_source")
    assert set(run.iterdir()) == {
        run / "JRL Bonds",
        run / "JRL MAP",
        *([final] if failure is None else []),
    }
    assert {path: path.read_bytes() for path in local_paths} == originals


@pytest.fixture
def excel_case(tmp_path: Path):
    """Use isolated Excel and copies; never save the supplied source template."""
    import xlwings as xw

    root = Path(__file__).resolve().parents[1]
    template = tmp_path / "Input template.xlsx"
    shutil.copy2(root / "Bonds Lite input template.xlsx", template)
    with xw.App(visible=False, add_book=True) as application:
        working = source = data_sheet = book = None
        resources: list[Any] = []
        try:
            application.api.AutomationSecurity = 3
            application.api.EnableEvents = False
            application.api.Calculation = -4135
            application.api.CalculateBeforeSave = False
            application.display_alerts = False
            working = application.books.open(str(template), update_links=False)
            data_sheet = working.sheets["I. Bonds Data"]
            source = application.books.add()
            source.sheets[0].name = "I. Bonds Data"
            headers = [data_sheet.range("B8:AUN8").value[:RAW_COLUMN_COUNT]]
            sample = [data_sheet.range("B9:AUN9").options(
                err_to_str=True
            ).value[:RAW_COLUMN_COUNT]]
            assert len(headers[0]) == len(sample[0]) == RAW_COLUMN_COUNT
            resources.extend((application, working, source, headers, sample, tmp_path))
            yield resources
        finally:
            try:
                with ExitStack() as cleanup:
                    for book in list(application.books):
                        cleanup.callback(book.close)
            finally:
                resources.clear()
                book = data_sheet = source = working = None


EXCEL_ENABLED = os.environ.get("BONDS_LITE_RUN_EXCEL_TESTS") == "1"
PERFORMANCE_ENABLED = os.environ.get("BONDS_LITE_RUN_PERFORMANCE_TESTS") == "1"


@pytest.fixture
def stage_panel_excel_book(tmp_path: Path) -> Iterator[Any]:
    """Own only a new local workbook and its isolated Excel application."""
    if os.environ.get("BONDS_LITE_RUN_STAGE_PANEL_EXCEL_TESTS") != "1":
        pytest.skip("Set BONDS_LITE_RUN_STAGE_PANEL_EXCEL_TESTS=1")
    import xlwings as xw
    from bonds_lite_input import initialise_stage_panel

    with xw.App(visible=False, add_book=True) as application:
        application.api.AutomationSecurity = 3
        application.api.EnableEvents = False
        application.api.Calculation = -4135
        application.api.CalculateBeforeSave = False
        application.display_alerts = False
        caller = application.books[0]
        try:
            sheet = caller.sheets[0]
            sheet.name = "I. Control"
            caller.save(str(tmp_path / "movable panel.xlsx"))
            initialise_stage_panel(caller)
            yield caller
        finally:
            for owned in list(application.books):
                owned.close()


def _stage_panel_contents(region: Any) -> tuple[Any, ...]:
    """Snapshot values, formatting and links without depending on panel addresses."""
    return (
        region.api.Value2,
        region.formula,
        tuple(
            (
                cell.number_format, cell.api.Style.Name, cell.api.Interior.Color,
                cell.api.Font.Name, cell.api.Font.Size, cell.api.Font.Bold,
                cell.api.Font.Color, cell.api.WrapText,
                cell.api.HorizontalAlignment, cell.api.VerticalAlignment,
                cell.api.MergeArea.Rows.Count, cell.api.MergeArea.Columns.Count,
                tuple(
                    (cell.api.Hyperlinks.Item(index).Address,
                     cell.api.Hyperlinks.Item(index).SubAddress)
                    for index in range(1, cell.api.Hyperlinks.Count + 1)
                ),
            )
            for cell in region
        ),
    )


@pytest.mark.excel
@pytest.mark.parametrize("resize", ["insert_row_and_column", "remerge"])
def test_excel_stage_panel_cut_resize_reopen_uses_names_only(
    stage_panel_excel_book: Any, tmp_path: Path, resize: str,
) -> None:
    """Moved and resized panels retain handoff and never recreate the old panel."""
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    application = caller.app
    sheet = caller.sheets["I. Control"]
    output = _complete_run(tmp_path / "Bonds Lite output" / "20261028_090000")
    relative = "Bonds Lite output/20261028_090000"
    fields = {
        "last_run": datetime(2026, 10, 28),
        "start": datetime(2026, 10, 28, 9),
        "finish": datetime(2026, 10, 28, 9, 0, 6),
        "elapsed": 6 / 86400, "user": "before-cut", "status": "Succeeded",
        "last_output": relative, "stage_two_input": relative,
    }
    workflow.write_stage_audit(caller, fields)
    sheet.range("H73:O83").api.Cut(Destination=sheet.range("Q100:X110").api)
    sheet.range("H73:O83").value = "Old position sentinel"
    old_area = _stage_panel_contents(sheet.range("H73:O83"))
    assert caller.names["SST_Stage1_Panel"].refers_to_range.address == "$Q$100:$X$110"
    assert workflow.require_stage_panel(caller) is None
    assert workflow._read_stage_fields(caller)["user"] == "before-cut"
    caller.names[workflow.STAGE_ONE_FIELD_NAMES["elapsed"]].refers_to = (
        "='I. Control'!$U$104:$X$104"
    )
    sheet.range("100:110").row_height = 24
    sheet.range("Q:X").column_width = 16
    if resize == "insert_row_and_column":
        sheet.range("103:103").api.Insert()
        sheet.range("V:V").api.Insert()
        assert caller.names["SST_Stage1_Panel"].refers_to_range.address == "$Q$100:$Y$111"
    else:
        sheet.range("U101:X101").unmerge()
        sheet.range("U105:X105").unmerge()
        sheet.range("U105:W105").merge()
        caller.names[workflow.STAGE_ONE_FIELD_NAMES["user"]].refers_to = (
            "='I. Control'!$V$105"
        )
    assert workflow.require_stage_panel(caller) is None
    cells = workflow._stage_field_cells(caller)
    assert cells["user"].column == 21
    fields["user"] = "after-resize"
    assert workflow._read_stage_fields(caller)["user"] == "before-cut"
    workflow.write_stage_audit(caller, fields)
    read_back = workflow._read_stage_fields(caller)
    for field, value in fields.items():
        assert read_back[field] == (pytest.approx(value) if field == "elapsed" else value)
    for field in ("last_output", "stage_two_input"):
        assert cells[field].api.Hyperlinks.Item(1).Address == relative
    assert _stage_panel_contents(sheet.range("H73:O83")) == old_area
    panel = caller.names["SST_Stage1_Panel"].refers_to_range
    installed = _stage_panel_contents(panel)
    names_before = [(name.name, name.refers_to) for name in caller.names]
    assert workflow.initialise_stage_panel(caller) is None
    assert _stage_panel_contents(panel) == installed
    assert [(name.name, name.refers_to) for name in caller.names] == names_before
    assert _stage_panel_contents(sheet.range("H73:O83")) == old_area
    caller.save()
    path = caller.fullname
    caller.close()
    reopened = application.books.open(path, update_links=False)
    assert workflow.require_stage_panel(reopened) is None
    reopened_values = workflow._read_stage_fields(reopened)
    for field, value in fields.items():
        assert reopened_values[field] == (
            pytest.approx(value) if field == "elapsed" else value
        )
    assert workflow.resolve_stage_input(reopened) == output.resolve()
    reopened_panel = reopened.names["SST_Stage1_Panel"].refers_to_range
    reopened_contents = _stage_panel_contents(reopened_panel)
    assert workflow.initialise_stage_panel(reopened) is None
    assert reopened.api.Saved, "Validation of a complete panel must not dirty it"
    assert _stage_panel_contents(reopened_panel) == reopened_contents
    assert [(name.name, name.refers_to) for name in reopened.names] == names_before
    assert _stage_panel_contents(reopened.sheets["I. Control"].range("H73:O83")) == old_area


@pytest.mark.excel
@pytest.mark.parametrize("defect", [
    "broken_pointer", "missing", "outside", "overlap", "same_merge",
    "partial_merge", "unmerged_multicell", "multi_panel", "multi_field",
    "foreign_sheet", "foreign_book", "foreign_panel", "read_only", "protected",
])
def test_excel_require_stage_panel_rejects_invalid_layout_before_write_or_save(
    stage_panel_excel_book: Any, defect: str,
) -> None:
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    sheet = caller.sheets["I. Control"]
    field_name = workflow.STAGE_ONE_FIELD_NAMES["user"]
    if defect == "broken_pointer":
        caller.names[field_name].refers_to = "=#REF!"
    elif defect == "missing":
        caller.names[field_name].delete()
    elif defect == "outside":
        caller.names[field_name].refers_to = "='I. Control'!$Q$1"
    elif defect in ("overlap", "same_merge"):
        caller.names[field_name].refers_to = (
            "='I. Control'!$L$79" if defect == "overlap" else "='I. Control'!$M$79"
        )
    elif defect == "partial_merge":
        sheet.range("L78:O78").unmerge()
        sheet.range("L78:P78").merge()
    elif defect == "unmerged_multicell":
        sheet.range("L78:O78").unmerge()
        caller.names[field_name].refers_to = "='I. Control'!$L$78:$M$78"
    elif defect in ("multi_panel", "multi_field"):
        name = "SST_Stage1_Panel" if defect == "multi_panel" else field_name
        first = "$H$73:$O$83" if defect == "multi_panel" else "$L$78"
        caller.names[name].refers_to = f"=('I. Control'!{first},'I. Control'!$Q$1)"
    elif defect == "foreign_sheet":
        other = caller.sheets.add("Other control")
        caller.names[field_name].refers_to = f"='{other.name}'!$L$78"
    elif defect in ("foreign_book", "foreign_panel"):
        foreign = caller.app.books.add()
        foreign.sheets[0].name = "I. Control"
        name = "SST_Stage1_Panel" if defect == "foreign_panel" else field_name
        address = "$H$73:$O$83" if defect == "foreign_panel" else "$L$78"
        caller.names[name].refers_to = f"='[{foreign.name}]I. Control'!{address}"
    elif defect == "protected":
        sheet.api.Protect()
    saves: list[bool] = []
    boundary = _ExcelForwarder(caller, save=lambda: saves.append(True))
    if defect == "read_only":
        boundary = _ExcelForwarder(
            caller, save=lambda: saves.append(True),
            api=_ExcelForwarder(caller.api, ReadOnly=True),
        )
    before = _stage_panel_contents(sheet.range("H73:P83"))
    names_before = [(name.name, name.refers_to) for name in caller.names]
    saved_before = caller.api.Saved
    for action in (
        lambda: workflow.require_stage_panel(boundary),
        lambda: workflow.write_stage_audit(boundary, {"status": "Must not write"}),
        lambda: workflow.initialise_stage_panel(boundary),
    ):
        with pytest.raises(BondsLiteError) as caught:
            action()
        assert str(caught.value).strip()
        assert _stage_panel_contents(sheet.range("H73:P83")) == before
        assert [(name.name, name.refers_to) for name in caller.names] == names_before
        assert caller.api.Saved == saved_before
    assert saves == []


@pytest.mark.excel
@pytest.mark.parametrize("defect", [
    None, "shape", "pointer", "broken_pointer", "outside_pointer", "missing_legacy",
])
def test_excel_initialise_stage_panel_upgrades_only_coherent_moved_legacy_names(
    stage_panel_excel_book: Any, defect: str | None,
) -> None:
    """Infer missing anchors only from an unchanged legacy panel's relative layout."""
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    sheet = caller.sheets["I. Control"]
    relative = "Bonds Lite output/20261027_090000"
    workflow.write_stage_audit(caller, {
        "last_run": datetime(2026, 10, 28),
        "start": datetime(2026, 10, 28, 8),
        "finish": datetime(2026, 10, 28, 11),
        "status": "Succeeded", "user": "legacy-user", "elapsed": 0.125,
        "last_output": relative, "stage_two_input": relative,
    })
    sheet.range("L78").color = (33, 144, 77)
    sheet.range("L78").font.bold = True
    legacy_names = {
        "SST_Stage1_Panel", "SST_Stage1_Status",
        "SST_Stage1_LastSuccessfulOutput", "SST_Stage2_Input",
    }
    new_names = set(workflow.STAGE_ONE_FIELD_NAMES.values()) - legacy_names
    assert len(new_names) == 5
    for name in new_names:
        caller.names[name].delete()
    sheet.range("H73:O83").api.Cut(Destination=sheet.range("Q100:X110").api)
    sheet.range("H73:O83").value = "Legacy old position sentinel"
    if defect == "shape":
        sheet.range("103:103").api.Insert()
    elif defect == "pointer":
        caller.names["SST_Stage1_Status"].refers_to = "='I. Control'!$U$105"
    elif defect == "broken_pointer":
        caller.names["SST_Stage1_LastSuccessfulOutput"].refers_to = "=#REF!"
    elif defect == "outside_pointer":
        caller.names["SST_Stage2_Input"].refers_to = "='I. Control'!$Q$1"
    elif defect == "missing_legacy":
        caller.names["SST_Stage2_Input"].delete()
    caller.save()
    panel = caller.names["SST_Stage1_Panel"].refers_to_range
    before = _stage_panel_contents(panel)
    old_area = _stage_panel_contents(sheet.range("H73:O83"))
    names_before = {name.name: name.refers_to for name in caller.names}
    saves: list[bool] = []
    boundary = _ExcelForwarder(caller, save=lambda: saves.append(True))
    if defect:
        with pytest.raises(BondsLiteError, match=r"(?i)legacy|infer|layout|incomplete|panel"):
            workflow.initialise_stage_panel(boundary)
        assert {name.name: name.refers_to for name in caller.names} == names_before
        assert caller.api.Saved
    else:
        assert workflow.initialise_stage_panel(boundary) is None
        after = {name.name: name.refers_to for name in caller.names}
        assert after.keys() - names_before.keys() == new_names
        assert all(after[name] == reference for name, reference in names_before.items())
        for field, row in {"last_run": 101, "start": 102, "finish": 103,
                           "elapsed": 104, "user": 105}.items():
            region = caller.names[workflow.STAGE_ONE_FIELD_NAMES[field]].refers_to_range
            assert (region.row, region.column) == (row, 21)
        assert workflow.require_stage_panel(caller) is None
        assert workflow._read_stage_fields(caller)["last_output"] == relative
        installed_names = {name.name: name.refers_to for name in caller.names}
        assert workflow.initialise_stage_panel(caller) is None
        assert {name.name: name.refers_to for name in caller.names} == installed_names
    assert saves == []
    assert _stage_panel_contents(panel) == before
    assert _stage_panel_contents(sheet.range("H73:O83")) == old_area


@pytest.mark.excel
@pytest.mark.skipif(
    os.environ.get("BONDS_LITE_RUN_STAGE_PANEL_EXCEL_TESTS") != "1",
    reason="Set BONDS_LITE_RUN_STAGE_PANEL_EXCEL_TESTS=1 for blank-workbook checks",
)
def test_excel_stage_panel_initialise_validate_write_and_reopen(
    tmp_path: Path,
) -> None:
    """Exercise real ranges/styles/names only in new, macro-free local copies."""
    import xlwings as xw
    from bonds_lite_input import (
        STAGE_ONE_FIELDS, _read_stage_fields, initialise_stage_panel,
        require_stage_panel, write_stage_audit,
    )

    original = tmp_path / "blank control.xlsx"
    copy = tmp_path / "panel copy.xlsx"
    expected_names = {
        "SST_Stage1_Panel": "H73:O83",
        "SST_Stage1_LastRun": "L74",
        "SST_Stage1_Start": "L75",
        "SST_Stage1_Finish": "L76",
        "SST_Stage1_Elapsed": "L77",
        "SST_Stage1_User": "L78",
        "SST_Stage1_Status": "L79",
        "SST_Stage1_LastSuccessfulOutput": "L80",
        "SST_Stage2_Input": "L82",
    }
    with xw.App(visible=False, add_book=True) as application:
        application.api.AutomationSecurity = 3
        application.api.EnableEvents = False
        application.api.Calculation = -4135
        application.api.CalculateBeforeSave = False
        application.display_alerts = False
        caller = application.books[0]
        try:
            sheet = caller.sheets[0]
            sheet.name = "I. Control"
            for other in list(caller.sheets)[1:]:
                other.delete()
            sheet.range("D74:F76").merge()
            sheet.range("D74").value = "Existing button sentinel"
            sheet.range("D9").value = "Entity sentinel"
            sheet.range("Q1").formula = "=1+2"
            sheet.range("G73:G83").value = "Left sentinel"
            sheet.range("P73:P83").value = "Right sentinel"
            caller.save(str(original))
            assert not caller.api.ReadOnly and caller.api.Saved

            with pytest.raises(BondsLiteError):
                require_stage_panel(caller)

            for obstacle in ("occupied", "formula", "merge", "protection", "shape", "name"):
                sheet.range("H73:O83").clear()
                if obstacle == "occupied":
                    sheet.range("H73").value = "Do not overwrite"
                elif obstacle == "formula":
                    sheet.range("H73").formula = '=""'
                elif obstacle == "merge":
                    sheet.range("G73:H73").merge()
                elif obstacle == "protection":
                    sheet.api.Protect()
                elif obstacle == "shape":
                    area = sheet.range("H73:O83")
                    shape = sheet.api.Shapes.AddShape(
                        1, area.left, area.top, 30, 20
                    )
                else:
                    caller.names.add("SST_Stage1_Panel", "='I. Control'!$Q$1:$X$11")
                before = sheet.range("G72:P84").formula
                names_before = [(name.name, name.refers_to) for name in caller.names]
                try:
                    with pytest.raises(BondsLiteError):
                        initialise_stage_panel(caller)
                    assert sheet.range("G72:P84").formula == before
                    assert [(name.name, name.refers_to) for name in caller.names] == names_before
                finally:
                    if obstacle == "protection":
                        sheet.api.Unprotect()
                    elif obstacle == "shape":
                        shape.Delete()
                    elif obstacle == "merge":
                        sheet.range("G73:H73").unmerge()
                    elif obstacle == "name":
                        caller.names["SST_Stage1_Panel"].delete()

            sheet.range("H73:O83").clear()
            caller.save()
            assert caller.api.Saved
            assert initialise_stage_panel(caller) is None
            assert not caller.api.Saved, "Initialisation must not auto-save Draft"
            assert require_stage_panel(caller) is None
            assert [item.name for item in caller.sheets] == ["I. Control"]
            assert sheet.range("D74").value == "Existing button sentinel"
            assert sheet.range("G73:G83").value == ["Left sentinel"] * 11
            assert sheet.range("P73:P83").value == ["Right sentinel"] * 11
            for name, address in expected_names.items():
                region = caller.names[name].refers_to_range
                assert region.sheet.name == "I. Control"
                expected = sheet.range(address)
                assert region.row == expected.row and region.column == expected.column
                if name == "SST_Stage1_Panel":
                    assert region.address == expected.address
                else:
                    assert region.last_cell.column <= 15
                    assert region.last_cell.row == expected.row
            for field, anchor in STAGE_ONE_FIELDS.items():
                row = sheet.range(anchor).row
                end_row = row + 1 if field in ("last_output", "stage_two_input") else row
                assert sheet.range((row, 8)).value
                assert sheet.range((row, 8)).api.MergeArea.Address == f"$H${row}:$K${end_row}"
                assert sheet.range(anchor).api.MergeArea.Address == f"$L${row}:$O${end_row}"
            assert sheet.range("H73").value
            assert sheet.range("H73").api.Font.Bold
            assert "[h]" in sheet.range("L77").number_format.lower()
            assert "d" in sheet.range("L74").number_format.lower()
            for address in ("L75", "L76"):
                assert "h" in sheet.range(address).number_format.lower()

            installed = sheet.range("H73:O83").formula
            names_installed = [(name.name, name.refers_to) for name in caller.names]
            assert initialise_stage_panel(caller) is None
            assert sheet.range("H73:O83").formula == installed
            assert [(name.name, name.refers_to) for name in caller.names] == names_installed

            for name in expected_names:
                installed_name = caller.names[name]
                reference = installed_name.refers_to
                installed_name.refers_to = "='I. Control'!$Q$1"
                try:
                    with pytest.raises(BondsLiteError):
                        require_stage_panel(caller)
                finally:
                    installed_name.refers_to = reference
            sheet.api.Protect()
            try:
                with pytest.raises(BondsLiteError):
                    require_stage_panel(caller)
            finally:
                sheet.api.Unprotect()

            output = _complete_run(tmp_path / "Bonds Lite output" / "20261028_090000")
            relative = "Bonds Lite output/20261028_090000"
            fields = {
                "last_run": datetime(2026, 10, 28, 9),
                "start": datetime(2026, 10, 28, 9),
                "finish": datetime(2026, 10, 28, 9, 0, 6),
                "elapsed": 6 / 86400, "user": "panel-user", "status": "Succeeded",
                "last_output": relative, "stage_two_input": relative,
            }
            save_events: list[bool] = []

            def save() -> None:
                save_events.append(bool(application.api.EnableEvents))
                caller.save()

            boundary = _ExcelForwarder(caller, save=save)
            for initial_events in (False, True):
                application.api.EnableEvents = initial_events
                assert write_stage_audit(boundary, fields) is None
                assert caller.api.Saved
                assert bool(application.api.EnableEvents) is initial_events
                assert save_events[-1] is False
            read_back = _read_stage_fields(caller)
            assert set(read_back) == set(STAGE_ONE_FIELDS)
            for key, value in fields.items():
                if key == "elapsed":
                    assert read_back[key] == pytest.approx(value)
                else:
                    assert read_back[key] == value
            for address in ("L80", "L82"):
                assert sheet.range(address).api.Hyperlinks.Count == 1
                assert sheet.range(address).api.Hyperlinks.Item(1).Address == relative
            before = sheet.range("G72:P84").formula
            write_stage_audit(boundary, {"status": "Running", "stage_two_input": ""})
            assert sheet.range("L82").api.Hyperlinks.Count == 0
            assert sheet.range("L80").value == relative
            after = sheet.range("G72:P84").formula
            for row_index, (old_row, new_row) in enumerate(zip(before, after), 72):
                for column_index, (old, new) in enumerate(zip(old_row, new_row), 7):
                    if (row_index, column_index) not in ((79, 12), (82, 12)):
                        assert old == new
            with pytest.raises(BondsLiteError):
                write_stage_audit(boundary, {"D9": "Forbidden write"})
            assert sheet.range("G72:P84").formula == after
            assert sheet.range("D9").value == "Entity sentinel"
            assert sheet.range("Q1").formula == "=1+2"

            def failing_save() -> None:
                assert not application.api.EnableEvents
                raise OSError("Panel save sentinel")

            for initial_events in (False, True):
                application.api.EnableEvents = initial_events
                with pytest.raises(
                    (BondsLiteError, OSError), match="Panel save sentinel"
                ):
                    write_stage_audit(
                        _ExcelForwarder(caller, save=failing_save),
                        {"status": "Failed"},
                    )
                assert bool(application.api.EnableEvents) is initial_events
                unsaved_boundary = _ExcelForwarder(
                    caller, save=save, api=_ExcelForwarder(caller.api, Saved=False)
                )
                with pytest.raises(BondsLiteError):
                    write_stage_audit(unsaved_boundary, {"status": "Failed"})
                assert bool(application.api.EnableEvents) is initial_events
            write_stage_audit(boundary, fields)
            caller.close()
            shutil.copy2(original, copy)
            caller = application.books.open(str(copy), update_links=False)
            assert not caller.api.ReadOnly and caller.api.Saved
            assert require_stage_panel(caller) is None
            assert _read_stage_fields(caller)["stage_two_input"] == relative
            assert [item.name for item in caller.sheets] == ["I. Control"]
            assert not caller.api.HasVBProject
            caller.close()
            caller = application.books.open(str(copy), update_links=False, read_only=True)
            assert caller.api.ReadOnly and caller.api.Saved
            with pytest.raises(BondsLiteError):
                require_stage_panel(caller)
        finally:
            application.api.EnableEvents = False
            for owned in list(application.books):
                owned.close()


@pytest.mark.skipif(
    os.environ.get("BONDS_LITE_RUN_DRAFT_LAYOUT_TESTS") != "1",
    reason="Set BONDS_LITE_RUN_DRAFT_LAYOUT_TESTS=1 for read-only Draft inspection",
)
def test_actual_draft_stage_panel_candidate_is_free_read_only() -> None:
    """Report original layout metadata without saving or opening desktop Excel."""
    import hashlib
    import posixpath
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    openpyxl = pytest.importorskip("openpyxl")
    path = Path(__file__).resolve().parents[1] / "Draft Input spreadsheet for SST.xlsm"
    before = hashlib.sha256(path.read_bytes()).digest()
    workbook = openpyxl.load_workbook(
        path, read_only=True, data_only=False, keep_vba=True
    )
    try:
        sheet = workbook["I. Control"]
        bounds = (8, 73, 15, 83)
        if "SST_Stage1_Panel" in workbook.defined_names:
            destinations = list(workbook.defined_names["SST_Stage1_Panel"].destinations)
            assert len(destinations) == 1 and destinations[0][0] == "I. Control"
            bounds = openpyxl.utils.range_boundaries(destinations[0][1])
        panel_left, panel_top, panel_right, panel_bottom = bounds
        occupied = [
            (cell.coordinate, cell.value)
            for row in sheet.iter_rows(
                min_row=panel_top, max_row=panel_bottom,
                min_col=panel_left, max_col=panel_right,
            )
            for cell in row if cell.value is not None
        ]
        print(f"PANEL_OCCUPIED={occupied!r}")
        with ZipFile(path) as archive:
            spreadsheet = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            relation = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            package = "http://schemas.openxmlformats.org/package/2006/relationships"
            document = ET.fromstring(archive.read("xl/workbook.xml"))
            entry = next(item for item in document.findall(f"{{{spreadsheet}}}sheets/{{{spreadsheet}}}sheet")
                         if item.get("name") == "I. Control")
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(item.get("Target") for item in relationships
                          if item.get("Id") == entry.get(f"{{{relation}}}id"))
            sheet_path = (target.lstrip("/") if target.startswith("/") else
                          posixpath.normpath(posixpath.join("xl", target)))
            xml = ET.fromstring(archive.read(sheet_path))
            merges = [item.get("ref") for item in xml.findall(
                f"{{{spreadsheet}}}mergeCells/{{{spreadsheet}}}mergeCell"
            )]
            protection = xml.find(f"{{{spreadsheet}}}sheetProtection")
            print(f"PANEL_MERGES={merges!r}")
            print(f"PANEL_PROTECTION={None if protection is None else protection.attrib!r}")
            drawing_ids = [item.get(f"{{{relation}}}id") for item in xml.findall(
                f"{{{spreadsheet}}}drawing"
            )]
            drawing_overlaps = []
            rel_path = posixpath.join(posixpath.dirname(sheet_path), "_rels",
                                     posixpath.basename(sheet_path) + ".rels")
            if drawing_ids and rel_path in archive.namelist():
                links = ET.fromstring(archive.read(rel_path))
                drawing_namespace = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
                for link in links.findall(f"{{{package}}}Relationship"):
                    if link.get("Id") not in drawing_ids:
                        continue
                    drawing_target = link.get("Target")
                    drawing_path = (drawing_target.lstrip("/") if drawing_target.startswith("/") else
                                    posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), drawing_target)))
                    drawing = ET.fromstring(archive.read(drawing_path))
                    for anchor in drawing:
                        start = anchor.find("xdr:from", drawing_namespace)
                        finish = anchor.find("xdr:to", drawing_namespace)
                        print(f"PANEL_DRAWING_ANCHOR={ET.tostring(anchor, encoding='unicode')}")
                        if start is not None:
                            row = int(start.find("xdr:row", drawing_namespace).text) + 1
                            column = int(start.find("xdr:col", drawing_namespace).text) + 1
                            last_row = row if finish is None else int(finish.find("xdr:row", drawing_namespace).text) + 1
                            last_column = column if finish is None else int(finish.find("xdr:col", drawing_namespace).text) + 1
                            if (row <= panel_bottom and last_row >= panel_top
                                    and column <= panel_right and last_column >= panel_left):
                                drawing_overlaps.append((row, column, last_row, last_column))
            overlapping_merges = []
            for address in merges:
                left, top, right, bottom = openpyxl.utils.range_boundaries(address)
                if (top <= panel_bottom and bottom >= panel_top
                    and left <= panel_right and right >= panel_left):
                    overlapping_merges.append(address)
            from bonds_lite_input import STAGE_ONE_NAMES

            if "SST_Stage1_Panel" in workbook.defined_names:
                field_regions = []
                field_anchors = {}
                for name in STAGE_ONE_NAMES:
                    definition = workbook.defined_names[name]
                    destinations = list(definition.destinations)
                    assert len(destinations) == 1 and destinations[0][0] == "I. Control"
                    if name == "SST_Stage1_Panel":
                        continue
                    named_bounds = openpyxl.utils.range_boundaries(destinations[0][1])
                    left, top, right, bottom = named_bounds
                    intersecting = [
                        merge_bounds
                        for merge_bounds in map(openpyxl.utils.range_boundaries, merges)
                        if (merge_bounds[0] <= right and merge_bounds[2] >= left
                            and merge_bounds[1] <= bottom and merge_bounds[3] >= top)
                    ]
                    assert len(intersecting) <= 1
                    if intersecting:
                        merged_bounds = intersecting[0]
                        assert (left == right and top == bottom) or named_bounds == merged_bounds
                        left, top, right, bottom = merged_bounds
                    else:
                        assert left == right and top == bottom
                    assert panel_left <= left <= right <= panel_right
                    assert panel_top <= top <= bottom <= panel_bottom
                    for old_left, old_top, old_right, old_bottom in field_regions:
                        assert (right < old_left or left > old_right
                                or bottom < old_top or top > old_bottom)
                    field_regions.append((left, top, right, bottom))
                    field_anchors[name] = sheet.cell(row=top, column=left)
                assert sheet.cell(row=panel_top, column=panel_left).value == "Stage 1 - Generate Bonds Lite Input"
                assert field_anchors["SST_Stage1_Status"].value in (
                    "Not Run", "Running", "Succeeded", "Failed", "Interrupted",
                )
                if bounds == (8, 73, 15, 83):
                    assert sheet["H74"].fill.fgColor.rgb[-6:] == "FF7F7F"
                    assert sheet["L74"].fill.fgColor.rgb[-6:] == "FFE6E6"
                assert field_anchors["SST_Stage1_Elapsed"].number_format == "[h]:mm:ss"
            else:
                assert not occupied, f"Occupied candidate cells: {occupied!r}"
                assert not overlapping_merges, f"Candidate merges: {overlapping_merges!r}"
            assert protection is None or protection.get("sheet", "0") in ("0", "false")
            assert not drawing_overlaps, f"Candidate drawing overlaps: {drawing_overlaps!r}"
    finally:
        workbook.close()
        assert hashlib.sha256(path.read_bytes()).digest() == before


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_supplied_bond_template_layout() -> None:
    """Rescan supplied table metadata and formula seeds without writing the files."""
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    import openpyxl

    root = Path(__file__).resolve().parents[1]
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    layouts = []
    for filename in ("Bonds Lite data template.xlsx", "Bonds Lite data template_MAP.xlsx",
                     "Bonds Lite input template.xlsx"):
        path = root / filename
        with ZipFile(path) as archive:
            tables = [
                ET.fromstring(archive.read(name))
                for name in archive.namelist()
                if name.startswith("xl/tables/") and name.endswith(".xml")
            ]
        table = next(item for item in tables if item.get("name") == "I_tbl_Monthly_Bonds")
        headers = [
            column.get("name")
            for column in table.findall("s:tableColumns/s:tableColumn", namespace)
        ]
        print(f"FILE={filename} TABLE={table.get('ref')} COLUMNS={len(headers)}")
        print(f"FIRST_HEADERS={headers[:16]}")
        print(f"LAST_HEADERS={headers[-23:]}")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook["I. Bonds Data"]
            seeds = next(sheet.iter_rows(min_row=9, max_row=9))
            print("FORMULA_SEEDS=" + repr([
                (cell.coordinate, str(getattr(cell.value, "text", cell.value))[:180])
                for cell in seeds if cell.data_type == "f"
            ]))
            if filename == "Bonds Lite input template.xlsx":
                named = workbook.defined_names["output_folder_path"]
                assert list(named.destinations) == [("I. Control", "$D$19")]
        finally:
            workbook.close()
        layouts.append((table.get("ref"), headers))
    assert layouts[0][0].split(":")[1].startswith("ATT")
    assert len(layouts[0][1]) == RAW_COLUMN_COUNT
    assert layouts[1][0].split(":")[1].startswith("ATT")
    assert len(layouts[1][1]) == RAW_COLUMN_COUNT
    assert len(layouts[2][1]) == 1235
    assert layouts[0][1] == layouts[1][1] == layouts[2][1][:RAW_COLUMN_COUNT]


@pytest.mark.excel
@pytest.mark.parametrize(
    "defect",
    [
        None,
        "merged",
        "missing",
        "broken",
        "multiple_cells",
        "multiple_areas",
        "other_sheet",
        "foreign_book",
        "missing_folder",
        "file_not_folder",
    ],
)
def test_set_model_output_folder_validates_local_named_cell_before_write(
    stage_panel_excel_book: Any,
    tmp_path: Path,
    defect: str | None,
) -> None:
    """Use only a blank owned book to verify Excel name and merge semantics."""
    from bonds_lite_input import set_model_output_folder

    book = stage_panel_excel_book
    sheet = book.sheets["I. Control"]
    sheet.range("D19").value = "Original output setting"
    sheet.range("C18:E20").color = (230, 240, 230)
    other = book.sheets.add("Other")
    other.range("D19").value = "Other sheet sentinel"
    folder = tmp_path / "20261028_090000" / "JRL MAP"
    folder.mkdir(parents=True)
    foreign = None
    reference = "='I. Control'!$D$19"
    if defect == "merged":
        sheet.range("D19:F19").merge()
        reference = "='I. Control'!$F$19"
    elif defect == "broken":
        reference = "=#REF!"
    elif defect == "multiple_cells":
        reference = "='I. Control'!$D$19:$D$20"
    elif defect == "multiple_areas":
        reference = "=('I. Control'!$D$19,'I. Control'!$F$19)"
    elif defect == "other_sheet":
        reference = "=Other!$D$19"
    elif defect == "foreign_book":
        foreign = book.app.books.add()
        foreign.sheets[0].name = "I. Control"
        foreign.sheets[0].range("D19").value = "Foreign sentinel"
        foreign.save(str(tmp_path / "Foreign.xlsx"))
        reference = f"='[{foreign.name}]I. Control'!$D$19"
    elif defect == "missing_folder":
        folder = folder / "missing"
    elif defect == "file_not_folder":
        folder = folder / "not a folder.xlsx"
        folder.write_bytes(b"Non-directory sentinel")
    if defect != "missing":
        book.names.add("output_folder_path", reference)
    before = sheet.range("C18:G20").formula
    try:
        if defect in (None, "merged"):
            assert set_model_output_folder(book, folder) is None
            assert sheet.range("D19").value == str(folder.resolve())
        else:
            with pytest.raises(BondsLiteError):
                set_model_output_folder(book, folder)
            assert sheet.range("C18:G20").formula == before
        assert other.range("D19").value == "Other sheet sentinel"
        if foreign is not None:
            assert foreign.sheets[0].range("D19").value == "Foreign sentinel"
    finally:
        if foreign is not None:
            foreign.close()


@pytest.mark.excel
@pytest.mark.parametrize("named_value", ["correct", "wrong", "blank", "missing"])
def test_verify_saved_copy_checks_optional_model_output_folder_after_save(
    stage_panel_excel_book: Any,
    local_paths: tuple[Path, Path, Path, Path],
    stress_values: dict,
    tmp_path: Path,
    named_value: str,
) -> None:
    """Check the persisted name value alongside snapshot data in a blank workbook."""
    from bonds_lite_input import StageOneInputs, verify_saved_copy

    book = stage_panel_excel_book
    settings = StageOneInputs(*local_paths, datetime(2026, 9, 30), "JRL", stress_values)
    sheet = book.sheets["I. Control"]
    sheet.range("D10").value = settings.valuation_date
    sheet.range("D13").value = settings.entity
    folder = tmp_path / "JRL Bonds"
    folder.mkdir()
    if named_value != "missing":
        book.names.add("output_folder_path", "='I. Control'!$D$19")
        sheet.range("D19").value = {
            "correct": str(folder.resolve()),
            "wrong": str(tmp_path / "Other user"),
            "blank": None,
        }[named_value]
    bonds = book.sheets.add("I. Bonds Data")
    bonds.range("B8:C9").value = [["ISIN", "Calculated"], ["BONDS_FIXTURE_1", None]]
    bonds.range("C9").formula = "=1+1"
    bonds.api.ListObjects.Add(1, bonds.range("B8:C9").api, None, 1).Name = (
        "I_tbl_Monthly_Bonds"
    )
    sensitivity = book.sheets.add("I. Sensitivity")
    for address, grid in settings.stresses.items():
        sensitivity.range(address).value = grid
    book.save()
    assert book.api.Saved
    if named_value == "correct":
        assert (
            verify_saved_copy(
                book, settings, 1, [3], settings.stresses, model_output_folder=folder
            )
            is None
        )
    else:
        with pytest.raises(BondsLiteError):
            verify_saved_copy(
                book, settings, 1, [3], settings.stresses, model_output_folder=folder
            )


@pytest.fixture
def native_excel_case(tmp_path: Path):
    """Use small blank workbooks for range copying, not the large bond template."""
    import xlwings as xw

    with xw.App(visible=False, add_book=True) as application:
        working = source = book = None
        resources: list[Any] = []
        try:
            application.api.AutomationSecurity = 3
            application.api.EnableEvents = False
            application.api.Calculation = -4135
            application.display_alerts = False
            working = application.books[0]
            source = application.books.add()
            resources.extend((application, working, source, None, None, tmp_path))
            yield resources
        finally:
            try:
                with ExitStack() as cleanup:
                    for book in list(application.books):
                        cleanup.callback(book.close)
            finally:
                resources.clear()
                book = source = working = None


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_native_copy_values_interface_exists() -> None:
    """Establish the missing helper without creating an Excel application."""
    from bonds_lite_input import native_copy_values

    assert callable(native_copy_values)


class _ExcelForwarder:
    """Forward to live Excel while observing selected transfer boundaries."""

    def __init__(self, wrapped: Any, **overrides: Any) -> None:
        self.wrapped = wrapped
        self.overrides = overrides

    def __getattr__(self, name: str) -> Any:
        if name in self.overrides:
            return self.overrides[name]
        if name.casefold() in {"value", "value2", "options"}:
            pytest.fail("Native copying must not read or write bulk Python values")
        return getattr(self.wrapped, name)

    def __setattr__(self, name: str, value: Any) -> None:
        if name.casefold() in {"value", "value2", "options"}:
            pytest.fail("Native copying must not read or write bulk Python values")
        object.__setattr__(self, name, value)


def _excel_copy_snapshot(region: Any) -> tuple[Any, ...]:
    """Retain source values, formulas and cell types for a small regression grid."""
    return (
        region.api.Value2,
        region.api.Formula2,
        tuple(
            region.sheet.api.Evaluate(f"TYPE({cell.address})")
            for cell in region
        ),
    )


def _raw_column_blocks(calculated: dict[int, str]) -> list[tuple[int, int]]:
    """Group the fixed bond schema's raw columns without any Excel calls."""
    raw = [
        column for column in range(RAW_FIRST_COLUMN, RAW_LAST_COLUMN + 1)
        if column not in calculated
    ]
    blocks = []
    for _, group in groupby(enumerate(raw), key=lambda item: item[1] - item[0]):
        columns = [column for _, column in group]
        blocks.append((columns[0], columns[-1]))
    return blocks


def test_raw_column_blocks_exclude_source_formulas_and_destination_tail() -> None:
    """Keep the old-loop baseline and native route bounded to B:I and L:ATT."""
    calculated = {10: "=source J", 11: "=source K"}
    calculated.update({column: "=tail" for column in range(1217, 1237)})
    assert _raw_column_blocks(calculated) == [(2, 9), (12, RAW_LAST_COLUMN)]


def _seed_source_formulas(sheet: Any, rows: int, folder: Path) -> None:
    """Give disposable sources unrefreshed formulas that must never be copied."""
    sheet.range(f"J9:J{8 + rows}").api.Formula2 = (
        f"='{folder}\\[Source-only-must-not-refresh.xlsx]Sheet1'!$A$1"
    )
    sheet.range(f"K9:K{8 + rows}").api.Formula2 = '="SOURCE_K_MUST_NOT_COPY"'


def _assert_excel_copy_equal(source: Any, destination: Any) -> None:
    """Compare values case-sensitively and distinguish blanks, numbers and text."""
    local = source.get_address()
    external = destination.get_address(external=True)
    for comparison in (
        f"--NOT(EXACT({local},{external}))",
        f"--(ISTEXT({local})<>ISTEXT({external}))",
        f"--(ISNUMBER({local})<>ISNUMBER({external}))",
        f"--(ISBLANK({local})<>ISBLANK({external}))",
    ):
        result = source.sheet.api.Evaluate(f"SUMPRODUCT({comparison})")
        assert isinstance(result, (int, float)) and result == 0, comparison


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_native_copy_values_preserves_values_types_formats_and_source(
    native_excel_case: list[Any],
) -> None:
    """Paste values only, preserving literal text, empty strings and formatting."""
    from bonds_lite_input import native_copy_values

    application, working, source, _, _, folder = native_excel_case
    working.save(str(folder / "working.xlsx"))
    source_sheet = source.sheets.add("Native source")
    target_sheet = working.sheets.add("Native target")
    source_range = source_sheet.range("A1:B7")
    target_range = target_sheet.range("D3:E9")
    source_range.api.Value2 = (
        (12.5, True), (0, False), (None, None), (None, None),
        (None, None), ("MiXeD", "mixed"), (None, "#N/A"),
    )
    source_sheet.range("A3").clear_contents()
    source_sheet.range("B3").api.Formula2 = '=IF(TRUE,"","unused")'
    for address, text in (
        ("A4", "=1+1"), ("B4", "'literal"), ("A5", "00123"),
        ("B5", "123"), ("B7", "#N/A"),
    ):
        source_sheet.range(address).number_format = "@"
        source_sheet.range(address).api.Value2 = f"'{text}"
        assert source_sheet.range(address).api.Value2 == text
        assert source_sheet.range(address).api.HasFormula is False
    source_sheet.range("A7").api.Formula2 = "=6*7"
    source_sheet.api.Calculate()
    assert source_sheet.range("B3").api.Value2 == ""
    assert source_sheet.api.Evaluate("ISTEXT(B3)") is True
    assert source_sheet.api.Evaluate("ISERROR(B7)") is False
    target_range.api.Value2 = "Overwrite even with a blank"
    target_range.api.NumberFormat = "0.0000"
    target_range.api.Font.Bold = True
    target_range.api.Interior.Color = 65535
    target_range.api.HorizontalAlignment = -4152
    formats = tuple(
        (cell.api.NumberFormat, cell.api.Font.Bold, cell.api.Interior.Color,
         cell.api.HorizontalAlignment)
        for cell in target_range
    )
    source.save(str(folder / "source.xlsx"))
    source_path = Path(source.fullname)
    original_bytes = source_path.read_bytes()
    original_values = _excel_copy_snapshot(source_range)
    names = {
        book.name: tuple(name.name for name in book.names)
        for book in (source, working)
    }
    links = {book.name: book.api.LinkSources(1) for book in (source, working)}

    assert native_copy_values(source_range, target_range) is None

    _assert_excel_copy_equal(source_range, target_range)
    assert target_range.api.HasFormula is False
    assert _excel_copy_snapshot(source_range) == original_values
    assert source_path.read_bytes() == original_bytes
    assert tuple(
        (cell.api.NumberFormat, cell.api.Font.Bold, cell.api.Interior.Color,
         cell.api.HorizontalAlignment)
        for cell in target_range
    ) == formats
    assert {
        book.name: tuple(name.name for name in book.names)
        for book in (source, working)
    } == names
    assert {
        book.name: book.api.LinkSources(1) for book in (source, working)
    } == links
    assert application.api.CutCopyMode == 0


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.parametrize("error_kind", ["constant", "formula"])
def test_native_copy_values_rejects_excel_errors_and_clears_copy_mode(
    native_excel_case: list[Any], error_kind: str,
) -> None:
    """Reject real Excel errors before changing the destination, not error text."""
    from bonds_lite_input import native_copy_values

    application, working, source, _, _, _ = native_excel_case
    sheet = source.sheets[0]
    source_range = sheet.range("A1:B2")
    source_range.api.Value2 = ((1, 2), (3, 4))
    sheet.range("B2").api.Formula2 = "=NA()"
    sheet.api.Calculate()
    if error_kind == "constant":
        sheet.range("B2").api.Copy()
        sheet.range("B2").api.PasteSpecial(
            Paste=-4163, SkipBlanks=False, Transpose=False
        )
        assert sheet.range("B2").api.HasFormula is False
    assert sheet.api.Evaluate("ISNA(B2)") is True
    destination = working.sheets.add("Reject error").range("A1:B2")
    destination.api.Value2 = (("keep", "these"), ("original", "values"))
    before = destination.api.Value2
    source_range.api.Copy()
    assert application.api.CutCopyMode != 0

    with pytest.raises(BondsLiteError, match=r"(?i)error|#N/A"):
        native_copy_values(source_range, destination)

    assert destination.api.Value2 == before
    assert sheet.api.Evaluate("ISNA(B2)") is True
    assert application.api.CutCopyMode == 0


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_native_copy_values_clears_copy_mode_when_paste_fails(
    native_excel_case: list[Any],
) -> None:
    """Clean up after a real protected-sheet PasteSpecial failure."""
    from bonds_lite_input import native_copy_values
    from pywintypes import com_error

    application, working, source, _, _, _ = native_excel_case
    source_range = source.sheets[0].range("A1:B1")
    source_range.api.Value2 = ((12, "value"),)
    target_sheet = working.sheets.add("Protected destination")
    destination = target_sheet.range("A1:B1")
    destination.api.Value2 = (("unchanged", "unchanged"),)
    before = destination.api.Value2
    target_sheet.api.Protect()
    try:
        with pytest.raises((BondsLiteError, com_error)):
            native_copy_values(source_range, destination)
        assert application.api.CutCopyMode == 0
        assert destination.api.Value2 == before
    finally:
        target_sheet.api.Unprotect()


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.parametrize(
    ("original", "changed"),
    [("MiXeD", "mixed"), ("123", 123), (None, 0), (None, "")],
    ids=["case", "numeric-text", "blank-to-zero", "blank-to-empty-text"],
)
def test_native_copy_values_rejects_postcopy_value_or_type_mismatch(
    native_excel_case: list[Any], original: object, changed: object,
) -> None:
    """Inject corruption after a real paste to exercise Excel-side validation."""
    from bonds_lite_input import native_copy_values

    application, working, source, _, _, _ = native_excel_case
    source_range = source.sheets[0].range("A1")
    if isinstance(original, str):
        source_range.number_format = "@"
    source_range.api.Value2 = original
    destination = working.sheets.add("Mismatch target").range("A1")

    def paste_then_corrupt(*args: Any, **kwargs: Any) -> None:
        destination.api.PasteSpecial(*args, **kwargs)
        if changed == "":
            destination.api.Formula2 = '=""'
            destination.sheet.api.Calculate()
            destination.api.Copy()
            destination.api.PasteSpecial(Paste=-4163)
        else:
            destination.api.Value2 = changed

    boundary = _ExcelForwarder(
        destination,
        api=_ExcelForwarder(destination.api, PasteSpecial=paste_then_corrupt),
    )
    with pytest.raises(BondsLiteError, match=r"(?i)match|differ|equal|verif|read.back"):
        native_copy_values(source_range, boundary)
    assert application.api.CutCopyMode == 0


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_native_copy_values_uses_native_paste_and_scalar_excel_validation(
    native_excel_case: list[Any],
) -> None:
    """Observe real COM calls without allowing a Python bulk-value transfer."""
    from bonds_lite_input import native_copy_values

    _, working, source, _, _, folder = native_excel_case
    working.save(str(folder / "working.xlsx"))
    source_range = source.sheets[0].range("A1:B2")
    source_range.api.Value2 = ((1, "MiXeD"), (None, True))
    destination = working.sheets.add("Native boundary").range("A1:B2")
    calls: list[str] = []
    evaluations: list[str] = []

    def copy(*args: Any, **kwargs: Any) -> Any:
        calls.append("copy")
        return source_range.api.Copy(*args, **kwargs)

    def paste(*args: Any, **kwargs: Any) -> Any:
        calls.append("paste")
        assert not args
        assert kwargs == {"Paste": -4163, "SkipBlanks": False, "Transpose": False}
        return destination.api.PasteSpecial(**kwargs)

    def evaluate(expression: str) -> Any:
        evaluations.append(expression)
        result = source_range.sheet.api.Evaluate(expression)
        assert isinstance(result, (int, float)), "Return only scalar error counts"
        return result

    source_boundary = _ExcelForwarder(
        source_range,
        api=_ExcelForwarder(source_range.api, Copy=copy),
        sheet=_ExcelForwarder(
            source_range.sheet,
            api=_ExcelForwarder(source_range.sheet.api, Evaluate=evaluate),
        ),
    )
    target_boundary = _ExcelForwarder(
        destination, api=_ExcelForwarder(destination.api, PasteSpecial=paste),
    )
    native_copy_values(source_boundary, target_boundary)

    assert calls == ["copy", "paste"]
    assert evaluations, "Validate source errors and postcopy equality inside Excel"
    for function in ("ISERROR", "EXACT", "ISTEXT", "ISNUMBER", "ISBLANK"):
        assert any(function in expression.upper() for expression in evaluations)
    _assert_excel_copy_equal(source_range, destination)


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_populate_bond_data_routes_only_raw_columns_through_native_copy(
    excel_case: list[Any], monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Copy every raw cell once while retaining metadata and formula filling."""
    import bonds_lite_input as workflow

    native_copy_values = workflow.native_copy_values
    _, working, source, headers, sample, folder = excel_case
    sheet = source.sheets[0]
    target = working.sheets["I. Bonds Data"]
    target_headers = target.range("B8:AUN8").options(ndim=2).value
    sheet.range("B8").value = headers
    grid = [sample[0].copy() for _ in range(5)]
    for index, row in enumerate(grid, start=1):
        row[0], row[1] = index, f"NATIVE_{index}"
    sheet.range("B9").value = grid
    sheet.api.ListObjects.Add(
        1, sheet.range("B8:ATT13").api, None, 1
    ).Name = "I_tbl_Monthly_Bonds"
    calculated = workflow.formula_columns(
        target.api.ListObjects("I_tbl_Monthly_Bonds")
    )
    assert set(calculated) == {10, 11, *range(1217, 1237)}
    raw_columns = set(range(RAW_FIRST_COLUMN, RAW_LAST_COLUMN + 1)) - calculated.keys()
    formula_raw_column = min(raw_columns - {2, 3})
    raw_region = sheet.range((9, formula_raw_column), (13, formula_raw_column))
    raw_region.api.Formula2 = "=40+2"
    raw_region.api.Calculate()
    _seed_source_formulas(sheet, 5, folder)
    source_formulas = sheet.range("J9:K13").api.Formula2
    source_values = sheet.range("J9:K13").api.Value2
    coverage: set[tuple[int, int]] = set()

    def observe_copy(source_range: Any, destination_range: Any) -> None:
        assert source_range.column >= RAW_FIRST_COLUMN
        assert source_range.column + source_range.shape[1] - 1 <= RAW_LAST_COLUMN
        assert (
            source_range.sheet.book.app.pid == destination_range.sheet.book.app.pid
        )
        assert source_range.shape == destination_range.shape
        first_row, first_column = destination_range.row, destination_range.column
        row_count, column_count = destination_range.shape
        locations = {
            (row, column)
            for row in range(first_row, first_row + row_count)
            for column in range(first_column, first_column + column_count)
        }
        assert all(column in raw_columns for _, column in locations)
        assert coverage.isdisjoint(locations)
        coverage.update(locations)
        native_copy_values(source_range, destination_range)

    monkeypatch.setattr(workflow, "native_copy_values", observe_copy)
    assert workflow.populate_bond_data(source, working) == (5, list(calculated))
    assert coverage == {
        (row, column) for row in range(9, 14) for column in raw_columns
    }
    assert target.range("B8:AUN8").options(ndim=2).value == target_headers
    assert target_headers[0][:RAW_COLUMN_COUNT] == headers[0]
    assert sheet.api.ListObjects("I_tbl_Monthly_Bonds").ListColumns.Count == RAW_COLUMN_COUNT
    assert sheet.range("J9:K13").api.Formula2 == source_formulas
    assert sheet.range("J9:K13").api.Value2 == source_values
    for first, last in _raw_column_blocks(calculated):
        _assert_excel_copy_equal(
            sheet.range((9, first), (13, last)),
            target.range((9, first), (13, last)),
        )
    assert raw_region.api.HasFormula is True
    pasted_raw = target.range((9, formula_raw_column), (13, formula_raw_column))
    assert pasted_raw.api.HasFormula is False
    for column in calculated:
        assert target.range((9, column), (13, column)).api.HasFormula is True


@pytest.mark.excel
@pytest.mark.performance
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.skipif(
    not PERFORMANCE_ENABLED, reason="Set BONDS_LITE_RUN_PERFORMANCE_TESTS=1"
)
def test_populate_bond_data_copy_benchmark(
    excel_case: list[Any], record_property: Any,
) -> None:
    """Compare old and native raw transfers, then time population without refresh."""
    from bonds_lite_input import formula_columns, native_copy_values, populate_bond_data

    _, working, source, headers, sample, folder = excel_case
    rows, columns = 998, RAW_COLUMN_COUNT
    sheet = source.sheets[0]
    target = working.sheets["I. Bonds Data"]
    calculated = formula_columns(target.api.ListObjects("I_tbl_Monthly_Bonds"))
    assert len(headers[0]) == columns
    sheet.range("B8").value = headers
    grid = [sample[0].copy() for _ in range(rows)]
    for index, row in enumerate(grid, start=1):
        row[0] = index
        row[1] = f"BENCH_{index:04d}"
    sheet.range((9, 2), (8 + rows, 1 + columns)).value = grid
    sheet.api.ListObjects.Add(
        1, sheet.range((8, 2), (8 + rows, 1 + columns)).api, None, 1
    ).Name = "I_tbl_Monthly_Bonds"
    _seed_source_formulas(sheet, rows, folder)
    blocks = _raw_column_blocks(calculated)
    assert blocks == [(2, 9), (12, RAW_LAST_COLUMN)]
    transfers = [
        (
            sheet.range((9, first), (8 + rows, last)),
            target.range((9, first), (8 + rows, last)),
        )
        for first, last in blocks
    ]
    excel_errors = {
        "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
        "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!",
        "#UNKNOWN!", "#CONNECT!",
    }
    for _, destination in transfers:
        destination.clear_contents()

    started = time.perf_counter()
    for first, last in blocks:
        for offset in range(0, rows, 100):
            stop = min(rows, offset + 100)
            origin = sheet.range((9 + offset, first), (8 + stop, last))
            destination = target.range((9 + offset, first), (8 + stop, last))
            values = origin.options(ndim=2, err_to_str=True).value
            assert not any(
                isinstance(value, str) and value in excel_errors
                for row in values for value in row
            ), "Excel error in old-loop raw inputs"
            destination.value = [
                [f"'{value}" if isinstance(value, str) else value for value in row]
                for row in values
            ]
            assert destination.options(ndim=2, err_to_str=True).value == values
    baseline_elapsed = time.perf_counter() - started

    for _, destination in transfers:
        destination.clear_contents()
    started = time.perf_counter()
    for origin, destination in transfers:
        native_copy_values(origin, destination)
    native_elapsed = time.perf_counter() - started
    for origin, destination in transfers:
        _assert_excel_copy_equal(origin, destination)

    started = time.perf_counter()
    result = populate_bond_data(source, working)
    elapsed = time.perf_counter() - started

    print(
        f"OLD_RAW_SECONDS={baseline_elapsed:.6f} "
        f"NATIVE_RAW_SECONDS={native_elapsed:.6f} "
        f"COPY_SECONDS={elapsed:.6f} ROWS={rows} COLUMNS={columns}"
    )
    record_property("OLD_RAW_SECONDS", baseline_elapsed)
    record_property("NATIVE_RAW_SECONDS", native_elapsed)
    record_property("COPY_SECONDS", elapsed)
    record_property("ROWS", rows)
    record_property("COLUMNS", columns)
    assert result == (rows, list(calculated))
    assert target.api.ListObjects("I_tbl_Monthly_Bonds").ListRows.Count == rows
    for offset in range(0, rows, 100):
        last = min(rows, offset + 100)
        expected = sheet.range((9 + offset, 2), (8 + last, 1 + columns)).options(
            ndim=2, err_to_str=True
        ).value
        actual = target.range((9 + offset, 2), (8 + last, 1 + columns)).options(
            ndim=2, err_to_str=True
        ).value
        for column in range(2, 2 + columns):
            if column not in calculated:
                assert [row[column - 2] for row in actual] == [
                    row[column - 2] for row in expected
                ], f"Raw-column mismatch at {column_name(column)}, row {9 + offset}"


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.parametrize("rows", [1, 5, 998, 1001, 1058])
def test_excel_populate_bond_data_dynamic_rows(excel_case, rows: int) -> None:
    """Resize the real template and retain all calculated columns without refresh."""
    from bonds_lite_input import populate_bond_data

    _, working, source, headers, sample, folder = excel_case
    sheet = source.sheets[0]
    target = working.sheets["I. Bonds Data"]
    target_headers = target.range("B8:AUN8").options(ndim=2).value
    local_formulas = target.range("J9:K9").api.Formula2
    sheet.range("B8").value = headers
    for start in range(0, rows, 100):
        grid = [sample[0].copy() for _ in range(min(100, rows - start))]
        for offset, row in enumerate(grid):
            row[0] = start + offset + 1
            row[1] = f"TEST_{start + offset + 1}"
        sheet.range((9 + start, 2)).value = grid
    sheet.api.ListObjects.Add(
        1, sheet.range((8, RAW_FIRST_COLUMN), (8 + rows, RAW_LAST_COLUMN)).api,
        None, 1,
    ).Name = "I_tbl_Monthly_Bonds"
    _seed_source_formulas(sheet, rows, folder)
    source_formulas = sheet.range(f"J9:K{8 + rows}").api.Formula2
    source_values = sheet.range(f"J9:K{8 + rows}").api.Value2
    count, calculated_columns = populate_bond_data(source, working)
    target = working.sheets["I. Bonds Data"]
    table = target.api.ListObjects("I_tbl_Monthly_Bonds")
    assert count == rows == table.ListRows.Count
    assert table.ListColumns.Count == 1235
    assert set(calculated_columns) == {10, 11, *range(1217, 1237)}
    assert sheet.api.ListObjects("I_tbl_Monthly_Bonds").ListColumns.Count == RAW_COLUMN_COUNT
    assert sheet.range(f"J9:K{8 + rows}").api.Formula2 == source_formulas
    assert sheet.range(f"J9:K{8 + rows}").api.Value2 == source_values
    assert target.range("J9:K9").api.Formula2 == local_formulas
    assert target.range(f"C{8 + rows}").value == f"TEST_{rows}"
    for column in calculated_columns:
        assert target.range((9, column)).api.HasFormula
        assert target.range((8 + rows, column)).api.HasFormula
    if rows < 998:
        assert target.range((9 + rows, 2)).value is None
    assert target.range("B8:AUN8").options(ndim=2).value == target_headers
    assert target_headers[0][:RAW_COLUMN_COUNT] == headers[0]


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.parametrize("raw_text", ["=1+1", "'TEST_LITERAL"])
def test_excel_populate_bond_data_preserves_raw_text(
    excel_case, raw_text: str,
) -> None:
    """Preserve literal raw text without interpreting it as an Excel formula."""
    from bonds_lite_input import populate_bond_data

    _, working, source, headers, sample, folder = excel_case
    sheet = source.sheets[0]
    sheet.range("B8").value = headers
    sheet.range("B9").value = sample
    sheet.api.ListObjects.Add(
        1, sheet.range("B8:ATT9").api, None, 1
    ).Name = "I_tbl_Monthly_Bonds"
    _seed_source_formulas(sheet, 1, folder)
    source_cell = sheet.range("C9").api
    source_cell.NumberFormat = "@"
    source_cell.Value2 = f"'{raw_text}"
    source_value = source_cell.Value2
    assert source_value == raw_text, (
        f"Text setup: input={raw_text!r}, source.Value2={source_value!r}"
    )
    assert source_cell.HasFormula is False
    destination = working.sheets["I. Bonds Data"].range("C9").api

    try:
        populate_bond_data(source, working)
    finally:
        print(
            f"input={raw_text!r}, source.Value2={source_cell.Value2!r}, "
            f"destination.Value2={destination.Value2!r}, "
            f"destination.HasFormula={destination.HasFormula!r}"
        )

    assert source_cell.Value2 == source_value
    assert source_cell.HasFormula is False
    assert destination.Value2 == source_value
    assert destination.HasFormula is False


@pytest.fixture
def excel_workflow_case(excel_case: list[Any], stress_values: dict) -> tuple:
    """Prepare distinct one-row Bonds and two-row MAP sources, without calculation."""
    application, working, source, headers, sample, folder = excel_case
    isin_index = headers[0].index("ISIN")
    pv_index = next(
        index
        for index, header in enumerate(headers[0])
        if re.search(r"(?i)^pv(?:$|[_ (])|present.?value", header)
    )
    bonds_row = list(sample[0])
    bonds_row[isin_index] = "BONDS_FIXTURE_1"
    bonds_row[pv_index] = 101.25
    source.sheets[0].range("B8").value = headers
    source.sheets[0].range("B9").value = [bonds_row]
    source.sheets[0].api.ListObjects.Add(
        1, source.sheets[0].range("B8:ATT9").api, None, 1
    ).Name = "I_tbl_Monthly_Bonds"
    _seed_source_formulas(source.sheets[0], 1, folder)
    data = folder / "Data.xlsx"
    source.save(str(data))
    source.close()
    map_data = folder / "MAP Data.xlsx"
    map_source = application.books.add()
    try:
        map_source.sheets[0].name = "I. Bonds Data"
        map_rows = [list(sample[0]), list(sample[0])]
        for index, row in enumerate(map_rows, 1):
            row[isin_index] = f"MAP_FIXTURE_{index}"
            row[pv_index] = (202.5, 303.75)[index - 1]
        map_source.sheets[0].range("B8").value = headers
        map_source.sheets[0].range("B9").value = map_rows
        map_source.sheets[0].api.ListObjects.Add(
            1, map_source.sheets[0].range("B8:ATT10").api, None, 1
        ).Name = "I_tbl_Monthly_Bonds"
        _seed_source_formulas(map_source.sheets[0], 2, folder)
        map_source.save(str(map_data))
    finally:
        map_source.close()
    template = Path(working.fullname)
    labels = working.sheets["I. Sensitivity"].range("B1:HH48").options(ndim=2).value
    caller = application.books.add()
    caller.sheets[0].name = "I. Control"
    credit = caller.sheets.add("I. Credit Stresses")
    credit.range("B1").value = labels
    for address, grid in stress_values.items():
        credit.range(address).value = grid
    control = caller.sheets["I. Control"]
    control.range("D39").value = str(data)
    control.range("D40").value = str(map_data)
    control.range("D41").value = None
    control.range("D42").value = str(template)
    control.range("D9").value = "JRL"
    control.range("D11").value = datetime(2026, 9, 30)
    caller.save(str(folder / "Caller.xlsx"))
    return application, working, caller, data, template, folder, map_data


@pytest.fixture
def local_escalation_dependency(excel_workflow_case: tuple) -> list[int]:
    """Populate only the disposable template's local escalation lookup values."""
    from bonds_lite_input import formula_columns, required_links

    _, working, _, _, _, _, _ = excel_workflow_case
    sheet = working.sheets["I. Bonds Data"]
    table = sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    formulas = sheet.range("B9:AUN9").api.Formula2
    expected = [3, 45001, 45002, 0, 12, 6, 12, 6, 0]
    escalation = working.sheets["P. Bond Escalation file"]
    escalation.range("A2:J61561").clear_contents()
    escalation.range("A2:J3").value = [
        ["BONDS", *expected],
        ["MAP", *expected],
    ]
    assert sheet.range("B9:AUN9").api.Formula2 == formulas
    assert required_links(working, list(formula_columns(table))) == []
    working.save()
    return expected


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_excel_audited_stage_persists_handoff_and_invalidates_failed_rerun(
    excel_workflow_case: tuple,
    local_escalation_dependency: list[int],
) -> None:
    """Persist a successful audit, then block Stage 2 after a failed new attempt."""
    from bonds_lite_input import (
        _read_stage_fields,
        initialise_stage_panel,
        resolve_stage_input,
        run_stage_one,
    )

    application, working, caller, data, template, folder, map_data = excel_workflow_case
    working.close()
    originals = {path: path.read_bytes() for path in (data, map_data, template)}
    initialise_stage_panel(caller)
    output = run_stage_one(caller)
    fields = _read_stage_fields(caller)
    relative = output.relative_to(folder).as_posix()
    assert fields["status"] == "Succeeded"
    assert fields["last_output"] == fields["stage_two_input"] == relative
    assert fields["start"] <= fields["finish"] and fields["elapsed"] >= 0
    assert fields["user"]
    assert caller.api.Saved
    assert resolve_stage_input(caller) == output
    output_bytes = _pair_bytes(output)
    caller_path = Path(caller.fullname)
    caller.close()
    reopened = application.books.open(str(caller_path), update_links=False)
    try:
        assert _read_stage_fields(reopened)["stage_two_input"] == relative
        assert resolve_stage_input(reopened) == output
        reopened.sheets["I. Control"].range("D39").value = str(folder / "Missing.xlsx")
        with pytest.raises(BondsLiteError, match="does not exist"):
            run_stage_one(reopened)
        failed = _read_stage_fields(reopened)
        assert failed["status"] == "Failed"
        assert not failed["stage_two_input"]
        assert failed["last_output"] == relative
        assert reopened.api.Saved
        with pytest.raises(BondsLiteError):
            resolve_stage_input(reopened)
        assert _pair_bytes(output) == output_bytes
        assert {path: path.read_bytes() for path in originals} == originals
        assert set(output.parent.iterdir()) == {output}
    finally:
        reopened.close()


def _inject_target_dependency(working: Any, dependency: Path) -> None:
    """Use one target seed and its uniform column to exercise link discovery."""
    sheet = working.sheets["I. Bonds Data"]
    table = sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    seed = sheet.range("ATV9")
    seed.api.Formula2 = f"='{dependency.parent}\\[{dependency.name}]Sheet1'!$A$1"
    sheet.range(f"ATV9:ATV{8 + table.ListRows.Count}").api.Formula2R1C1 = (
        seed.api.Formula2R1C1
    )
    working.save()


class _SourceGuard:
    """Forward normal reads while overriding forbidden source operations."""

    def __init__(self, wrapped: Any, **overrides: Any) -> None:
        self.wrapped = wrapped
        self.overrides = overrides

    def __getattr__(self, name: str) -> Any:
        if name in self.overrides:
            return self.overrides[name]
        return getattr(self.wrapped, name)

    def __getitem__(self, key: Any) -> Any:
        item = self.wrapped[key]
        return self.overrides.get(item.name, item)

    def __iter__(self) -> Iterator[Any]:
        for item in self.wrapped:
            yield self.overrides.get(item.name, item)


def _assert_excel_workflow_publishes_validated_copy(
    excel_workflow_case: tuple,
    stress_values: dict[str, list[list[object]]],
    dependency: Path | None = None,
    expected_lookup: list[int] | None = None,
) -> None:
    """Exercise real refresh, calculation, persisted validation and publication."""
    from bonds_lite_input import generate_bonds_lite_input

    application, working, caller, data, template, folder, map_data = excel_workflow_case
    formulas = working.sheets["I. Bonds Data"].range("B9:AUN9").api.Formula2[0]
    assert all(
        isinstance(formulas[index], str) and formulas[index].startswith("=")
        for index in (8, 9)
    )
    working.save()
    working.close()
    paths = [
        data,
        map_data,
        template,
        Path(caller.fullname),
    ]
    if dependency is not None:
        paths.append(dependency)
    originals = {path: path.read_bytes() for path in paths}
    output = generate_bonds_lite_input(caller)
    assert output.parent == folder / "Bonds Lite output"
    assert output.is_dir()
    assert set(path.name for path in output.iterdir()) == {
        *EXPECTED_INPUT_FILENAMES.values(),
        "JRL Bonds",
        "JRL MAP",
    }
    for dataset, filename in EXPECTED_INPUT_FILENAMES.items():
        path = output / filename
        assert path.stat().st_size > 0
        assert filename not in [book.name for book in application.books]
        check = application.books.open(str(path), update_links=False, read_only=True)
        try:
            for address, grid in stress_values.items():
                assert (
                    check.sheets["I. Sensitivity"].range(address).options(ndim=2).value
                    == grid
                )
            assert check.sheets["I. Control"].range("D10").value == datetime(
                2026, 9, 30
            )
            assert check.sheets["I. Control"].range("D13").value == "JRL"
            model_folder = output / f"JRL {dataset}"
            named = check.names["output_folder_path"].refers_to_range
            assert named.sheet.name == "I. Control"
            assert named.value == str(model_folder.resolve())
            assert not list(model_folder.iterdir())
            bonds = check.sheets["I. Bonds Data"]
            table = bonds.api.ListObjects("I_tbl_Monthly_Bonds")
            rows = 1 if dataset == "Bonds" else 2
            assert table.ListRows.Count == rows and table.ListColumns.Count == 1235
            identifiers = (
                bonds.range(
                    (9, table.ListColumns("ISIN").Range.Column),
                    (8 + rows, table.ListColumns("ISIN").Range.Column),
                )
                .options(ndim=2)
                .value
            )
            prefix = "BONDS" if dataset == "Bonds" else "MAP"
            assert identifiers == [
                [f"{prefix}_FIXTURE_{index}"] for index in range(1, rows + 1)
            ]
            pv_column = next(
                column
                for column in table.ListColumns
                if re.search(r"(?i)^pv(?:$|[_ (])|present.?value", column.Name)
            )
            pvs = (
                bonds.range(
                    (9, pv_column.Range.Column), (8 + rows, pv_column.Range.Column)
                )
                .options(ndim=2)
                .value
            )
            assert pvs == ([[101.25]] if dataset == "Bonds" else [[202.5], [303.75]])
            assert bonds.range("J9:K9").api.Formula2[0] == formulas[8:10]
            assert bonds.range(f"J9:K{8 + rows}").api.HasFormula is True
            for column, formula in enumerate(formulas, start=2):
                if isinstance(formula, str) and formula.startswith("="):
                    assert (
                        bonds.range((9, column), (8 + rows, column)).api.HasFormula
                        is True
                    )
            assert bonds.range(f"ATU9:AUN{8 + rows}").api.HasFormula is True
            if expected_lookup is not None:
                assert (
                    [list(row) for row in bonds.range(f"ATV9:AUD{8 + rows}").api.Value2]
                    == [expected_lookup] * rows
                )
        finally:
            check.close()
    assert {path: path.read_bytes() for path in originals} == originals
    assert set(output.parent.iterdir()) == {output}
    assert not [path for path in output.iterdir() if path.name.startswith(".")]


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_excel_full_workflow_publishes_validated_copy(
    excel_workflow_case: tuple,
    local_escalation_dependency: list[int],
    stress_values: dict[str, list[list[object]]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Compute local lookups without refreshing or calculating source formulas."""
    import bonds_lite_input as workflow

    original_owned_book = workflow.owned_book
    opens: list[tuple[Path, bool]] = []
    protected_sources: list[Path] = []
    source_paths = {
        path.name: path for path in (excel_workflow_case[3], excel_workflow_case[6])
    }

    def forbidden(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Never refresh source links or calculate a source worksheet")

    @contextmanager
    def observe_open(
        application: Any,
        path: Path,
        read_only: bool = True,
    ) -> Iterator[Any]:
        opens.append((Path(path), read_only))
        with original_owned_book(application, path, read_only=read_only) as book:
            source_path = source_paths.get(Path(path).name) if read_only else None
            if source_path is None:
                yield book
                return
            protected_sources.append(source_path)
            sheet = book.sheets["I. Bonds Data"]
            rows = 1 if source_path == excel_workflow_case[3] else 2
            region = sheet.range(f"B8:ATT{8 + rows}")
            original = (region.api.Formula2, region.api.Value2)
            assert sheet.range("J9:K9").api.HasFormula is True

            def source_range(*args: Any, **kwargs: Any) -> Any:
                result = sheet.range(*args, **kwargs)
                assert result.column >= RAW_FIRST_COLUMN
                assert result.column + result.shape[1] - 1 <= RAW_LAST_COLUMN
                return result

            guarded_sheet = _SourceGuard(
                sheet,
                range=source_range,
                api=_SourceGuard(sheet.api, Calculate=forbidden),
            )
            try:
                yield _SourceGuard(
                    book,
                    api=_SourceGuard(book.api, UpdateLink=forbidden),
                    sheets=_SourceGuard(
                        book.sheets, **{"I. Bonds Data": guarded_sheet}
                    ),
                )
            finally:
                assert (region.api.Formula2, region.api.Value2) == original

    monkeypatch.setattr(workflow, "owned_book", observe_open)
    _assert_excel_workflow_publishes_validated_copy(
        excel_workflow_case,
        stress_values,
        expected_lookup=local_escalation_dependency,
    )
    assert protected_sources == [excel_workflow_case[3], excel_workflow_case[6]]
    assert len(opens) == 4, f"Expected two source/copy pairs, observed: {opens}"
    assert opens[0][0] != excel_workflow_case[3]
    assert opens[0][0].name == excel_workflow_case[3].name and opens[0][1] is True
    assert opens[2][0] != excel_workflow_case[6]
    assert opens[2][0].name == excel_workflow_case[6].name and opens[2][1] is True
    assert opens[1][1] is False and opens[3][1] is False
    assert opens[1][0].name != opens[3][0].name
    assert not opens[0][0].exists() and not opens[2][0].exists()
    for path, _ in (opens[1], opens[3]):
        assert path.name != "working.xlsx"
        assert path.name not in EXPECTED_INPUT_FILENAMES.values()
        assert path.is_relative_to(excel_workflow_case[5] / "Bonds Lite output")
        assert not path.exists()


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_excel_full_workflow_refreshes_injected_local_dependency(
    excel_workflow_case: tuple,
    stress_values: dict[str, list[list[object]]],
    local_escalation_dependency: list[int],
) -> None:
    """Exercise required-link refresh without relying on any real network path."""
    from bonds_lite_input import formula_columns, required_links

    application, working, _, _, _, folder, _ = excel_workflow_case
    dependency = folder / "Synthetic dependency.xlsx"
    local = application.books.add()
    try:
        local.sheets[0].name = "Sheet1"
        local.sheets[0].range("A1").value = local_escalation_dependency[0]
        local.save(str(dependency))
    finally:
        local.close()
    _inject_target_dependency(working, dependency)
    columns = formula_columns(
        working.sheets["I. Bonds Data"].api.ListObjects("I_tbl_Monthly_Bonds")
    )
    assert required_links(working, list(columns)) == [str(dependency)]
    _assert_excel_workflow_publishes_validated_copy(
        excel_workflow_case,
        stress_values,
        dependency,
        local_escalation_dependency,
    )


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
def test_excel_full_workflow_missing_dependency_does_not_publish(
    excel_workflow_case: tuple,
    local_escalation_dependency: list[int],
) -> None:
    """Report the missing local dependency and remove every temporary output."""
    from bonds_lite_input import generate_bonds_lite_input

    _, working, caller, data, template, folder, map_data = excel_workflow_case
    dependency = folder / "Missing dependency.xlsx"
    _inject_target_dependency(working, dependency)
    working.close()
    assert not dependency.exists()
    originals = {
        path: path.read_bytes()
        for path in (
            data,
            map_data,
            template,
            Path(caller.fullname),
        )
    }
    with pytest.raises(BondsLiteError) as caught:
        generate_bonds_lite_input(caller)
    assert "Required formula dependency is unavailable" in str(caught.value)
    assert dependency.name in str(caught.value)
    assert {path: path.read_bytes() for path in originals} == originals
    output_folder = folder / "Bonds Lite output"
    assert not output_folder.exists() or not list(output_folder.iterdir())
    assert not list(folder.rglob("working.xlsx"))


def test_workbook_run_guard_rejects_duplicate_and_releases(tmp_path: Path) -> None:
    from bonds_lite_input import workbook_run_guard

    with workbook_run_guard(tmp_path / "Caller.xlsm"):
        with pytest.raises(BondsLiteError, match="already"):
            with workbook_run_guard(tmp_path / "Caller.xlsm"):
                pytest.fail("Duplicate run entered")
    with workbook_run_guard(tmp_path / "Caller.xlsm"):
        pass


@pytest.mark.excel
@pytest.mark.skipif(not EXCEL_ENABLED, reason="Set BONDS_LITE_RUN_EXCEL_TESTS=1")
@pytest.mark.parametrize("controls", ["defaults", "entered"])
def test_install_button_preserves_backup_and_saves_launcher(
    tmp_path: Path,
    controls: str,
) -> None:
    """Install on a disposable Draft copy without modifying the supplied file."""
    from zipfile import ZipFile

    import xlwings as xw

    from bonds_lite_input import (
        STAGE_ONE_FIELD_NAMES,
        initialise_stage_panel,
        install_button,
    )

    root = Path(__file__).resolve().parents[1]
    caller = tmp_path / "Draft caller.xlsm"
    shutil.copy2(root / "Draft Input spreadsheet for SST.xlsm", caller)
    with xw.App(visible=False, add_book=False) as application:
        application.api.AutomationSecurity = 3
        application.api.EnableEvents = False
        application.display_alerts = False
        book = application.books.open(str(caller), update_links=False)
        try:
            initialise_stage_panel(book)
            sheet = book.sheets["I. Control"]
            panel = book.names["SST_Stage1_Panel"].refers_to_range
            panel.api.Cut(Destination=sheet.range("Q100:X110").api)
            book.names[STAGE_ONE_FIELD_NAMES["user"]].refers_to_range[
                0, 0
            ].value = "Installer user sentinel"
            entered = {
                "D39": "Custom Bonds.xlsx",
                "D40": "Custom MAP.xlsx",
                "D42": "Custom shared template.xlsx",
            }
            for address in ("D39", "D40", "D41", "D42"):
                sheet.range(address).value = (
                    entered.get(address) if controls == "entered" else None
                )
            book.save()
            panel_values = book.names["SST_Stage1_Panel"].refers_to_range.formula
            panel_names = {
                name: book.names[name].refers_to
                for name in ("SST_Stage1_Panel", *STAGE_ONE_FIELD_NAMES.values())
            }
        finally:
            book.close()
    original = caller.read_bytes()
    backup = install_button(caller)
    assert backup.read_bytes() == original
    with ZipFile(caller) as archive:
        assert "xl/vbaProject.bin" in archive.namelist()
        drawings = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/drawings/") and name.endswith((".xml", ".vml"))
        )
        assert b"Generate Bonds Lite Input" in drawings
        assert b"GenerateBondsLiteInput" in drawings
    with xw.App(visible=False, add_book=False) as application:
        application.api.AutomationSecurity = 3
        application.api.EnableEvents = False
        application.display_alerts = False
        book = application.books.open(str(caller), update_links=False, read_only=True)
        module = None
        try:
            sheet = book.sheets["I. Control"]
            assert sheet.range("D41").value in (None, "")
            if controls == "defaults":
                assert sheet.range("D40").value in (None, "")
                assert sheet.range("D42").value == "Bonds Lite input template.xlsx"
            else:
                assert {address: sheet.range(address).value for address in entered} == entered
            assert {name: book.names[name].refers_to for name in panel_names} == panel_names
            assert book.names["SST_Stage1_Panel"].refers_to_range.formula == panel_values
            module = book.api.VBProject.VBComponents.Item(
                "modBondsLiteLauncher"
            ).CodeModule
            macro = module.Lines(1, module.CountOfLines)
            assert "shell.Run command, 1, False" in macro
            assert "-u -m bonds_lite_input" in macro
            assert "--pause" in macro
        finally:
            module = None
            book.close()


def test_stage_two_pinned_interfaces_exist() -> None:
    """Keep missing Stage 2 APIs as test failures, not collection errors."""
    import bonds_lite_input as workflow

    expected = (
        "ModelConfig", "ModelInput", "load_model_config", "read_model_inputs",
        "preflight_model_environment", "run_model_batch", "validate_model_results",
        "model_results_transaction", "confirm_model_overwrite", "run_stage_two",
        "initialise_stage_two_panel", "require_stage_two_panel",
        "_read_stage_two_fields", "write_stage_two_audit",
    )
    missing = [name for name in expected if not callable(getattr(workflow, name, None))]
    assert not missing, f"Missing Stage 2 interfaces: {', '.join(missing)}"


def test_stage_two_field_names_do_not_repurpose_stage_one_handoff() -> None:
    """Give attempts their own identity without overwriting the Stage 1 link."""
    import bonds_lite_input as workflow

    assert workflow.STAGE_TWO_FIELD_NAMES == {
        "last_run": "SST_Stage2_LastRun",
        "start": "SST_Stage2_Start",
        "finish": "SST_Stage2_Finish",
        "elapsed": "SST_Stage2_Elapsed",
        "user": "SST_Stage2_User",
        "status": "SST_Stage2_Status",
        "input_run": "SST_Stage2_InputRun",
        "last_output": "SST_Stage2_LastSuccessfulResults",
    }
    assert "SST_Stage2_Input" not in workflow.STAGE_TWO_FIELD_NAMES.values()
    assert workflow.STAGE_ONE_FIELD_NAMES["stage_two_input"] == "SST_Stage2_Input"


@pytest.fixture
def stage_two_no_excel(monkeypatch: pytest.MonkeyPatch) -> None:
    """Reject accidental Excel creation in every ordinary Stage 2 fixture."""
    def forbidden(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Ordinary Stage 2 tests must not open or attach to real Excel")

    monkeypatch.setitem(
        sys.modules, "xlwings", SimpleNamespace(App=forbidden, Book=forbidden, apps=[]),
    )


@pytest.fixture
def stage_two_config_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, stage_two_no_excel: None,
) -> Any:
    """Use local path sentinels, never the configured external model environment."""
    import bonds_lite_input as workflow

    monkeypatch.setattr(workflow, "_user_environment_value", lambda name: "")
    caller = tmp_path / "Sarah Walker" / "Testing" / "Draft.xlsm"
    caller.parent.mkdir(parents=True)
    caller.write_bytes(b"Caller sentinel")
    for filename in EXPECTED_INPUT_FILENAMES.values():
        (caller.parent / filename).write_bytes(filename.encode("ascii"))
    environment = tmp_path / "separate py311 environment"
    environment.mkdir()
    executable = environment / "python.exe"
    executable.write_bytes(b"Not executable: configuration tests only")
    model = tmp_path / "external model with spaces"
    model.mkdir()
    script = model / "run_bonds_batch.py"
    script.write_bytes(b"raise AssertionError('Do not execute configuration fixtures')")
    monkeypatch.setenv("BONDS_LITE_MODEL_PYTHON", str(executable))
    monkeypatch.setenv("BONDS_LITE_MODEL_DIR", str(model))
    monkeypatch.setenv("BONDS_LITE_PYTHON", "controller-must-not-be-used.exe")
    try:
        import psutil
    except ImportError:
        pass
    else:
        monkeypatch.setattr(
            psutil, "Process",
            lambda *args, **kwargs: pytest.fail("A fake child has no real process-tree PID"),
        )
    return SimpleNamespace(
        caller=caller, environment=environment, executable=executable,
        model=model, script=script,
    )


@pytest.mark.parametrize("environment_directory", [False, True])
def test_load_model_config_resolves_separate_interpreter_and_frozen_default_timeout(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch,
    environment_directory: bool,
) -> None:
    """Resolve an environment directory explicitly and keep the config immutable."""
    from dataclasses import FrozenInstanceError

    import bonds_lite_input as workflow

    state = stage_two_config_paths
    if environment_directory:
        monkeypatch.setenv("BONDS_LITE_MODEL_PYTHON", str(state.environment))
    config = workflow.load_model_config(state.caller)
    assert isinstance(config, workflow.ModelConfig)
    assert config.python_executable == state.executable.resolve()
    assert config.model_dir == state.model.resolve()
    assert config.timeout_seconds == 3600.0
    with pytest.raises(FrozenInstanceError):
        config.timeout_seconds = 1.0


def test_load_model_config_uses_user_registry_when_excel_environment_is_stale(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_config_paths
    monkeypatch.delenv("BONDS_LITE_MODEL_PYTHON")
    monkeypatch.delenv("BONDS_LITE_MODEL_DIR")
    user_settings = {
        "BONDS_LITE_MODEL_PYTHON": str(state.executable),
        "BONDS_LITE_MODEL_DIR": str(state.model),
    }
    monkeypatch.setattr(
        workflow, "_user_environment_value",
        lambda name: user_settings.get(name, ""), raising=False,
    )

    config = workflow.load_model_config(state.caller)

    assert config.python_executable == state.executable.resolve()
    assert config.model_dir == state.model.resolve()


@pytest.mark.parametrize("setting", [None, "", "   ", "missing.exe", "pythonw.exe"])
def test_load_model_config_rejects_missing_or_invalid_python_without_fallback(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch,
    setting: str | None,
) -> None:
    """Never fall back to the controller interpreter or silently pick pythonw."""
    import bonds_lite_input as workflow

    state = stage_two_config_paths
    if setting is None:
        monkeypatch.delenv("BONDS_LITE_MODEL_PYTHON")
    elif setting == "pythonw.exe":
        wrong = state.environment / setting
        wrong.write_bytes(b"Wrong executable")
        monkeypatch.setenv("BONDS_LITE_MODEL_PYTHON", str(wrong))
    else:
        monkeypatch.setenv("BONDS_LITE_MODEL_PYTHON", setting)
    with pytest.raises(BondsLiteError, match=r"(?i)python|interpreter"):
        workflow.load_model_config(state.caller)


@pytest.mark.parametrize("defect", ["directory_without_python", "missing_model", "missing_script"])
def test_load_model_config_rejects_incomplete_environment_or_model(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch, defect: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_config_paths
    if defect == "directory_without_python":
        state.executable.unlink()
        monkeypatch.setenv("BONDS_LITE_MODEL_PYTHON", str(state.environment))
    elif defect == "missing_model":
        monkeypatch.setenv("BONDS_LITE_MODEL_DIR", str(state.model / "missing"))
    else:
        state.script.unlink()
    with pytest.raises(BondsLiteError):
        workflow.load_model_config(state.caller)


def test_load_model_config_default_model_is_relative_to_caller_not_cwd(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_config_paths
    expected = (
        state.caller.parent / "../../02 Sensitivities/2025/05 ORSA/Models/"
        "Bonds Lite Batch/bonds-hm-cst"
    ).resolve()
    expected.mkdir(parents=True)
    (expected / "run_bonds_batch.py").write_bytes(b"Do not run")
    monkeypatch.delenv("BONDS_LITE_MODEL_DIR")
    monkeypatch.chdir(tmp_path)
    assert workflow.load_model_config(state.caller).model_dir == expected


@pytest.fixture
def stage_two_saved_inputs(tmp_path: Path, stage_two_no_excel: None) -> Any:
    """Reproduce only verified named scalars and the four-row filename mapping."""
    import hashlib

    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    folder = tmp_path / "Bonds Lite output" / "20261028_090000"
    folder.mkdir(parents=True)
    scenario_folder = tmp_path / "controlled model scenarios"
    scenario_folder.mkdir()
    (scenario_folder / "GBP_RPI.csv").write_text("year,scenario\n0,0\n", encoding="utf-8")
    mapping = {
        "ISINs": "fixture_identifiers",
        "Grouped": "fixture_grouped",
        "BaseScenario": "fixture_base",
        "GroupedNoZeros": "fixture_conditional",
    }
    addresses = {
        "entity_name": "D13", "output_folder_path": "D19",
        "output_suffix": "D25", "bool_inv_exp": "D28",
        "risk_factor_path": "D16",
    }
    for dataset, filename in EXPECTED_INPUT_FILENAMES.items():
        destination = folder / f"JRL {dataset}"
        destination.mkdir()
        book = openpyxl.Workbook()
        try:
            control = book.active
            control.title = "I. Control"
            values = {
                "entity_name": "JRL", "output_folder_path": str(destination),
                "output_suffix": f"fixture_{dataset}", "bool_inv_exp": False,
                "risk_factor_path": str(scenario_folder),
            }
            for name, address in addresses.items():
                control[address] = values[name]
                book.defined_names.add(DefinedName(
                    name, attr_text=f"'I. Control'!${address[0]}${address[1:]}",
                ))
            parameters = book.create_sheet("P. Parameter")
            parameters.append(["OutputType", "Filename"])
            for key, value in mapping.items():
                parameters.append([key, value])
            parameters.add_table(Table(displayName="P_tbl_OutputFilenames", ref="A1:B5"))
            parameters["D1"] = "Curve Name"
            parameters["E1"] = "File Name"
            parameters["D2"] = "RPI"
            parameters["E2"] = "GBP_RPI"
            parameters.add_table(Table(displayName="P_tbl_OtherCurves", ref="D1:E2"))
            parameters["G1"] = "Curve Name"
            parameters["H1"] = "File Name"
            parameters.add_table(Table(displayName="P_tbl_SpreadCurves", ref="G1:H1"))
            base_curves = book.create_sheet("P. Base Curves")
            base_curves.append(["Tenor", "GBP SONIA"])
            base_curves.append([1, 0.01])
            base_curves.append([2, 0.02])
            base_curves.add_table(Table(displayName="P_tbl_BaseCurves", ref="A1:B3"))
            book.save(folder / filename)
        finally:
            book.close()
    hashes = {
        dataset: hashlib.sha256((folder / filename).read_bytes()).hexdigest()
        for dataset, filename in EXPECTED_INPUT_FILENAMES.items()
    }
    return SimpleNamespace(
        folder=folder, mapping=mapping, addresses=addresses, hashes=hashes,
        scenario_folder=scenario_folder,
    )


def test_read_model_inputs_uses_saved_names_mapping_and_never_changes_sources(
    stage_two_saved_inputs: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_saved_inputs
    before = _pair_bytes(state.folder)
    timestamps = {path: path.stat().st_mtime_ns for path in state.folder.glob("*.xlsx")}
    inputs = workflow.read_model_inputs(state.folder)
    assert set(inputs) == {"Bonds", "MAP"}
    for dataset, metadata in inputs.items():
        assert isinstance(metadata, workflow.ModelInput)
        assert metadata.dataset == dataset
        assert metadata.input_path == state.folder / EXPECTED_INPUT_FILENAMES[dataset]
        assert metadata.output_folder == state.folder / f"JRL {dataset}"
        assert metadata.output_suffix == f"fixture_{dataset}"
        assert metadata.input_hash == state.hashes[dataset]
        assert isinstance(metadata.required_files, tuple)
        required = set(metadata.required_files)
        for key in ("ISINs", "Grouped", "BaseScenario"):
            assert f"{state.mapping[key]}_fixture_{dataset}.csv" in required
        assert any("fixture_base" in name and "nozeros" in name.lower().replace("_", "")
                   for name in required)
        assert f"py_bonds_input_file_fixture_{dataset}.xlsx" in required
        assert any("invexp_base" in name.lower() for name in required)
        assert not any("fixture_conditional" in name for name in required)
        assert all(Path(name).name == name for name in required)
    assert _pair_bytes(state.folder) == before
    assert {path: path.stat().st_mtime_ns for path in timestamps} == timestamps


@pytest.mark.parametrize("bad_value", [None, "#NAME?"])
def test_read_model_inputs_rejects_missing_or_error_cached_gbp_sonia(
    stage_two_saved_inputs: Any, bad_value: object,
) -> None:
    import openpyxl

    import bonds_lite_input as workflow

    state = stage_two_saved_inputs
    path = state.folder / EXPECTED_INPUT_FILENAMES["Bonds"]
    workbook = openpyxl.load_workbook(path)
    try:
        cell = workbook["P. Base Curves"]["B2"]
        cell.value = bad_value
        if bad_value == "#NAME?":
            cell.data_type = "e"
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(BondsLiteError, match=r"(?i)GBP SONIA.*numeric"):
        workflow.read_model_inputs(state.folder)


@pytest.mark.parametrize("extra", ["lock", "nested", "root"])
def test_read_model_inputs_only_allows_the_exact_two_processable_root_workbooks(
    stage_two_saved_inputs: Any, extra: str,
) -> None:
    import bonds_lite_input as workflow

    folder = stage_two_saved_inputs.folder
    extra_path = {
        "lock": folder / "~$Bonds Lite input_Bonds.xlsx",
        "nested": folder / "JRL Bonds" / "py_bonds_input_file_old.xlsx",
        "root": folder / "third input.xlsx",
    }[extra]
    extra_path.write_bytes(b"Not an input workbook")
    before = _pair_bytes(folder)
    if extra == "root":
        with pytest.raises(BondsLiteError, match=r"(?i)two|extra|unexpected|exact"):
            workflow.read_model_inputs(folder)
    else:
        assert set(workflow.read_model_inputs(folder)) == {"Bonds", "MAP"}
    assert extra_path.read_bytes() == b"Not an input workbook"
    assert _pair_bytes(folder) == before


@pytest.mark.parametrize("dataset", ["Bonds", "MAP"])
@pytest.mark.parametrize("defect", ["missing", "empty", "directory"])
def test_read_model_inputs_rejects_missing_or_invalid_pair_member(
    stage_two_saved_inputs: Any, dataset: str, defect: str,
) -> None:
    import bonds_lite_input as workflow

    path = stage_two_saved_inputs.folder / EXPECTED_INPUT_FILENAMES[dataset]
    path.unlink()
    if defect == "empty":
        path.touch()
    elif defect == "directory":
        path.mkdir()
    with pytest.raises(BondsLiteError):
        workflow.read_model_inputs(stage_two_saved_inputs.folder)


@pytest.mark.parametrize("defect", [
    "outside", "wrong_dataset", "wrong_entity", "entity_traversal",
    "suffix_traversal", "suffix_absolute", "mapping_traversal", "mapping_absolute",
    "missing_name", "multi_cell", "foreign_name", "formula_name", "missing_table",
])
def test_read_model_inputs_rejects_unsafe_saved_metadata_without_rewriting(
    stage_two_saved_inputs: Any, defect: str, tmp_path: Path,
) -> None:
    import openpyxl

    import bonds_lite_input as workflow

    state = stage_two_saved_inputs
    path = state.folder / EXPECTED_INPUT_FILENAMES["Bonds"]
    book = openpyxl.load_workbook(path)
    try:
        control = book["I. Control"]
        if defect == "outside":
            outside = tmp_path / "Other user" / "JRL Bonds"
            outside.mkdir(parents=True)
            control["D19"] = str(outside)
        elif defect == "wrong_dataset":
            control["D19"] = str(state.folder / "JRL MAP")
        elif defect == "wrong_entity":
            control["D13"] = "Another entity"
        elif defect == "entity_traversal":
            control["D13"] = "../JRL"
        elif defect.startswith("suffix_"):
            control["D25"] = "../escape" if defect.endswith("traversal") else "C:/escape"
        elif defect.startswith("mapping_"):
            book["P. Parameter"]["B2"] = (
                "../escape" if defect.endswith("traversal") else "C:/escape"
            )
        elif defect == "missing_name":
            del book.defined_names["output_folder_path"]
        elif defect == "multi_cell":
            book.defined_names["output_folder_path"].attr_text = "'I. Control'!$D$19:$D$20"
        elif defect == "foreign_name":
            book.defined_names["output_folder_path"].attr_text = "'[Other.xlsx]I. Control'!$D$19"
        elif defect == "formula_name":
            book.defined_names["output_folder_path"].attr_text = "=1+1"
        else:
            del book["P. Parameter"].tables["P_tbl_OutputFilenames"]
        book.save(path)
    finally:
        book.close()
    before = _pair_bytes(state.folder)
    with pytest.raises(BondsLiteError):
        workflow.read_model_inputs(state.folder)
    assert _pair_bytes(state.folder) == before
    assert not list(state.folder.parent.glob("*.csv"))


@pytest.fixture
def stage_two_results_case(tmp_path: Path, stage_two_no_excel: None) -> Any:
    """Use explicit ModelInput records to isolate result policy from XLSX parsing."""
    import hashlib

    import bonds_lite_input as workflow

    folder = _complete_run(tmp_path / "Bonds Lite output" / "20261028_090000")
    inputs = {}
    for dataset, filename in EXPECTED_INPUT_FILENAMES.items():
        target = folder / f"JRL {dataset}"
        target.mkdir()
        suffix = f"fixture_{dataset}"
        path = folder / filename
        required = (
            f"custom_ISINs_{suffix}.csv", f"custom_Grouped_{suffix}.csv",
            f"custom_BaseScenario_{suffix}.csv",
            f"custom_BaseScenario_nozeros_{suffix}.csv",
            f"py_bonds_InvExp_Base_{suffix}.csv",
            f"py_bonds_input_file_{suffix}.xlsx",
        )
        inputs[dataset] = workflow.ModelInput(
            dataset=dataset, input_path=path, output_folder=target,
            output_suffix=suffix, required_files=required,
            input_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
        )
    return SimpleNamespace(folder=folder, inputs=inputs)


def _stage_two_summary_rows(inputs: dict[str, Any]) -> list[dict[str, str]]:
    """Mirror the verified external runner's six summary columns."""
    return [
        {
            "input_file": item.input_path.name, "status": "success", "seconds": "0.01",
            "output_folder_path": str(item.output_folder),
            "output_suffix": item.output_suffix, "error": "",
        }
        for item in inputs.values()
    ]


def _write_stage_two_summary(
    folder: Path, rows: list[dict[str, str]],
    name: str = "bonds_batch_summary_20261028_090000.csv",
    columns: tuple[str, ...] | None = None,
) -> Path:
    """Write structured CSV, including quoted paths and error messages."""
    import csv

    path = folder / name
    headers = columns or (
        "input_file", "status", "seconds", "output_folder_path", "output_suffix", "error",
    )
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


def _write_stage_two_artifacts(inputs: dict[str, Any]) -> None:
    """Create required files only; conditional diagnostics are not success gates."""
    for item in inputs.values():
        for name in item.required_files:
            (item.output_folder / name).write_bytes(
                b"\n" if "InvExp_Base" in name else b"Fresh result sentinel"
            )


def _stage_two_tree(folder: Path) -> dict[str, bytes | None]:
    """Capture files and empty directories for byte-for-byte rollback assertions."""
    return {
        path.relative_to(folder).as_posix(): path.read_bytes() if path.is_file() else None
        for path in folder.rglob("*")
    }


def test_stage_two_model_input_is_frozen(stage_two_results_case: Any) -> None:
    from dataclasses import FrozenInstanceError

    item = stage_two_results_case.inputs["Bonds"]
    with pytest.raises(FrozenInstanceError):
        item.output_folder = item.output_folder.parent


def test_validate_model_results_requires_both_datasets_but_not_conditional_outputs(
    stage_two_results_case: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    before = _pair_bytes(state.folder)
    _write_stage_two_artifacts(state.inputs)
    summary = _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    assert workflow.validate_model_results(state.folder, state.inputs) == summary
    assert _pair_bytes(state.folder) == before
    assert not list(state.folder.glob("*.txt"))


@pytest.mark.parametrize("defect", [
    "missing", "duplicate_summary", "missing_column", "missing_map", "duplicate_bonds",
    "extra_row", "failed_map", "nonblank_error", "wrong_input", "wrong_target",
    "wrong_suffix", "input_changed", "summary_directory",
])
def test_validate_model_results_rejects_false_success_even_after_zero_exit(
    stage_two_results_case: Any, defect: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    _write_stage_two_artifacts(state.inputs)
    rows = _stage_two_summary_rows(state.inputs)
    if defect == "missing_map":
        rows.pop()
    elif defect == "duplicate_bonds":
        rows[1] = rows[0].copy()
    elif defect == "extra_row":
        rows.append({**rows[0], "input_file": "Unexpected.xlsx"})
    elif defect == "failed_map":
        rows[1].update(status="failed", error="Model failed, input invalid")
    elif defect == "nonblank_error":
        rows[1]["error"] = "Caught model exception"
    elif defect == "wrong_input":
        rows[1]["input_file"] = "another input.xlsx"
    elif defect == "wrong_target":
        rows[1]["output_folder_path"] = str(state.inputs["Bonds"].output_folder)
    elif defect == "wrong_suffix":
        rows[1]["output_suffix"] = "another suffix"
    elif defect == "input_changed":
        state.inputs["MAP"].input_path.write_bytes(b"Changed by model")
    columns = tuple(rows[0])[:-1] if defect == "missing_column" else None
    if defect != "missing":
        summary = _write_stage_two_summary(state.folder, rows, columns=columns)
        if defect == "duplicate_summary":
            _write_stage_two_summary(
                state.folder, rows, name="bonds_batch_summary_20261028_090001.csv",
            )
        elif defect == "summary_directory":
            summary.unlink()
            summary.mkdir()
    before = _stage_two_tree(state.folder)
    with pytest.raises(BondsLiteError) as caught:
        workflow.validate_model_results(state.folder, state.inputs)
    assert str(caught.value).strip()
    assert _stage_two_tree(state.folder) == before


@pytest.mark.parametrize("dataset", ["Bonds", "MAP"])
@pytest.mark.parametrize("artifact", range(6))
@pytest.mark.parametrize("defect", ["missing", "empty", "directory"])
def test_validate_model_results_rejects_each_missing_empty_or_nonfile_artifact(
    stage_two_results_case: Any, dataset: str, artifact: int, defect: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    _write_stage_two_artifacts(state.inputs)
    _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    item = state.inputs[dataset]
    path = item.output_folder / item.required_files[artifact]
    path.unlink()
    if defect == "empty":
        path.touch()
    elif defect == "directory":
        path.mkdir()
    with pytest.raises(BondsLiteError, match=re.escape(path.name)):
        workflow.validate_model_results(state.folder, state.inputs)


@pytest.mark.parametrize("answer", ["", "n", "no", "yes", "Y"])
@pytest.mark.parametrize("existing", ["bonds", "map", "summary", "none"])
def test_confirm_model_overwrite_defaults_no_and_only_prompts_for_existing_results(
    stage_two_results_case: Any, monkeypatch: pytest.MonkeyPatch,
    answer: str, existing: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    if existing in ("bonds", "map"):
        item = state.inputs["Bonds" if existing == "bonds" else "MAP"]
        (item.output_folder / "old.csv").write_bytes(b"Previous result")
    elif existing == "summary":
        _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    prompts = []

    def respond(prompt: str = "") -> str:
        prompts.append(prompt)
        return answer

    monkeypatch.setattr("builtins.input", respond)
    before = _stage_two_tree(state.folder)
    expected = existing == "none" or answer.casefold() in {"y", "yes"}
    assert workflow.confirm_model_overwrite(state.inputs, state.folder) is expected
    assert len(prompts) == (0 if existing == "none" else 1)
    assert _stage_two_tree(state.folder) == before


@pytest.mark.parametrize("committed", [False, True])
@pytest.mark.parametrize("failure", [None, RuntimeError, KeyboardInterrupt])
@pytest.mark.parametrize("rerun", [False, True])
def test_model_results_transaction_restores_old_pair_or_commits_only_on_clean_exit(
    stage_two_results_case: Any, committed: bool,
    failure: type[BaseException] | None, rerun: bool,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    unrelated = state.folder / "unrelated folder"
    unrelated.mkdir()
    (unrelated / "keep.csv").write_bytes(b"Unrelated nested file")
    (state.folder / "keep.txt").write_bytes(b"Unrelated root file")
    if rerun:
        for dataset, item in state.inputs.items():
            nested = item.output_folder / "old nested"
            nested.mkdir()
            (nested / "old.bin").write_bytes(f"Old {dataset}".encode())
            (item.output_folder / "old.csv").write_bytes(dataset.encode())
        _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    before = _stage_two_tree(state.folder)
    parent_before = set(state.folder.parent.iterdir())
    original = failure("Transaction sentinel") if failure else None
    try:
        with workflow.model_results_transaction(state.folder, state.inputs) as transaction:
            assert all(item.output_folder.is_dir() for item in state.inputs.values())
            assert all(not list(item.output_folder.iterdir()) for item in state.inputs.values())
            assert not list(state.folder.glob("bonds_batch_summary*.csv"))
            assert _pair_bytes(state.folder) == {
                name: before[name] for name in EXPECTED_INPUT_FILENAMES.values()
            }
            if rerun:
                siblings = set(state.folder.parent.iterdir()) - parent_before
                assert siblings, "Rollback storage must be outside the selected run folder"
                assert all(path.stat().st_dev == state.folder.stat().st_dev for path in siblings)
                assert not any(path.name.startswith(".") for path in state.folder.iterdir())
            _write_stage_two_artifacts(state.inputs)
            _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
            if committed:
                assert transaction.commit() is None
            if original:
                raise original
    except BaseException as caught:
        if original is None or caught is not original:
            raise
    else:
        assert original is None, "A transaction must not suppress the original failure"
    if committed and original is None:
        assert not list(state.folder.glob("*/old.csv"))
        assert all((item.output_folder / name).is_file()
                   for item in state.inputs.values() for name in item.required_files)
        assert len(list(state.folder.glob("bonds_batch_summary*.csv"))) == 1
        assert (state.folder / "keep.txt").read_bytes() == before["keep.txt"]
        assert (unrelated / "keep.csv").read_bytes() == before["unrelated folder/keep.csv"]
        assert _pair_bytes(state.folder) == {
            name: before[name] for name in EXPECTED_INPUT_FILENAMES.values()
        }
    else:
        assert _stage_two_tree(state.folder) == before
    assert set(state.folder.parent.iterdir()) == parent_before


def test_model_results_transaction_rollback_failure_retains_recoverable_backup(
    stage_two_results_case: Any, monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture, capsys: pytest.CaptureFixture[str],
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    for dataset, item in state.inputs.items():
        (item.output_folder / "old.csv").write_bytes(f"Old {dataset}".encode())
    old_summary = _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    summary_bytes = old_summary.read_bytes()
    before_siblings = set(state.folder.parent.iterdir())
    original_rename, original_replace, original_move = Path.rename, Path.replace, shutil.move

    def deny_restore(source: Path) -> None:
        if not Path(source).is_relative_to(state.folder):
            raise PermissionError("Rollback restoration denied sentinel")

    def rename(source: Path, target: Path) -> Path:
        deny_restore(source)
        return original_rename(source, target)

    def replace(source: Path, target: Path) -> Path:
        deny_restore(source)
        return original_replace(source, target)

    def move(source: Any, target: Any, *args: Any, **kwargs: Any) -> Any:
        deny_restore(Path(source))
        return original_move(source, target, *args, **kwargs)

    with pytest.raises((BondsLiteError, PermissionError, RuntimeError)):
        with workflow.model_results_transaction(state.folder, state.inputs):
            monkeypatch.setattr(Path, "rename", rename)
            monkeypatch.setattr(Path, "replace", replace)
            monkeypatch.setattr(shutil, "move", move)
            raise RuntimeError("Model failed sentinel")
    backups = set(state.folder.parent.iterdir()) - before_siblings
    assert backups, "A failed rollback must retain the only recovery copy"
    recovered_bytes = [path.read_bytes() for backup in backups for path in backup.rglob("*")
                       if path.is_file()]
    assert b"Old Bonds" in recovered_bytes and b"Old MAP" in recovered_bytes
    assert summary_bytes in recovered_bytes
    captured = capsys.readouterr()
    message = caplog.text + captured.out + captured.err
    assert any(str(path) in message for path in backups), "Report the recovery location"


@pytest.fixture
def stage_two_attempt(
    stage_two_results_case: Any, monkeypatch: pytest.MonkeyPatch,
) -> Any:
    """Fake owned Excel boundaries while retaining the real result transaction."""
    import getpass
    import subprocess

    import bonds_lite_input as workflow

    state = stage_two_results_case
    state.caller_path = state.folder.parent.parent / "Draft.xlsm"
    state.caller_path.write_bytes(b"Caller sentinel")
    state.relative = state.folder.relative_to(state.caller_path.parent).as_posix()
    state.previous = "Bonds Lite output/20261027_090000"
    state.fields = {
        "status": "Succeeded", "input_run": state.previous,
        "last_output": state.previous, "finish": datetime(2026, 10, 27), "elapsed": 0.1,
    }
    state.legacy_handoff = state.relative
    state.clock = 100.0
    state.events, state.writes, state.saves, state.guard_calls = [], [], [], []
    state.locked, state.approved = False, True
    state.fail_saves = set()
    state.preflight_error = state.batch_error = state.validation_error = None
    state.config = workflow.ModelConfig(
        python_executable=Path(sys.executable), model_dir=state.folder.parent,
    )
    caller = SimpleNamespace(
        fullname=str(state.caller_path), api=SimpleNamespace(Saved=True),
        app=SimpleNamespace(),
    )
    state.caller = caller
    for dataset, item in state.inputs.items():
        (item.output_folder / "old.csv").write_bytes(f"Old {dataset}".encode())
    _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    state.before = _stage_two_tree(state.folder)

    @contextmanager
    def guard(path: Path) -> Iterator[None]:
        assert path == state.caller_path and not state.locked
        state.guard_calls.append(path)
        state.locked = True
        state.events.append("lock")
        try:
            yield
        finally:
            state.events.append("release")
            state.locked = False

    def require(book: Any) -> None:
        assert book is caller and state.locked
        state.events.append("require")

    def read(book: Any) -> dict[str, Any]:
        assert book is caller and state.locked
        state.events.append("read")
        return state.fields.copy()

    def resolve(book: Any) -> Path:
        assert book is caller and state.locked
        assert caller.api.Saved, "Resolve the committed Stage 1 handoff BEFORE dirtying Draft"
        assert not state.writes
        state.events.append("resolve")
        return state.folder

    def config(path: Path) -> Any:
        assert Path(path) == state.caller_path and state.locked
        assert not state.writes
        state.events.append("config")
        return state.config

    def metadata(folder: Path) -> dict[str, Any]:
        assert folder == state.folder and state.locked and not state.writes
        state.events.append("metadata")
        return state.inputs

    def preflight(settings: Any) -> None:
        assert settings is state.config and state.locked and not state.writes
        state.events.append("preflight")
        if state.preflight_error:
            raise state.preflight_error

    def confirm(inputs: dict[str, Any], folder: Path) -> bool:
        assert inputs is state.inputs and folder == state.folder and state.locked
        assert not state.writes
        state.clock = 200.0
        state.events.append("confirm")
        return state.approved

    def save() -> None:
        assert state.locked
        state.saves.append(deepcopy(state.fields))
        state.events.append("save")
        if len(state.saves) in state.fail_saves:
            raise BondsLiteError("Stage 2 audit save sentinel")
        caller.api.Saved = True

    def write(book: Any, fields: dict[str, Any]) -> None:
        assert book is caller and state.locked
        assert set(fields) <= {
            "last_run", "start", "finish", "elapsed", "user", "status",
            "input_run", "last_output",
        }, "Never write Stage 1's SST_Stage2_Input handoff"
        state.writes.append(deepcopy(fields))
        state.fields.update(fields)
        caller.api.Saved = False
        state.events.append(f"write:{fields.get('status', '')}")
        caller.save()

    transaction_factory = workflow.model_results_transaction

    @contextmanager
    def transaction(folder: Path, inputs: dict[str, Any]) -> Iterator[Any]:
        assert state.saves[-1]["status"] == "Running" and caller.api.Saved
        state.events.append("transaction")
        try:
            with transaction_factory(folder, inputs) as owned:
                def commit() -> None:
                    assert state.saves[-1]["status"] == "Succeeded" and caller.api.Saved
                    assert state.fields["last_output"] == state.relative
                    assert "validate" in state.events
                    state.events.append("commit")
                    return owned.commit()

                yield SimpleNamespace(commit=commit)
        finally:
            state.events.append("transaction_closed")

    def batch(folder: Path, settings: Any, application: Any = None) -> int:
        assert folder == state.folder and settings is state.config and state.locked
        assert application is caller.app
        running = state.saves[0]
        assert running["status"] == "Running" and running["input_run"] == state.relative
        assert running["last_output"] == state.previous
        assert not running.get("finish") and not running.get("elapsed")
        assert all(not list(item.output_folder.iterdir()) for item in state.inputs.values())
        assert not list(folder.glob("bonds_batch_summary*.csv"))
        state.events.append("batch")
        _write_stage_two_artifacts(state.inputs)
        _write_stage_two_summary(folder, _stage_two_summary_rows(state.inputs))
        state.clock = 206.5
        try:
            if state.batch_error:
                raise state.batch_error
            return 0
        finally:
            state.events.append("child_reaped")

    validate_results = workflow.validate_model_results

    def validate(folder: Path, inputs: dict[str, Any]) -> Path:
        assert folder == state.folder and inputs is state.inputs and state.locked
        assert "child_reaped" in state.events
        state.events.append("validate")
        if state.validation_error:
            raise state.validation_error
        return validate_results(folder, inputs)

    class AttemptDatetime(datetime):
        @classmethod
        def now(cls, tz: Any = None) -> datetime:
            return datetime(2026, 10, 28, 9, 0, 0 if state.clock <= 200 else 7)

    caller.save = save
    monkeypatch.setattr(getpass, "getuser", lambda: "stage-two-user")
    monkeypatch.setattr(workflow, "datetime", AttemptDatetime)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: state.clock)
    for name, replacement in {
        "workbook_run_guard": guard, "require_stage_two_panel": require,
        "_read_stage_two_fields": read, "resolve_stage_input": resolve,
        "load_model_config": config, "read_model_inputs": metadata,
        "preflight_model_environment": preflight, "confirm_model_overwrite": confirm,
        "write_stage_two_audit": write, "model_results_transaction": transaction,
        "run_model_batch": batch, "validate_model_results": validate,
    }.items():
        monkeypatch.setattr(workflow, name, replacement)
    monkeypatch.setattr(
        subprocess, "Popen", lambda *args, **kwargs: pytest.fail("Never launch the real model"),
    )
    return state


def test_run_stage_two_saves_verified_audit_before_commit_and_releases_one_guard(
    stage_two_attempt: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    assert workflow.run_stage_two(state.caller) == state.folder
    assert state.guard_calls == [state.caller_path]
    assert state.events.count("batch") == 1 and state.events.count("commit") == 1
    assert state.events.index("resolve") < state.events.index("write:Running")
    assert state.events.index("metadata") < state.events.index("write:Running")
    assert state.events.index("preflight") < state.events.index("batch")
    assert state.events.index("child_reaped") < state.events.index("validate")
    assert state.events.index("write:Succeeded") < state.events.index("commit")
    assert state.events[-2:] == ["transaction_closed", "release"]
    running, success = state.saves
    assert isinstance(running["last_run"], datetime)
    assert running["start"] == datetime(2026, 10, 28, 9)
    assert running["user"] == "stage-two-user"
    assert success["finish"] == datetime(2026, 10, 28, 9, 0, 7)
    assert success["elapsed"] == pytest.approx(6.5 / 86400)
    assert success["status"] == "Succeeded"
    assert success["input_run"] == success["last_output"] == state.relative
    assert state.legacy_handoff == state.relative and state.caller.api.Saved
    assert not list(state.folder.parent.glob(".*"))


@pytest.mark.parametrize("error", [
    BondsLiteError("Batch failure sentinel"), RuntimeError("Unexpected batch sentinel"),
    TimeoutError("Silent child deadline sentinel"), KeyboardInterrupt("Interrupt sentinel"),
])
@pytest.mark.parametrize("audit_failure", [False, True])
def test_run_stage_two_failure_restores_old_pair_and_preserves_original_exception(
    stage_two_attempt: Any, error: BaseException, audit_failure: bool,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    state.batch_error = error
    if audit_failure:
        state.fail_saves.add(2)
    with pytest.raises(type(error)) as caught:
        workflow.run_stage_two(state.caller)
    assert caught.value is error
    expected_status = (
        "Interrupted" if isinstance(error, KeyboardInterrupt)
        else "Timed Out" if isinstance(error, TimeoutError) else "Failed"
    )
    assert state.fields["status"] == expected_status
    assert state.fields["input_run"] == state.relative
    assert state.fields["last_output"] == state.previous
    assert state.fields["elapsed"] == pytest.approx(6.5 / 86400)
    assert "commit" not in state.events and "validate" not in state.events
    assert state.events.index("child_reaped") < state.events.index("transaction_closed")
    assert _stage_two_tree(state.folder) == state.before
    assert state.guard_calls == [state.caller_path] and not state.locked


def test_run_stage_two_validation_failure_rolls_back_before_reporting_success(
    stage_two_attempt: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    state.validation_error = BondsLiteError("MAP summary failed sentinel")
    with pytest.raises(BondsLiteError, match="MAP summary failed"):
        workflow.run_stage_two(state.caller)
    assert state.fields["status"] == "Failed"
    assert state.fields["last_output"] == state.previous
    assert _stage_two_tree(state.folder) == state.before
    assert "write:Succeeded" not in state.events and "commit" not in state.events


@pytest.mark.parametrize("recovery_save_fails", [False, True])
def test_run_stage_two_final_audit_failure_rolls_back_new_outputs_and_old_summary(
    stage_two_attempt: Any, recovery_save_fails: bool,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    state.fail_saves = {2, 3} if recovery_save_fails else {2}
    with pytest.raises(BondsLiteError, match=r"(?i)audit|sav"):
        workflow.run_stage_two(state.caller)
    assert state.fields["status"] == "Failed"
    assert state.fields["input_run"] == state.relative
    assert state.fields["last_output"] == state.previous
    assert _stage_two_tree(state.folder) == state.before
    assert "commit" not in state.events and not state.locked


def test_run_stage_two_initial_audit_failure_launches_no_child_or_transaction(
    stage_two_attempt: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    state.fail_saves.add(1)
    with pytest.raises(BondsLiteError, match="audit save sentinel"):
        workflow.run_stage_two(state.caller)
    assert "batch" not in state.events and "transaction" not in state.events
    assert _stage_two_tree(state.folder) == state.before
    assert not state.locked


def test_run_stage_two_declined_overwrite_changes_neither_audit_nor_results(
    stage_two_attempt: Any,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    state.approved = False
    before_fields = deepcopy(state.fields)
    assert workflow.run_stage_two(state.caller) is None
    assert not state.writes and not state.saves
    assert state.fields == before_fields and state.caller.api.Saved
    assert _stage_two_tree(state.folder) == state.before
    assert "batch" not in state.events and "transaction" not in state.events
    assert not state.locked


@pytest.mark.parametrize("boundary", [
    "require_stage_two_panel", "load_model_config", "resolve_stage_input",
    "read_model_inputs", "preflight_model_environment",
])
def test_run_stage_two_preflight_rejection_does_not_dirty_draft_or_touch_results(
    stage_two_attempt: Any, monkeypatch: pytest.MonkeyPatch, boundary: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt
    original = BondsLiteError(f"{boundary} rejected sentinel")

    def reject(*args: Any, **kwargs: Any) -> None:
        assert state.locked and not state.writes
        raise original

    monkeypatch.setattr(workflow, boundary, reject)
    with pytest.raises(BondsLiteError) as caught:
        workflow.run_stage_two(state.caller)
    assert caught.value is original
    assert not state.writes and not state.saves and state.caller.api.Saved
    assert "transaction" not in state.events and "batch" not in state.events
    assert _stage_two_tree(state.folder) == state.before
    assert not state.locked


def test_run_stage_two_duplicate_guard_prevents_preflight_and_audit(
    stage_two_attempt: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_attempt

    @contextmanager
    def reject(path: Path) -> Iterator[None]:
        assert path == state.caller_path
        raise BondsLiteError("Already in progress sentinel")
        yield

    monkeypatch.setattr(workflow, "workbook_run_guard", reject)
    with pytest.raises(BondsLiteError, match="Already in progress"):
        workflow.run_stage_two(state.caller)
    assert state.events == [] and not state.writes
    assert _stage_two_tree(state.folder) == state.before


@pytest.mark.parametrize("stage", [None, "1", "2"])
@pytest.mark.parametrize("cancelled", [False, True])
def test_main_stage_dispatch_keeps_default_one_and_treats_stage_two_cancellation_as_cancelled(
    monkeypatch: pytest.MonkeyPatch, console_logger: logging.Logger,
    reporting_threads: list[Any], capsys: pytest.CaptureFixture[str],
    stage: str | None, cancelled: bool,
) -> None:
    import bonds_lite_input as workflow

    caller = SimpleNamespace(fullname="Draft.xlsm", api=SimpleNamespace(FullName="Draft.xlsm"))
    app = SimpleNamespace(api=SimpleNamespace(Hwnd=1234), books=[caller])
    monkeypatch.setitem(sys.modules, "xlwings", SimpleNamespace(apps=[app]))
    arguments = ["bonds_lite_input", "--caller", "Draft.xlsm", "--excel-hwnd", "1234"]
    if stage is not None:
        arguments.extend(["--stage", stage])
    monkeypatch.setattr(sys, "argv", arguments)
    received = []

    def first(book: Any) -> Path:
        assert book is caller
        received.append(1)
        return Path("Bonds Lite output/20261028_090000")

    def second(book: Any) -> Path | None:
        assert book is caller
        received.append(2)
        return None if cancelled else Path("Bonds Lite output/20261028_090000")

    monkeypatch.setattr(workflow, "run_stage_one", first)
    monkeypatch.setattr(workflow, "run_stage_two", second, raising=False)
    monkeypatch.setattr("builtins.input", lambda *args: pytest.fail("No implicit pause"))
    assert workflow.main() == 0
    assert received == [2 if stage == "2" else 1]
    captured = capsys.readouterr()
    console = (captured.out + captured.err).lower()
    if stage == "2" and cancelled:
        assert "cancel" in console and "succeeded" not in console
        assert "saved bonds lite" not in console and "\nnone" not in console
    else:
        assert "20261028_090000" in console
    assert all(reporter.joined for reporter in reporting_threads)


@pytest.mark.parametrize("stage", ["0", "3", "two"])
def test_main_rejects_invalid_stage_before_connecting_to_excel(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], stage: str,
) -> None:
    import bonds_lite_input as workflow

    monkeypatch.setattr(sys, "argv", ["bonds_lite_input", "--stage", stage])
    with pytest.raises(SystemExit) as caught:
        workflow.main()
    assert caught.value.code == 2
    error = capsys.readouterr().err.lower()
    assert "--stage" in error and "invalid choice" in error


def test_stage_two_launcher_preserves_stage_one_entry_and_visible_console() -> None:
    """Allow a shared launcher while pinning both public VBA entry points."""
    import bonds_lite_input as workflow

    launcher = workflow.VBA_LAUNCHER
    assert "Sub GenerateBondsLiteInput" in launcher
    assert "Sub RunBondsLite" in launcher
    assert "--stage" in launcher and re.search(r"\b2\b", launcher)
    assert "--pause" in launcher and "-u -m bonds_lite_input" in launcher
    assert "shell.Run command, 1, False" in launcher


@pytest.fixture
def stage_two_controlled_runner(stage_two_results_case: Any, tmp_path: Path) -> Any:
    """Copy the tiny fixture, never refer to the real external runner directory."""
    import bonds_lite_input as workflow

    state = stage_two_results_case
    model = tmp_path / "controlled model with spaces \u00e9"
    model.mkdir()
    shutil.copy2(Path(__file__).parent / "fixtures" / "stage_two_runner.py",
                 model / "run_bonds_batch.py")
    state.config = workflow.ModelConfig(
        python_executable=Path(sys.executable), model_dir=model, timeout_seconds=30.0,
    )
    state.manifest = {
        "exit_code": 0, "rows": _stage_two_summary_rows(state.inputs),
        "output_folders": {
            dataset: str(item.output_folder) for dataset, item in state.inputs.items()
        },
        "artifacts": {
            dataset: list(item.required_files) for dataset, item in state.inputs.items()
        },
        "observed_path": str(state.folder / "runner_observed.json"),
    }
    return state


def _curve_freeze_workbook(curves: dict[str, list[list[object]]]) -> Any:
    ranges = {}
    wrapped_ranges = {}

    def make_range(name: str, values: list[list[object]]) -> Any:
        formula_state = SimpleNamespace(HasFormula=True)
        address = f"${len(wrapped_ranges) + 3}$6:${len(wrapped_ranges) + 3}${5 + len(values)}"

        class CurveRange:
            api = formula_state

            def options(self, **kwargs: Any) -> Any:
                assert kwargs == {"ndim": 2}
                return self

            @property
            def value(self) -> list[list[object]]:
                return values

            @value.setter
            def value(self, new_values: list[list[object]]) -> None:
                values[:] = new_values
                self.api.HasFormula = False

        wrapped = CurveRange()
        wrapped_ranges[name] = wrapped
        return SimpleNamespace(Address=address)

    for name, values in curves.items():
        ranges[name] = SimpleNamespace(DataBodyRange=make_range(name, values))
    table = SimpleNamespace(ListColumns=lambda name: ranges[name])
    sheet = SimpleNamespace(
        api=SimpleNamespace(ListObjects=lambda name: table),
        range=lambda address: next(
            value for value in wrapped_ranges.values() if value.api.HasFormula
        ),
        wrapped_ranges=wrapped_ranges,
    )
    book = SimpleNamespace(sheets={"P. Base Curves": sheet}, api=SimpleNamespace(Saved=True))
    book.save = lambda: None
    return book


def test_freeze_stage_two_local_curves_uses_validated_values_and_leaves_source_unchanged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    source = tmp_path / "Bonds Lite input_Bonds.xlsx"
    source.write_bytes(b"Original workbook sentinel")
    original = source.read_bytes()
    curves = {
        "GBP SONIA": [[0.01], [0.02]],
        "GBP GOV": [[0.03], [0.04]],
        "GBP RPI": [[0.05], [0.06]],
    }
    workbook = _curve_freeze_workbook(curves)

    @contextmanager
    def owned_book(
        application: Any, path: Path, read_only: bool = True,
    ) -> Iterator[Any]:
        assert path == source and read_only is False
        yield workbook

    monkeypatch.setattr(workflow, "owned_book", owned_book)
    workflow._freeze_stage_two_curve_values(object(), source)

    assert curves == {
        "GBP SONIA": [[0.01], [0.02]],
        "GBP GOV": [[0.03], [0.04]],
        "GBP RPI": [[0.05], [0.06]],
    }
    assert all(
        not curve.api.HasFormula
        for curve in workbook.sheets["P. Base Curves"].wrapped_ranges.values()
    )
    assert source.read_bytes() == original


@pytest.mark.parametrize("bad_value", [None, "#NAME?"])
def test_freeze_stage_two_local_curves_rejects_bad_cached_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, bad_value: object,
) -> None:
    import bonds_lite_input as workflow

    source = tmp_path / "Bonds Lite input_Bonds.xlsx"
    source.write_bytes(b"Temporary workbook sentinel")
    curves = {
        "GBP SONIA": [[0.01], [bad_value]],
        "GBP GOV": [[0.03], [0.04]],
        "GBP RPI": [[0.05], [0.06]],
    }
    workbook = _curve_freeze_workbook(curves)

    @contextmanager
    def owned_book(
        application: Any, path: Path, read_only: bool = True,
    ) -> Iterator[Any]:
        yield workbook

    monkeypatch.setattr(workflow, "owned_book", owned_book)
    with pytest.raises(BondsLiteError, match=r"(?i)GBP SONIA.*numeric"):
        workflow._freeze_stage_two_curve_values(object(), source)


@pytest.mark.parametrize("outcome", ["success", "nonzero", "failed_map", "missing_map"])
def test_run_model_batch_one_controlled_child_exact_arguments_merged_console_and_reaped(
    stage_two_controlled_runner: Any, monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str], caplog: pytest.LogCaptureFixture,
    outcome: str,
) -> None:
    import subprocess

    import bonds_lite_input as workflow

    state = stage_two_controlled_runner
    if outcome == "nonzero":
        state.manifest["exit_code"] = 7
    elif outcome == "failed_map":
        state.manifest["rows"][1].update(status="failed", error="MAP sentinel, controlled")
    elif outcome == "missing_map":
        state.manifest["rows"].pop()
    (state.folder / "runner_case.json").write_text(json.dumps(state.manifest), encoding="utf-8")
    original_popen = subprocess.Popen
    children, waits = [], []

    def launch(arguments: list[str], **kwargs: Any) -> Any:
        assert not children, "One subprocess for the complete two-input folder"
        assert arguments[:3] == [
            str(state.config.python_executable), "-u",
            str(state.config.model_dir / "run_bonds_batch.py"),
        ]
        batch_folder = Path(arguments[3])
        assert batch_folder != state.folder and batch_folder.is_dir()
        assert {path.name for path in batch_folder.glob("*.xlsx")} == set(
            EXPECTED_INPUT_FILENAMES.values()
        )
        assert all(
            (batch_folder / filename).read_bytes() ==
            (state.folder / filename).read_bytes()
            for filename in EXPECTED_INPUT_FILENAMES.values()
        )
        shutil.copy2(state.folder / "runner_case.json", batch_folder / "runner_case.json")
        state.batch_folder = batch_folder
        assert kwargs.get("shell", False) is False
        assert Path(kwargs["cwd"]) == state.config.model_dir
        assert kwargs["stdout"] == subprocess.PIPE
        assert kwargs["stderr"] == subprocess.STDOUT
        child = original_popen(arguments, **kwargs)
        children.append(child)
        wait = child.wait

        def reaped(*args: Any, **options: Any) -> int:
            result = wait(*args, **options)
            waits.append(result)
            return result

        monkeypatch.setattr(child, "wait", reaped)
        return child

    monkeypatch.setattr(subprocess, "Popen", launch)
    before = _pair_bytes(state.folder)
    try:
        with caplog.at_level(logging.INFO, logger="bonds_lite_input"):
            if outcome == "nonzero":
                with pytest.raises(BondsLiteError, match="7"):
                    workflow.run_model_batch(state.folder, state.config)
            else:
                assert workflow.run_model_batch(state.folder, state.config) == 0
        assert len(children) == 1 and children[0].poll() is not None
        assert waits, "Wait/reap the owned child before returning or raising"
        captured = capsys.readouterr()
        console = captured.out + captured.err + caplog.text
        for marker in ("STAGE_TWO_STDOUT", "STAGE_TWO_STDERR", "STAGE_TWO_PARTIAL"):
            assert marker in console
        observed = json.loads((state.folder / "runner_observed.json").read_text(encoding="utf-8"))
        assert observed["argv"] == [str(state.batch_folder)]
        assert Path(observed["cwd"]) == state.config.model_dir
        assert observed["pid"] == children[0].pid
        if outcome in ("failed_map", "missing_map"):
            with pytest.raises(BondsLiteError):
                workflow.validate_model_results(state.folder, state.inputs)
        elif outcome == "success":
            assert workflow.validate_model_results(state.folder, state.inputs).is_file()
        assert not state.batch_folder.exists(), "Local input staging must be cleaned after the child exits"
        assert _pair_bytes(state.folder) == before
    finally:
        for child in children:
            if child.poll() is None:
                child.kill()
                child.wait(timeout=5.0)
            if child.stdout is not None:
                child.stdout.close()
        (state.folder / "runner_observed.json").unlink(missing_ok=True)


@pytest.mark.parametrize("mode", ["silent", "partial_line", "interrupt"])
def test_run_model_batch_deadline_and_interrupt_reap_only_owned_child_without_sleep(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch, mode: str,
) -> None:
    """A blocked reader cannot own the deadline; no real PID is present to kill."""
    import subprocess

    import bonds_lite_input as workflow

    paths = stage_two_config_paths
    config = workflow.ModelConfig(
        python_executable=paths.executable, model_dir=paths.model, timeout_seconds=10.0,
    )
    released = threading.Event()
    readers: list[threading.Thread] = []
    events: list[str] = []
    clock = SimpleNamespace(now=0.0)

    class BlockedOutput:
        """Block only an owned reader thread until the fake child is stopped."""

        closed = False
        first = True

        def readline(self, *args: Any) -> str:
            assert threading.current_thread() is not threading.main_thread(), (
                "Do not block the deadline controller in stdout.readline"
            )
            if threading.current_thread() not in readers:
                readers.append(threading.current_thread())
            if self.first and mode == "partial_line":
                self.first = False
                return "No trailing newline"
            assert released.wait(2.0), "Stop the child before joining its blocked reader"
            return ""

        def __iter__(self) -> Any:
            return self

        def __next__(self) -> str:
            value = self.readline()
            if not value:
                raise StopIteration
            return value

        def close(self) -> None:
            self.closed = True
            released.set()

    class OwnedChild:
        """Expose process lifecycle only, with no user or system PID to terminate."""

        pid = None
        returncode = None
        stdout = BlockedOutput()
        interrupted = False

        def poll(self) -> int | None:
            clock.now += 20.0
            if mode == "interrupt" and not self.interrupted:
                self.interrupted = True
                raise KeyboardInterrupt("Controlled child interruption")
            return self.returncode

        def wait(self, timeout: float | None = None) -> int:
            if self.returncode is None:
                clock.now += 20.0
                if mode == "interrupt" and not self.interrupted:
                    self.interrupted = True
                    raise KeyboardInterrupt("Controlled child interruption")
                raise subprocess.TimeoutExpired("controlled fixture", timeout)
            events.append("reaped")
            return self.returncode

        def terminate(self) -> None:
            events.append("terminate")
            self.returncode = -1
            released.set()

        def kill(self) -> None:
            events.append("kill")
            self.returncode = -1
            released.set()

        def __enter__(self) -> Any:
            return self

        def __exit__(self, *args: Any) -> None:
            self.wait()

    child = OwnedChild()
    calls = []

    def launch(*args: Any, **kwargs: Any) -> Any:
        calls.append(args)
        assert len(calls) == 1
        return child

    monkeypatch.setattr(workflow.time, "monotonic", lambda: clock.now)
    monkeypatch.setattr(workflow.time, "perf_counter", lambda: clock.now)
    monkeypatch.setattr(subprocess, "Popen", launch)
    monkeypatch.setattr(workflow.time, "sleep", lambda *args: pytest.fail("No sleep-based timeout"))
    monkeypatch.setattr(os, "kill", lambda *args: pytest.fail("Never signal a user PID"))
    monkeypatch.setattr(subprocess, "run", lambda *args, **kwargs: pytest.fail("No global taskkill"))
    expected_error = KeyboardInterrupt if mode == "interrupt" else TimeoutError
    try:
        with pytest.raises(expected_error):
            workflow.run_model_batch(paths.caller.parent, config)
        assert len(calls) == 1
        assert child.returncode is not None and events[-1] == "reaped"
        assert set(events) <= {"terminate", "kill", "reaped"}
        assert all(not reader.is_alive() for reader in readers)
    finally:
        released.set()
        for reader in readers:
            reader.join(timeout=2.0)


def test_run_model_batch_streams_progress_before_child_can_finish(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch,
    console_logger: logging.Logger,
) -> None:
    """The controlled child cannot finish until its first message is visible."""
    import io
    import subprocess

    import bonds_lite_input as workflow

    paths = stage_two_config_paths
    seen = threading.Event()
    waits = []

    class ProgressHandler(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            if "STREAM_BEFORE_EXIT" in record.getMessage():
                seen.set()

    class VisibleOutput:
        def __init__(self, wrapped: Any) -> None:
            self.wrapped = wrapped

        def write(self, text: str) -> int:
            if "STREAM_BEFORE_EXIT" in text:
                seen.set()
            return self.wrapped.write(text)

        def __getattr__(self, name: str) -> Any:
            return getattr(self.wrapped, name)

    class Child:
        pid = None
        stdout = io.StringIO("STREAM_BEFORE_EXIT\n")
        returncode = None

        def poll(self) -> int | None:
            if seen.is_set():
                self.returncode = 0
            return self.returncode

        def wait(self, timeout: float | None = None) -> int:
            assert seen.wait(2.0), "Progress was buffered until after child completion"
            self.returncode = 0
            waits.append(True)
            return 0

        def __enter__(self) -> Any:
            return self

        def __exit__(self, *args: Any) -> None:
            self.wait()

    child = Child()
    handler = ProgressHandler()
    console_logger.addHandler(handler)
    monkeypatch.setattr(sys, "stdout", VisibleOutput(sys.stdout))
    monkeypatch.setattr(subprocess, "Popen", lambda *args, **kwargs: child)
    try:
        config = workflow.ModelConfig(
            python_executable=paths.executable, model_dir=paths.model, timeout_seconds=5.0,
        )
        assert workflow.run_model_batch(paths.caller.parent, config) == 0
        assert seen.is_set() and waits
    finally:
        seen.set()
        child.stdout.close()
        console_logger.removeHandler(handler)


def test_preflight_model_environment_uses_bounded_separate_import_probe_not_model(
    stage_two_config_paths: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import subprocess

    import bonds_lite_input as workflow

    paths = stage_two_config_paths
    config = workflow.ModelConfig(python_executable=paths.executable, model_dir=paths.model)
    probes = []

    def probe(arguments: list[str], **kwargs: Any) -> Any:
        probes.append(arguments)
        assert arguments[0] == str(paths.executable)
        assert "-c" in arguments
        assert not any("run_bonds_batch.py" in argument for argument in arguments)
        assert kwargs.get("shell", False) is False
        timeout = kwargs.get("timeout")
        assert isinstance(timeout, (int, float)) and 0 < timeout < float("inf")
        return subprocess.CompletedProcess(arguments, 1, stdout="", stderr="Missing dependency sentinel")

    monkeypatch.setattr(subprocess, "run", probe)
    monkeypatch.setattr(subprocess, "Popen", lambda *args, **kwargs: pytest.fail("Do not launch the batch as a probe"))
    with pytest.raises(BondsLiteError, match=r"(?i)dependency|environment|import|preflight"):
        workflow.preflight_model_environment(config)
    assert len(probes) == 1


def test_read_model_inputs_follows_moved_names_instead_of_cell_addresses(
    stage_two_saved_inputs: Any,
) -> None:
    import openpyxl

    import bonds_lite_input as workflow

    state = stage_two_saved_inputs
    for filename in EXPECTED_INPUT_FILENAMES.values():
        path = state.folder / filename
        book = openpyxl.load_workbook(path)
        try:
            sheet = book["I. Control"]
            for offset, (name, old_address) in enumerate(state.addresses.items(), 40):
                sheet.cell(offset, 13, sheet[old_address].value)
                sheet[old_address] = "Fixed-address trap"
                book.defined_names[name].attr_text = f"'I. Control'!$M${offset}"
            book.save(path)
        finally:
            book.close()
    before = _pair_bytes(state.folder)
    inputs = workflow.read_model_inputs(state.folder)
    assert {dataset: item.output_folder for dataset, item in inputs.items()} == {
        dataset: state.folder / f"JRL {dataset}" for dataset in EXPECTED_INPUT_FILENAMES
    }
    assert _pair_bytes(state.folder) == before


@pytest.mark.parametrize("name", ["entity_name", "output_suffix", "bool_inv_exp"])
def test_read_model_inputs_rejects_nonlocal_or_multicell_scalar_names(
    stage_two_saved_inputs: Any, name: str,
) -> None:
    import openpyxl

    import bonds_lite_input as workflow

    state = stage_two_saved_inputs
    path = state.folder / EXPECTED_INPUT_FILENAMES["MAP"]
    book = openpyxl.load_workbook(path)
    try:
        book.defined_names[name].attr_text = "'I. Control'!$D$13:$D$14"
        book.save(path)
    finally:
        book.close()
    before = _pair_bytes(state.folder)
    with pytest.raises(BondsLiteError, match=rf"(?i){name}|single|cell|name"):
        workflow.read_model_inputs(state.folder)
    assert _pair_bytes(state.folder) == before


@pytest.mark.parametrize("new_outputs", ["none", "summary_only", "input_copy_only"])
def test_validate_model_results_cannot_reuse_old_success_staged_by_transaction(
    stage_two_results_case: Any, new_outputs: str,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    _write_stage_two_artifacts(state.inputs)
    _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    before = _stage_two_tree(state.folder)
    with workflow.model_results_transaction(state.folder, state.inputs):
        if new_outputs == "summary_only":
            _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
        elif new_outputs == "input_copy_only":
            for item in state.inputs.values():
                shutil.copy2(item.input_path, item.output_folder / item.required_files[-1])
        with pytest.raises(BondsLiteError):
            workflow.validate_model_results(state.folder, state.inputs)
    assert _stage_two_tree(state.folder) == before


def test_model_results_transaction_backup_preparation_failure_restores_partial_moves(
    stage_two_results_case: Any, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import bonds_lite_input as workflow

    state = stage_two_results_case
    for dataset, item in state.inputs.items():
        (item.output_folder / "old.csv").write_bytes(dataset.encode())
    _write_stage_two_summary(state.folder, _stage_two_summary_rows(state.inputs))
    before = _stage_two_tree(state.folder)
    siblings = set(state.folder.parent.iterdir())
    transfer_count = []
    original_rename, original_replace, original_move = Path.rename, Path.replace, shutil.move

    def reject_second_move(source: Any, target: Any) -> None:
        if Path(source).is_relative_to(state.folder) and not Path(target).is_relative_to(state.folder):
            transfer_count.append(Path(source))
            if len(transfer_count) == 2:
                raise PermissionError("Backup preparation denied sentinel")

    def rename(source: Path, target: Path) -> Path:
        reject_second_move(source, target)
        return original_rename(source, target)

    def replace(source: Path, target: Path) -> Path:
        reject_second_move(source, target)
        return original_replace(source, target)

    def move(source: Any, target: Any, *args: Any, **kwargs: Any) -> Any:
        reject_second_move(source, target)
        return original_move(source, target, *args, **kwargs)

    monkeypatch.setattr(Path, "rename", rename)
    monkeypatch.setattr(Path, "replace", replace)
    monkeypatch.setattr(shutil, "move", move)
    with pytest.raises((BondsLiteError, PermissionError), match=r"(?i)backup|preparation|denied"):
        with workflow.model_results_transaction(state.folder, state.inputs):
            pytest.fail("Backup failure must prevent model launch")
    assert len(transfer_count) >= 2
    assert _stage_two_tree(state.folder) == before
    assert set(state.folder.parent.iterdir()) == siblings


@pytest.mark.parametrize("outcome", ["success", "cancelled", "failed", "timeout", "interrupted"])
def test_main_stage_two_cleans_up_and_releases_guard_before_pause(
    monkeypatch: pytest.MonkeyPatch, console_logger: logging.Logger,
    reporting_threads: list[Any], capsys: pytest.CaptureFixture[str], outcome: str,
) -> None:
    import bonds_lite_input as workflow

    caller = SimpleNamespace(fullname="Draft.xlsm", api=SimpleNamespace(FullName="Draft.xlsm"))
    app = SimpleNamespace(api=SimpleNamespace(Hwnd=1234), books=[caller])
    monkeypatch.setitem(sys.modules, "xlwings", SimpleNamespace(apps=[app]))
    monkeypatch.setattr(sys, "argv", [
        "bonds_lite_input", "--caller", "Draft.xlsm", "--excel-hwnd", "1234",
        "--stage", "2", "--pause",
    ])
    events = []
    errors = {
        "failed": BondsLiteError("Stage 2 failure sentinel"),
        "timeout": TimeoutError("Stage 2 timeout sentinel"),
        "interrupted": KeyboardInterrupt("Stage 2 interrupted sentinel"),
    }

    def second(book: Any) -> Path | None:
        assert book is caller
        events.append("lock")
        try:
            if outcome in errors:
                raise errors[outcome]
            return None if outcome == "cancelled" else Path("Bonds Lite output/20261028_090000")
        finally:
            events.extend(["child_reaped", "transaction_closed", "release"])

    def pause(prompt: str = "") -> str:
        assert events == ["lock", "child_reaped", "transaction_closed", "release"]
        assert reporting_threads and all(thread.joined for thread in reporting_threads)
        captured = capsys.readouterr()
        console = (captured.out + captured.err).lower()
        assert "total" in console
        if outcome == "cancelled":
            assert "cancel" in console and "succeeded" not in console
        events.append("pause")
        return ""

    monkeypatch.setattr(workflow, "run_stage_two", second, raising=False)
    monkeypatch.setattr(workflow, "run_stage_one", lambda *args: pytest.fail("Wrong stage"))
    monkeypatch.setattr(workflow, "notify_user", lambda *args, **kwargs: pytest.fail("No modal popup"))
    monkeypatch.setattr("builtins.input", pause)
    assert workflow.main() == (1 if outcome in errors else 0)
    assert events[-1] == "pause"


@pytest.mark.excel
@pytest.mark.parametrize("resize", [False, True])
def test_excel_stage_two_panel_preserves_moved_stage_one_and_survives_move_resize_reopen(
    stage_panel_excel_book: Any, tmp_path: Path, resize: bool,
) -> None:
    """Use the shared opt-in blank book; never install into the original Draft."""
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    sheet = caller.sheets["I. Control"]
    stage_one = caller.names["SST_Stage1_Panel"].refers_to_range
    stage_one.api.Cut(Destination=sheet.range("Q100").api)
    button = sheet.api.Buttons().Add(20, sheet.range("D120").top, 220, 48)
    button.Name = "btnGenerateBondsLiteInput"
    button.Caption = "Generate Bonds Lite Input"
    geometry = (button.Left, button.Top, button.Width, button.Height)
    legacy = caller.names["SST_Stage2_Input"].refers_to_range
    legacy.value = "Bonds Lite output/20261027_090000"
    stage_one = caller.names["SST_Stage1_Panel"].refers_to_range
    before = _stage_panel_contents(stage_one)
    stage_one_names = {
        name: caller.names[name].refers_to
        for name in ("SST_Stage1_Panel", *workflow.STAGE_ONE_FIELD_NAMES.values())
    }
    assert workflow.initialise_stage_two_panel(caller) is None
    panel = caller.names["SST_Stage2_Panel"].refers_to_range
    assert panel.top >= max(stage_one.top + stage_one.height, button.Top + button.Height)
    assert workflow.require_stage_two_panel(caller) is None
    assert _stage_panel_contents(stage_one) == before
    assert (button.Left, button.Top, button.Width, button.Height) == geometry
    old_address = panel.address
    panel.api.Cut(Destination=sheet.range("AB160").api)
    sheet.range(old_address).value = "Old Stage 2 position sentinel"
    old_area = _stage_panel_contents(sheet.range(old_address))
    if resize:
        sheet.range("163:163").api.Insert()
        sheet.range("AF:AF").api.Insert()
    output = _complete_run(tmp_path / "Bonds Lite output" / "20261028_090000")
    relative = output.relative_to(tmp_path).as_posix()
    fields = {
        "last_run": datetime(2026, 10, 28), "start": datetime(2026, 10, 28, 9),
        "finish": datetime(2026, 10, 28, 9, 0, 7), "elapsed": 7 / 86400,
        "user": "movable-stage-two", "status": "Succeeded",
        "input_run": relative, "last_output": relative,
    }
    assert workflow.write_stage_two_audit(caller, fields) is None
    read_back = workflow._read_stage_two_fields(caller)
    for field, value in fields.items():
        assert read_back[field] == (pytest.approx(value) if field == "elapsed" else value)
    destination = caller.names[workflow.STAGE_TWO_FIELD_NAMES["last_output"]].refers_to_range
    assert destination.api.Hyperlinks.Item(1).Address == relative
    assert workflow._read_stage_fields(caller)["stage_two_input"] == "Bonds Lite output/20261027_090000"
    names = {name.name: name.refers_to for name in caller.names}
    installed = _stage_panel_contents(caller.names["SST_Stage2_Panel"].refers_to_range)
    assert workflow.initialise_stage_two_panel(caller) is None
    assert {name.name: name.refers_to for name in caller.names} == names
    assert _stage_panel_contents(caller.names["SST_Stage2_Panel"].refers_to_range) == installed
    assert _stage_panel_contents(sheet.range(old_address)) == old_area
    assert {name: caller.names[name].refers_to for name in stage_one_names} == stage_one_names
    caller.save()
    application, path = caller.app, caller.fullname
    caller.close()
    reopened = application.books.open(path, update_links=False)
    assert workflow.require_stage_two_panel(reopened) is None
    assert workflow._read_stage_two_fields(reopened)["last_output"] == relative
    assert workflow.initialise_stage_two_panel(reopened) is None
    assert reopened.api.Saved
    assert {name.name: name.refers_to for name in reopened.names} == names


@pytest.mark.excel
@pytest.mark.parametrize("defect", ["missing", "outside", "overlap", "broken", "protected"])
def test_excel_stage_two_panel_rejects_invalid_names_before_audit_write_or_save(
    stage_panel_excel_book: Any, defect: str,
) -> None:
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    workflow.initialise_stage_two_panel(caller)
    sheet = caller.sheets["I. Control"]
    name = workflow.STAGE_TWO_FIELD_NAMES["user"]
    if defect == "missing":
        caller.names[name].delete()
    elif defect == "outside":
        caller.names[name].refers_to = "='I. Control'!$A$1"
    elif defect == "overlap":
        caller.names[name].refers_to = caller.names[
            workflow.STAGE_TWO_FIELD_NAMES["status"]
        ].refers_to
    elif defect == "broken":
        caller.names[name].refers_to = "=#REF!"
    else:
        sheet.api.Protect()
    panel = caller.names["SST_Stage2_Panel"].refers_to_range
    before, saved_before = _stage_panel_contents(panel), caller.api.Saved
    names_before = {name.name: name.refers_to for name in caller.names}
    saves = []
    boundary = _ExcelForwarder(caller, save=lambda: saves.append(True))
    for action in (
        lambda: workflow.require_stage_two_panel(boundary),
        lambda: workflow.write_stage_two_audit(boundary, {"status": "Must not write"}),
        lambda: workflow.initialise_stage_two_panel(boundary),
    ):
        with pytest.raises(BondsLiteError):
            action()
        assert _stage_panel_contents(panel) == before
        assert {name.name: name.refers_to for name in caller.names} == names_before
        assert caller.api.Saved == saved_before
    assert saves == []


@pytest.mark.excel
def test_excel_stage_two_installer_preserves_moved_controls_and_one_button_per_stage(
    stage_panel_excel_book: Any, tmp_path: Path,
) -> None:
    """Install twice into a new disposable XLSM, never a copy of the user's Draft."""
    import bonds_lite_input as workflow

    caller = stage_panel_excel_book
    application = caller.app
    sheet = caller.sheets["I. Control"]
    caller.names["SST_Stage1_Panel"].refers_to_range.api.Cut(Destination=sheet.range("Q100").api)
    sheet.range("D39").value = "Custom Bonds.xlsx"
    sheet.range("D40").value = "Custom MAP.xlsx"
    sheet.range("D42").value = "Custom template.xlsx"
    button = sheet.api.Buttons().Add(25, sheet.range("D120").top, 240, 52)
    button.Name = "btnGenerateBondsLiteInput"
    button.Caption = "Generate Bonds Lite Input"
    original_geometry = (button.Left, button.Top, button.Width, button.Height)
    stage_one_names = {
        name: caller.names[name].refers_to
        for name in ("SST_Stage1_Panel", *workflow.STAGE_ONE_FIELD_NAMES.values())
    }
    path = tmp_path / "Disposable installer caller.xlsm"
    caller.api.SaveAs(str(path), FileFormat=52)
    caller.close()
    previous_names = None
    previous_geometry = None
    for _attempt in range(2):
        before = path.read_bytes()
        backup = workflow.install_button(path)
        assert backup.read_bytes() == before
        book = application.books.open(str(path), update_links=False)
        try:
            sheet = book.sheets["I. Control"]
            buttons = sheet.api.Buttons()
            installed = [buttons.Item(index) for index in range(1, buttons.Count + 1)]
            assert [item.Caption for item in installed].count("Generate Bonds Lite Input") == 1
            assert [item.Caption for item in installed].count("Run Bonds Lite") == 1
            first = next(item for item in installed if item.Caption == "Generate Bonds Lite Input")
            second = next(item for item in installed if item.Caption == "Run Bonds Lite")
            assert (first.Left, first.Top, first.Width, first.Height) == original_geometry
            assert "GenerateBondsLiteInput" in first.OnAction
            assert "RunBondsLite" in second.OnAction
            assert {name: book.names[name].refers_to for name in stage_one_names} == stage_one_names
            assert {address: sheet.range(address).value for address in ("D39", "D40", "D42")} == {
                "D39": "Custom Bonds.xlsx", "D40": "Custom MAP.xlsx",
                "D42": "Custom template.xlsx",
            }
            names = {name.name: name.refers_to for name in book.names}
            geometry = (second.Left, second.Top, second.Width, second.Height)
            if previous_names is not None:
                assert names == previous_names and geometry == previous_geometry
            previous_names, previous_geometry = names, geometry
            assert workflow.require_stage_two_panel(book) is None
            module = book.api.VBProject.VBComponents.Item("modBondsLiteLauncher").CodeModule
            assert "Sub RunBondsLite" in module.Lines(1, module.CountOfLines)
        finally:
            book.close()
