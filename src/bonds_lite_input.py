"""Prepare Bonds Lite input copies without modifying their source workbooks."""

from __future__ import annotations

import argparse
import csv
import ctypes
import getpass
import hashlib
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, field
from datetime import datetime
from numbers import Real
from pathlib import Path
from typing import TYPE_CHECKING, Iterator
from xml.etree import ElementTree as ET
from zipfile import ZipFile

if TYPE_CHECKING:
    import xlwings as xw

logger = logging.getLogger(__name__)


@dataclass
class _ProgressState:
    """Plain Python timing state shared with the console reporter."""

    started: float
    stages: list[tuple[str, float]] = field(default_factory=list)
    lock: threading.Lock = field(default_factory=threading.Lock)


_progress_state: ContextVar[_ProgressState | None] = ContextVar(
    "bonds_lite_progress", default=None
)


@dataclass(frozen=True)
class StageOneInputs:
    """One validated pair of data paths and shared controls for a complete run."""

    caller_path: Path
    bonds_path: Path
    map_path: Path
    template_path: Path
    valuation_date: datetime
    entity: str
    stresses: dict[str, list[list[float]]]


@dataclass(frozen=True)
class ModelConfig:
    """Separate model interpreter, project directory and complete-batch deadline."""

    python_executable: Path
    model_dir: Path
    timeout_seconds: float = 3600.0


@dataclass(frozen=True)
class ModelInput:
    """Read-only execution contract captured from one saved Stage 1 workbook."""

    dataset: str
    input_path: Path
    output_folder: Path
    output_suffix: str
    required_files: tuple[str, ...]
    input_hash: str


def _user_environment_value(name: str) -> str:
    """Read a per-user Windows environment setting when the process lacks it."""
    if os.name != "nt":
        return ""
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as key:
            value, _ = winreg.QueryValueEx(key, name)
    except OSError:
        return ""
    return os.path.expandvars(value).strip() if isinstance(value, str) else ""


def load_model_config(caller_path: Path) -> ModelConfig:
    """Resolve per-user model settings without borrowing the controller Python."""
    setting = os.environ.get("BONDS_LITE_MODEL_PYTHON", "").strip()
    if not setting:
        setting = _user_environment_value("BONDS_LITE_MODEL_PYTHON")
    if not setting:
        raise BondsLiteError("Set BONDS_LITE_MODEL_PYTHON to the model environment's python.exe")
    executable = Path(os.path.expandvars(setting)).expanduser()
    if executable.is_dir():
        executable = executable / "python.exe"
    if executable.name.casefold() != "python.exe" or not executable.is_file():
        raise BondsLiteError(f"Model Python interpreter does not exist: {executable}")
    default = (
        Path(caller_path).resolve().parent
        / "../../02 Sensitivities/2025/05 ORSA/Models/Bonds Lite Batch/bonds-hm-cst"
    )
    model_setting = os.environ.get("BONDS_LITE_MODEL_DIR", "").strip()
    if not model_setting:
        model_setting = _user_environment_value("BONDS_LITE_MODEL_DIR")
    model_dir = Path(os.path.expandvars(model_setting or str(default))).expanduser().resolve()
    if not model_dir.is_dir() or not (model_dir / "run_bonds_batch.py").is_file():
        raise BondsLiteError(f"Model directory has no run_bonds_batch.py: {model_dir}")
    return ModelConfig(executable.resolve(), model_dir)


def preflight_model_environment(config: ModelConfig) -> None:
    """Probe numerical imports in the separate environment without running a model."""
    probe = (
        "import sys,numpy,pandas,xlwings,scipy,dateutil; "
        "assert sys.version_info[:2] == (3,11), 'Model Python 3.11 is required'; "
        "print('Model environment ready')"
    )
    try:
        result = subprocess.run(
            [str(config.python_executable), "-c", probe],
            cwd=str(config.model_dir), capture_output=True, text=True,
            timeout=30, check=True, shell=False, env=_model_environment(config),
        )
    except (OSError, subprocess.SubprocessError) as error:
        detail = getattr(error, "stderr", "") or str(error)
        raise BondsLiteError(f"Model environment preflight failed: {detail}") from error
    if result.returncode != 0:
        raise BondsLiteError(f"Model environment preflight failed: {result.stderr}")
    logger.info(result.stdout.strip())


def _file_hash(path: Path) -> str:
    """Hash saved inputs without holding their entire contents in memory."""
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _safe_model_filename(value: str) -> str:
    """Reject configured output names that could escape their result directory."""
    if (
        not isinstance(value, str) or not value or value != value.strip()
        or value in (".", "..")
        or re.search(r'[<>:"/\\|?*\x00-\x1f\x7f]', value)
    ):
        raise BondsLiteError(f"Unsafe model output filename or suffix: {value!r}")
    return value


def read_model_inputs(run_folder: Path) -> dict[str, ModelInput]:
    """Read saved names and filename mappings without Excel, mutation or model import."""
    import openpyxl

    run_folder = run_folder.resolve()
    expected = set(INPUT_FILENAMES.values())
    actual = {path.name for path in run_folder.glob("*.xlsx") if not path.name.startswith("~$")}
    if actual != expected:
        raise BondsLiteError("Stage 2 requires exactly the two expected Bonds/MAP input workbooks")
    result = {}
    for dataset, filename in INPUT_FILENAMES.items():
        path = run_folder / filename
        if path.is_symlink() or not path.is_file() or path.stat().st_size == 0:
            raise BondsLiteError(f"Invalid {dataset} input workbook: {path}")
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                values = {}
                for name in (
                    "entity_name", "output_folder_path", "output_suffix",
                    "bool_inv_exp", "risk_factor_path",
                ):
                    definition = workbook.defined_names[name]
                    destinations = list(definition.destinations)
                    if len(destinations) != 1 or destinations[0][0] != "I. Control":
                        raise BondsLiteError(f"Model control name must be local to I. Control: {name}")
                    left, top, right, bottom = openpyxl.utils.range_boundaries(destinations[0][1])
                    if (left, top) != (right, bottom):
                        raise BondsLiteError(f"Model control name must identify one cell: {name}")
                    cell = workbook["I. Control"].cell(top, left)
                    if cell.data_type == "f":
                        raise BondsLiteError(f"Model control must be a saved scalar value: {name}")
                    values[name] = cell.value
                entity = validate_entity(values["entity_name"])
                folder = Path(values["output_folder_path"]).resolve()
                if folder != run_folder / f"{entity} {dataset}" or not folder.is_dir() or folder.is_symlink():
                    raise BondsLiteError(f"{dataset} output folder does not match this run: {folder}")
                suffix = values["output_suffix"]
                suffix = "" if suffix is None else suffix
                if suffix:
                    _safe_model_filename(suffix)
                if not isinstance(values["bool_inv_exp"], bool):
                    raise BondsLiteError("bool_inv_exp must be a saved Boolean value")
                scenario_folder = Path(os.path.expandvars(str(values["risk_factor_path"]))).expanduser()
                if not scenario_folder.is_absolute():
                    scenario_folder = Path.cwd() / scenario_folder
                if not scenario_folder.is_dir():
                    raise BondsLiteError(f"Model scenario folder is inaccessible: {scenario_folder}")
                with ZipFile(path) as archive:
                    tables = [ET.fromstring(archive.read(item)) for item in archive.namelist()
                              if item.startswith("xl/tables/") and item.endswith(".xml")]
                base_curve_table = next((
                    item for item in tables if item.get("name") == "P_tbl_BaseCurves"
                ), None)
                if base_curve_table is None:
                    raise BondsLiteError("Input has no P_tbl_BaseCurves table")
                base_curve_ref = base_curve_table.get("ref")
                min_col, min_row, _, _ = openpyxl.utils.range_boundaries(base_curve_ref)
                curve_rows = workbook["P. Base Curves"][base_curve_ref]
                curve_headers = [cell.value for cell in curve_rows[0]]
                if "GBP SONIA" not in curve_headers:
                    raise BondsLiteError("P_tbl_BaseCurves has no GBP SONIA column")
                curve_index = curve_headers.index("GBP SONIA")
                curve_values_book = openpyxl.load_workbook(
                    path, read_only=True, data_only=True
                )
                try:
                    cached_curve_rows = curve_values_book["P. Base Curves"][base_curve_ref]
                    curve_values = [row[curve_index].value for row in cached_curve_rows[1:]]
                finally:
                    curve_values_book.close()
                invalid_curve_rows = [
                    min_row + offset + 1
                    for offset, value in enumerate(curve_values)
                    if isinstance(value, bool)
                    or not isinstance(value, Real)
                    or not math.isfinite(value)
                ]
                if not curve_values or invalid_curve_rows:
                    raise BondsLiteError(
                        "P. Base Curves[GBP SONIA] must contain finite numeric cached "
                        f"values; invalid rows: {invalid_curve_rows[:20]}"
                    )
                table = next((item for item in tables if item.get("name") == "P_tbl_OutputFilenames"), None)
                if table is None:
                    raise BondsLiteError("Input has no P_tbl_OutputFilenames table")
                header_row, *body = workbook["P. Parameter"][table.get("ref")]
                headers = [cell.value for cell in header_row]
                key_index, value_index = headers.index("OutputType"), headers.index("Filename")
                mapping = {
                    row[key_index].value: _safe_model_filename(row[value_index].value)
                    for row in body if row[key_index].value is not None
                }
                scenario_names = {"GBP_RPI"}
                for table_name in ("P_tbl_OtherCurves", "P_tbl_SpreadCurves"):
                    scenario_table = next((
                        item for item in tables if item.get("name") == table_name
                    ), None)
                    if scenario_table is None:
                        raise BondsLiteError(f"Input has no {table_name} scenario table")
                    table_rows = workbook["P. Parameter"][scenario_table.get("ref")]
                    scenario_headers = [cell.value for cell in table_rows[0]]
                    if "File Name" not in scenario_headers:
                        raise BondsLiteError(f"{table_name} has no File Name column")
                    filename_index = scenario_headers.index("File Name")
                    for row in table_rows[1:]:
                        stub = row[filename_index].value
                        if stub not in (None, ""):
                            scenario_names.add(_safe_model_filename(stub))
                missing_scenarios = [
                    name for name in sorted(scenario_names)
                    if not (scenario_folder / f"{name}.csv").is_file()
                ]
                if missing_scenarios:
                    raise BondsLiteError(
                        f"Model scenario files are missing from {scenario_folder}: "
                        + ", ".join(missing_scenarios)
                    )
                def output_name(key: str) -> str:
                    base = mapping.get(key, f"py_bonds_{key}")
                    return _safe_model_filename(f"{base}_{suffix}.csv" if suffix else f"{base}.csv")

                base = mapping.get("BaseScenario", "py_bonds_BaseScenario")
                nozeros = base.rstrip("_") + "_nozeros" + ("_" if base.endswith("_") else "")
                required = (
                    *(output_name(key) for key in ("ISINs", "Grouped", "BaseScenario")),
                    _safe_model_filename(f"{nozeros}_{suffix}.csv" if suffix else f"{nozeros}.csv"),
                    output_name("InvExp_Base"),
                    _safe_model_filename(f"py_bonds_input_file_{suffix}.xlsx"),
                )
                if len(set(required)) != len(required):
                    raise BondsLiteError("Configured model result filenames collide")
            finally:
                workbook.close()
        except BondsLiteError:
            raise
        except Exception as error:
            raise BondsLiteError(f"Cannot read safe model metadata from {path}: {error}") from error
        result[dataset] = ModelInput(dataset, path, folder, suffix, required, _file_hash(path))
    return result


def validate_model_results(run_folder: Path, inputs: dict[str, ModelInput]) -> Path:
    """Require both successful summary rows, fresh artifacts and unchanged inputs."""
    summaries = list(run_folder.glob("bonds_batch_summary*.csv"))
    if len(summaries) != 1 or not summaries[0].is_file() or summaries[0].is_symlink():
        raise BondsLiteError("Expected exactly one new regular batch summary CSV")
    summary = summaries[0]
    try:
        with summary.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            columns = {"input_file", "status", "seconds", "output_folder_path", "output_suffix", "error"}
            if set(reader.fieldnames or ()) != columns:
                raise BondsLiteError(f"Invalid batch summary columns: {summary}")
            rows = list(reader)
        expected = {item.input_path.name: item for item in inputs.values()}
        if len(rows) != 2 or {row["input_file"] for row in rows} != set(expected):
            raise BondsLiteError("Batch summary must contain exactly one Bonds and one MAP row")
        for row in rows:
            item = expected[row["input_file"]]
            if row["status"] != "success" or row["error"].strip():
                raise BondsLiteError(f"{item.dataset} batch failed: {row['error'] or row['status']}")
            if Path(row["output_folder_path"]).resolve() != item.output_folder or row["output_suffix"] != item.output_suffix:
                raise BondsLiteError(f"Batch summary target/suffix mismatch for {item.dataset}")
        for item in inputs.values():
            for name in item.required_files:
                artifact = item.output_folder / _safe_model_filename(name)
                if artifact.is_symlink() or not artifact.is_file() or artifact.stat().st_size == 0:
                    raise BondsLiteError(f"Missing or empty {item.dataset} model artifact: {artifact}")
            if _file_hash(item.input_path) != item.input_hash:
                raise BondsLiteError(f"Stage 2 changed its input workbook: {item.input_path}")
    except BondsLiteError:
        raise
    except (OSError, ValueError, KeyError, TypeError, csv.Error) as error:
        raise BondsLiteError(f"Cannot validate model results: {error}") from error
    return summary


def confirm_model_overwrite(inputs: dict[str, ModelInput], run_folder: Path) -> bool:
    """Obtain explicit console consent only when previous results exist."""
    existing = any(any(item.output_folder.iterdir()) for item in inputs.values())
    if existing or any(run_folder.glob("bonds_batch_summary*.csv")):
        return input("Replace existing Bonds and MAP model results? [y/N]: ").strip().casefold() in {"y", "yes"}
    return True


@dataclass
class _ModelResultsCommit:
    committed: bool = False

    def commit(self) -> None:
        self.committed = True


def _remove_model_artifact(path: Path) -> None:
    if path.is_symlink() or path.is_file():
        path.unlink()
    elif path.is_dir():
        shutil.rmtree(path)


@contextmanager
def model_results_transaction(run_folder: Path, inputs: dict[str, ModelInput]) -> Iterator[_ModelResultsCommit]:
    """Stage old results outside the batch folder and restore unless safely committed."""
    backup = Path(tempfile.mkdtemp(prefix=f".bonds-rollback-{run_folder.name}-", dir=run_folder.parent))
    moved = []
    prepared = False
    decision = _ModelResultsCommit()
    try:
        for item in inputs.values():
            destination = backup / item.dataset
            destination.mkdir()
            for source in item.output_folder.iterdir():
                target = destination / source.name
                source.rename(target)
                moved.append((source, target))
        summaries = backup / "summaries"
        summaries.mkdir()
        for source in run_folder.glob("bonds_batch_summary*.csv"):
            target = summaries / source.name
            source.rename(target)
            moved.append((source, target))
        prepared = True
        yield decision
    except BaseException:
        decision.committed = False
        raise
    finally:
        if not decision.committed:
            try:
                if prepared:
                    for item in inputs.values():
                        if item.output_folder.is_symlink() or not item.output_folder.is_dir():
                            raise BondsLiteError(f"Model output directory changed unexpectedly: {item.output_folder}")
                        for artifact in item.output_folder.iterdir():
                            _remove_model_artifact(artifact)
                    for artifact in run_folder.glob("bonds_batch_summary*.csv"):
                        _remove_model_artifact(artifact)
                for original, stored in reversed(moved):
                    stored.rename(original)
            except BaseException as error:
                message = f"Model result recovery blocked; rollback backup retained at {backup}: {error}"
                logger.error(message)
                raise BondsLiteError(message) from error
        try:
            shutil.rmtree(backup)
        except OSError:
            logger.warning("Results settled, but rollback backup could not be removed: %s", backup, exc_info=True)


def _model_environment(config: ModelConfig) -> dict[str, str]:
    environment = os.environ.copy()
    root = config.python_executable.parent
    environment.pop("PYTHONHOME", None)
    environment.pop("PYTHONPATH", None)
    environment["CONDA_PREFIX"] = str(root)
    environment["PATH"] = os.pathsep.join((str(root), str(root / "Library" / "bin"), str(root / "Scripts"), environment.get("PATH", "")))
    environment["PYTHONUNBUFFERED"] = "1"
    return environment


def _attach_model_job(child: subprocess.Popen) -> object | None:
    if os.name != "nt" or child.pid is None:
        return None
    import win32job

    job = win32job.CreateJobObject(None, f"Local\\BondsLite_Model_{os.getpid()}_{time.time_ns()}")
    try:
        information = win32job.QueryInformationJobObject(job, win32job.JobObjectExtendedLimitInformation)
        information["BasicLimitInformation"]["LimitFlags"] |= win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        win32job.SetInformationJobObject(job, win32job.JobObjectExtendedLimitInformation, information)
        win32job.AssignProcessToJobObject(job, int(child._handle))
        resume = ctypes.WinDLL("ntdll").NtResumeProcess
        resume.argtypes = [ctypes.c_void_p]
        resume.restype = ctypes.c_long
        if resume(int(child._handle)) != 0:
            raise BondsLiteError("Cannot resume the contained model process")
    except BaseException:
        job.Close()
        raise
    return job


def _stop_model_child(child: subprocess.Popen) -> None:
    if child.poll() is None:
        child.terminate()
    try:
        child.wait(timeout=5)
    except subprocess.TimeoutExpired:
        child.kill()
        child.wait(timeout=5)


def _freeze_stage_two_curve_values(application: xw.App, input_path: Path) -> None:
    """Freeze validated XLL curve caches in a temporary Stage 2 input copy."""
    with owned_book(application, input_path, read_only=False) as workbook:
        curve_sheet = workbook.sheets["P. Base Curves"]
        table = curve_sheet.api.ListObjects("P_tbl_BaseCurves")
        for column_name in ("GBP SONIA", "GBP GOV", "GBP RPI"):
            range_api = table.ListColumns(column_name).DataBodyRange
            if range_api is None:
                raise BondsLiteError(f"P. Base Curves[{column_name}] has no cached values")
            curve_range = curve_sheet.range(range_api.Address)
            values = curve_range.options(ndim=2).value
            invalid_rows = [
                index + 1
                for index, row in enumerate(values)
                if len(row) != 1
                or isinstance(row[0], bool)
                or not isinstance(row[0], Real)
                or not math.isfinite(row[0])
            ]
            if not values or invalid_rows:
                raise BondsLiteError(
                    f"P. Base Curves[{column_name}] must contain finite numeric values; "
                    f"invalid rows: {invalid_rows[:20]}"
                )
            curve_range.value = values
            if curve_range.api.HasFormula is not False:
                raise BondsLiteError(
                    f"Could not freeze P. Base Curves[{column_name}] in temporary input"
                )
        workbook.save()
        if not workbook.api.Saved:
            raise BondsLiteError("Excel did not finish saving the staged Stage 2 input")


@contextmanager
def _stage_batch_inputs_locally(
    run_folder: Path, excel_application: xw.App | None = None,
) -> Iterator[Path]:
    """Copy the exact Bonds/MAP inputs to local storage for Excel COM access."""
    run_folder = run_folder.resolve()
    expected = set(INPUT_FILENAMES.values())
    actual = {
        path.name for path in run_folder.glob("*.xlsx")
        if not path.name.startswith("~$")
    }
    if actual != expected:
        raise BondsLiteError("Stage 2 requires exactly the two expected Bonds/MAP input workbooks")
    local_root = Path(os.environ.get("LOCALAPPDATA", tempfile.gettempdir())) / "Temp"
    if not local_root.is_dir():
        local_root = Path(tempfile.gettempdir())
    with tempfile.TemporaryDirectory(prefix="bonds-stage2-input-", dir=local_root) as temporary:
        batch_folder = Path(temporary)
        for filename in INPUT_FILENAMES.values():
            source = run_folder / filename
            local_copy = batch_folder / filename
            if source.is_symlink() or not source.is_file() or source.stat().st_size == 0:
                raise BondsLiteError(f"Invalid Stage 2 input workbook: {source}")
            shutil.copy2(source, local_copy)
            if _file_hash(source) != _file_hash(local_copy):
                raise BondsLiteError(f"Local Stage 2 input copy verification failed: {filename}")
            if excel_application is not None:
                with timed_stage(f"Freeze XLL curve values in {filename}"):
                    _freeze_stage_two_curve_values(excel_application, local_copy)
        yield batch_folder


def run_model_batch(
    run_folder: Path,
    config: ModelConfig,
    excel_application: xw.App | None = None,
) -> int:
    """Stream one contained folder batch with a deadline independent of stdout reads."""
    if not math.isfinite(config.timeout_seconds) or config.timeout_seconds <= 0:
        raise BondsLiteError("Model timeout must be a positive finite duration")
    run_folder = run_folder.resolve()
    started = time.monotonic()
    with _stage_batch_inputs_locally(run_folder, excel_application) as batch_folder:
        child = subprocess.Popen(
            [
                str(config.python_executable), "-u",
                str(config.model_dir / "run_bonds_batch.py"), str(batch_folder),
            ],
            cwd=str(config.model_dir), shell=False, stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT, text=True, errors="replace", bufsize=1,
            env=_model_environment(config),
            creationflags=0x00000004 if os.name == "nt" else 0,
        )
        job = None
        reader = None
        reader_errors = []

        def stream_output() -> None:
            try:
                for line in iter(child.stdout.readline, ""):
                    print(line, end="", flush=True)
            except (OSError, ValueError) as error:
                reader_errors.append(error)

        try:
            job = _attach_model_job(child)
            reader = threading.Thread(target=stream_output, name="bonds-model-output", daemon=True)
            reader.start()
            while True:
                remaining = config.timeout_seconds - (time.monotonic() - started)
                if remaining <= 0:
                    raise TimeoutError(f"Bonds Lite batch exceeded {config.timeout_seconds:g} seconds")
                try:
                    result = child.wait(timeout=min(0.25, remaining))
                    break
                except subprocess.TimeoutExpired:
                    continue
        except BaseException:
            if job is not None:
                job.Close()
                job = None
            _stop_model_child(child)
            raise
        finally:
            if job is not None:
                job.Close()
            if reader is not None:
                reader.join(timeout=5)
                if reader.is_alive():
                    raise BondsLiteError("Model output reader did not stop; recovery requires attention")
            if child.stdout is not None:
                child.stdout.close()

        for summary in batch_folder.glob("bonds_batch_summary*.csv"):
            summary.rename(run_folder / summary.name)
        if reader_errors:
            raise BondsLiteError(f"Cannot read model console output: {reader_errors[0]}")
        if result != 0:
            raise BondsLiteError(f"Bonds Lite batch exited with code {result}")
        return result


def configure_console_logging() -> None:
    """Configure one owned console handler without adding files or root handlers."""
    for handler in logger.handlers[:]:
        if handler.name == "bonds_lite_console":
            logger.removeHandler(handler)
            handler.close()
    console = logging.StreamHandler(sys.stdout)
    console.name = "bonds_lite_console"
    console.setLevel(logging.INFO)
    console.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(console)
    logger.setLevel(logging.INFO)
    logger.propagate = False


def _report_waiting(
    stop_event: threading.Event, state: _ProgressState, interval: float
) -> None:
    """Report elapsed waiting time without accessing any Excel objects."""
    while not stop_event.wait(interval):
        with state.lock:
            stage, started = (
                state.stages[-1] if state.stages else ("Processing", state.started)
            )
            now = time.perf_counter()
            logger.info(
                f"Waiting: {stage} | stage {now - started:.1f}s | "
                f"total {now - state.started:.1f}s"
            )


@contextmanager
def run_progress(heartbeat_interval: float = 5.0) -> Iterator[_ProgressState]:
    """Own one console reporter and stop it before returning or propagating errors."""
    state = _ProgressState(started=time.perf_counter())
    token = _progress_state.set(state)
    stop_event = threading.Event()
    reporter = threading.Thread(
        target=_report_waiting,
        args=(stop_event, state, heartbeat_interval),
        name="BondsLiteProgress",
        daemon=True,
    )
    try:
        reporter.start()
        try:
            yield state
        finally:
            stop_event.set()
            reporter.join()
    finally:
        _progress_state.reset(token)


@contextmanager
def timed_stage(name: str) -> Iterator[None]:
    """Log stage boundaries and durations while preserving the original exception."""
    started = time.perf_counter()
    state = _progress_state.get()
    token = None
    if state is None:
        state = _ProgressState(started=started)
        token = _progress_state.set(state)
    with state.lock:
        state.stages.append((name, started))
    logger.info(f"Start: {name} | total {started - state.started:.1f}s")
    try:
        yield
    except BaseException as error:
        now = time.perf_counter()
        logger.error(
            f"Failed: {name} | stage {now - started:.1f}s | "
            f"total {now - state.started:.1f}s | {error}"
        )
        raise
    else:
        now = time.perf_counter()
        logger.info(
            f"Done: {name} | stage {now - started:.1f}s | "
            f"total {now - state.started:.1f}s"
        )
    finally:
        with state.lock:
            state.stages.pop()
        if token is not None:
            _progress_state.reset(token)


STRESS_BLOCKS = tuple(
    f"{first}{row}:{last}{row + 6}"
    for row in (9, 20, 31, 42)
    for first, last in (("C", "J"), ("M", "DI"), ("DL", "HH"))
)

RAW_FIRST_COLUMN = 2
RAW_LAST_COLUMN = 1216
INPUT_FILENAMES = {
    "Bonds": "Bonds Lite input_Bonds.xlsx",
    "MAP": "Bonds Lite input_MAP.xlsx",
}

STAGE_ONE_FIELDS = {
    "last_run": "L74", "start": "L75", "finish": "L76", "elapsed": "L77",
    "user": "L78", "status": "L79", "last_output": "L80", "stage_two_input": "L82",
}
STAGE_ONE_PANEL = "H73:O83"
STAGE_ONE_FIELD_NAMES = {
    "last_run": "SST_Stage1_LastRun", "start": "SST_Stage1_Start",
    "finish": "SST_Stage1_Finish", "elapsed": "SST_Stage1_Elapsed",
    "user": "SST_Stage1_User", "status": "SST_Stage1_Status",
    "last_output": "SST_Stage1_LastSuccessfulOutput",
    "stage_two_input": "SST_Stage2_Input",
}
STAGE_ONE_NAMES = {
    "SST_Stage1_Panel": STAGE_ONE_PANEL,
    **{name: STAGE_ONE_FIELDS[field] for field, name in STAGE_ONE_FIELD_NAMES.items()},
}
STAGE_ONE_LABELS = {
    "last_run": "Last Run", "start": "Start:", "finish": "Finish:",
    "elapsed": "Time", "user": "User", "status": "Status",
    "last_output": "Last Successful Output", "stage_two_input": "Stage 2 Input",
}

STAGE_TWO_FIELD_NAMES = {
    "last_run": "SST_Stage2_LastRun", "start": "SST_Stage2_Start",
    "finish": "SST_Stage2_Finish", "elapsed": "SST_Stage2_Elapsed",
    "user": "SST_Stage2_User", "status": "SST_Stage2_Status",
    "input_run": "SST_Stage2_InputRun", "last_output": "SST_Stage2_LastSuccessfulResults",
}

VBA_LAUNCHER = '''Option Explicit

Private Function QuoteArgument(ByVal value As String) As String
    If InStr(value, Chr$(34)) > 0 Then
        Err.Raise vbObjectError + 401, , "A file path contains a quote character"
    End If
    QuoteArgument = Chr$(34) & value & Chr$(34)
End Function

Public Sub GenerateBondsLiteInput()
    LaunchBondsLite 1
End Sub

Public Sub RunBondsLite()
    LaunchBondsLite 2
End Sub

Private Sub LaunchBondsLite(ByVal stage As Long)
    On Error GoTo Failed
    Dim shell As Object
    Dim executable As String
    Dim command As String
    Set shell = CreateObject("WScript.Shell")
    executable = shell.Environment("User")("BONDS_LITE_PYTHON")
    If Len(executable) = 0 Then executable = Environ$("BONDS_LITE_PYTHON")
    If Len(executable) = 0 Then
        Err.Raise vbObjectError + 402, , _
            "Set BONDS_LITE_PYTHON to your Anaconda python.exe before running"
    End If
    If Not CreateObject("Scripting.FileSystemObject").FileExists(executable) Then
        Err.Raise vbObjectError + 403, , "Python executable not found: " & executable
    End If
    If Len(ThisWorkbook.Path) = 0 Then
        Err.Raise vbObjectError + 404, , "Save the Draft workbook before running"
    End If
    command = QuoteArgument(executable) & " -u -m bonds_lite_input --pause --caller " & _
        QuoteArgument(ThisWorkbook.FullName) & " --excel-hwnd " & CStr(Application.Hwnd) & _
        " --stage " & CStr(stage)
    shell.Run command, 1, False
    Exit Sub
Failed:
    MsgBox Err.Description, vbExclamation, "Bonds Lite input"
End Sub
'''


class BondsLiteError(ValueError):
    """An actionable input or Excel workflow validation failure."""


def column_name(number: int) -> str:
    """Return the Excel column label for a positive one-based column number."""
    if number < 1:
        raise BondsLiteError("Column number must be positive")
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def range_bounds(address: str) -> tuple[int, int, int, int]:
    """Return first row, first column, last row and last column for a range."""
    match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", address)
    if match is None:
        raise BondsLiteError(f"Invalid cell range: {address}")
    first_label, first_row, last_label, last_row = match.groups()
    columns = []
    for label in (first_label, last_label):
        number = 0
        for character in label:
            number = 26 * number + ord(character) - 64
        columns.append(number)
    return int(first_row), columns[0], int(last_row), columns[1]


def validate_stress_values(
    blocks: dict[str, list[list[object]]],
) -> dict[str, list[list[float]]]:
    """Validate complete numeric stress grids without changing decimal units."""
    errors = []
    result = {}
    for address in STRESS_BLOCKS:
        if address not in blocks:
            errors.append(f"{address}: missing stress block")
            continue
        first_row, first_column, last_row, last_column = range_bounds(address)
        grid = blocks[address]
        if len(grid) != last_row - first_row + 1 or any(
            len(row) != last_column - first_column + 1 for row in grid
        ):
            errors.append(f"{address}: incorrect row or column dimensions")
            continue
        converted = []
        for row_offset, row in enumerate(grid):
            converted_row = []
            for column_offset, value in enumerate(row):
                cell = (
                    f"{column_name(first_column + column_offset)}"
                    f"{first_row + row_offset}"
                )
                if value is None or isinstance(value, str) and not value.strip():
                    errors.append(f"{cell}: blank stress input")
                elif isinstance(value, bool) or not isinstance(value, Real):
                    errors.append(f"{cell}: a real numeric input is required")
                elif not math.isfinite(value):
                    errors.append(f"{cell}: a finite number is required")
                else:
                    converted_row.append(float(value))
            converted.append(converted_row)
        result[address] = converted
    if errors:
        raise BondsLiteError("Invalid credit stresses:\n" + "\n".join(errors))
    return result


def resolve_paths(
    caller_path: str | Path, bonds_path: str, map_path: str, template_path: str
) -> tuple[Path, Path, Path, Path]:
    """Resolve existing distinct workbooks against the saved caller's folder."""
    values = (caller_path, bonds_path, map_path, template_path)
    for value in values:
        if not str(value).strip():
            raise BondsLiteError("A saved workbook, Bonds data, MAP data and template paths are required")
        if re.match(r"^[a-z]+://", str(value), re.IGNORECASE):
            raise BondsLiteError("Use a local synchronised path, not a workbook URL")
    caller = Path(str(caller_path).strip()).expanduser().resolve()
    resolved = [caller]
    for value in (bonds_path, map_path, template_path):
        path = Path(str(value).strip()).expanduser()
        resolved.append((caller.parent / path).resolve())
    for path in resolved:
        if not path.is_file():
            raise BondsLiteError(f"Workbook does not exist or is not a file: {path}")
    for index, path in enumerate(resolved):
        if any(path.samefile(other) for other in resolved[:index]):
            raise BondsLiteError("Caller, Bonds data, MAP data and template must be distinct files")
    return resolved[0], resolved[1], resolved[2], resolved[3]


def validate_headers(source: list[str], destination: list[str]) -> None:
    """Require the complete source schema to match the destination in order."""
    if not source or source != destination or any(
        not isinstance(header, str) or not header.strip() for header in source
    ):
        raise BondsLiteError("Bond column headers do not match the template schema")


def validate_bond_headers(source: list[str], destination: list[str]) -> None:
    """Require raw B:ATT headers to match the destination's raw-data prefix."""
    raw_width = RAW_LAST_COLUMN - RAW_FIRST_COLUMN + 1
    if len(source) != raw_width or len(destination) < raw_width:
        raise BondsLiteError(
            "Bond raw-data headers must cover B:ATT (1,215 columns); "
            "calculated columns belong only in the Bonds Lite input template"
        )
    validate_headers(source, destination[:raw_width])


def validate_identifiers(values: list[object]) -> None:
    """Reject an empty dataset and bond records without identifiers."""
    if not values:
        raise BondsLiteError("Bond data is empty")
    rows = [
        str(index + 1)
        for index, value in enumerate(values)
        if value is None or isinstance(value, str) and not value.strip()
    ]
    if rows:
        raise BondsLiteError("Missing bond identifiers in body rows: " + ", ".join(rows))


def choose_output_path(folder: Path, timestamp: datetime) -> Path:
    """Find an unused timestamped filename without creating or replacing files."""
    stem = f"Bonds Lite input_{timestamp:%Y%m%d_%H%M%S}"
    candidate = folder / f"{stem}.xlsx"
    suffix = 0
    while candidate.exists():
        suffix += 1
        candidate = folder / f"{stem}_{suffix:03d}.xlsx"
    return candidate


def allocate_run_folder(output_root: Path, timestamp: datetime) -> Path:
    """Reserve one new timestamped run directory without replacing earlier runs."""
    output_root.mkdir(parents=True, exist_ok=True)
    stem = timestamp.strftime("%Y%m%d_%H%M%S")
    suffix = 0
    while True:
        candidate = output_root / (stem if suffix == 0 else f"{stem}_{suffix:03d}")
        try:
            candidate.mkdir()
        except FileExistsError:
            suffix += 1
            continue
        return candidate


def publish_output(working_path: Path, output_folder: Path, timestamp: datetime) -> Path:
    """Move a validated working file without replacing another Windows output."""
    if not working_path.is_file():
        raise BondsLiteError(f"Working workbook is missing or is not a file: {working_path}")
    while True:
        candidate = choose_output_path(output_folder, timestamp)
        try:
            working_path.rename(candidate)
        except FileExistsError:
            continue
        return candidate


def _caller_path(caller: xw.Book) -> Path:
    """Use the caller's local synchronised path, not its native SharePoint URL."""
    value = caller.fullname
    if not value or re.match(r"^[a-z]+://", value, re.IGNORECASE):
        raise BondsLiteError("Save and open Draft from its local synchronised folder")
    return Path(value).expanduser().resolve()


def portable_output_path(caller_path: Path, output_path: Path) -> str:
    """Return the relative folder of a complete, published Bonds/MAP input pair."""
    path = Path(output_path)
    folder = Path(caller_path).resolve().parent
    output_folder = (folder / "Bonds Lite output").resolve()
    if (
        ".." in path.parts
        or path.is_symlink()
        or re.fullmatch(r"\d{8}_\d{6}(?:_\d{3,})?", path.name) is None
    ):
        raise BondsLiteError("Stage output must be a timestamped run folder; rerun Stage 1 for both inputs")
    resolved = path.resolve()
    if (
        resolved.parent != output_folder
        or not resolved.is_dir()
    ):
        raise BondsLiteError(f"Stage run folder is missing or outside Bonds Lite output: {path}")
    for dataset, filename in INPUT_FILENAMES.items():
        member = resolved / filename
        if (
            member.is_symlink()
            or not member.is_file()
            or member.resolve().parent != resolved
            or member.stat().st_size == 0
        ):
            raise BondsLiteError(f"Stage run folder has no complete {dataset} input: {member}")
    try:
        return resolved.relative_to(folder).as_posix()
    except ValueError as error:
        raise BondsLiteError("Stage run folder must remain inside Draft's output tree") from error


def _named_audit_region(caller: xw.Book, name: str) -> xw.Range:
    """Resolve one valid local-sheet audit name without creating fallback cells."""
    try:
        region = caller.names[name].refers_to_range
        if (
            region.sheet.name != "I. Control"
            or os.path.normcase(region.sheet.book.fullname) != os.path.normcase(caller.fullname)
            or region.api.Areas.Count != 1
        ):
            raise BondsLiteError(f"Stage audit panel name is not a single range in this workbook: {name}")
        return region
    except BondsLiteError:
        raise
    except Exception as error:
        raise BondsLiteError(f"Stage audit panel name is missing or invalid: {name}") from error


def _audit_bounds(region: xw.Range) -> tuple[int, int, int, int]:
    """Return inclusive bounds without assuming where the panel was installed."""
    return (
        region.row, region.column,
        region.row + region.rows.count - 1,
        region.column + region.columns.count - 1,
    )


def _stage_field_cells(caller: xw.Book) -> dict[str, xw.Range]:
    """Resolve each named field to the top-left cell of its current merged area."""
    cells = {}
    for field, name in STAGE_ONE_FIELD_NAMES.items():
        region = caller.names[name].refers_to_range
        merged = region.api.Cells(1, 1).MergeArea
        cells[field] = region.sheet.range((merged.Row, merged.Column))
    return cells


def require_stage_panel(caller: xw.Book) -> None:
    """Validate named field ownership and geometry at the panel's current location."""
    _require_named_stage_panel(caller, "SST_Stage1_Panel", STAGE_ONE_FIELD_NAMES)


def _require_named_stage_panel(caller: xw.Book, panel_name: str, field_names: dict[str, str]) -> None:
    if caller.api.ReadOnly:
        raise BondsLiteError("Draft is read-only; its stage audit cannot be saved")
    sheet = caller.sheets["I. Control"]
    if sheet.api.ProtectContents:
        raise BondsLiteError("I. Control is protected; its stage audit cannot be updated")
    panel = _named_audit_region(caller, panel_name)
    top, left, bottom, right = _audit_bounds(panel)
    occupied = []
    for name in field_names.values():
        region = _named_audit_region(caller, name)
        first = region.api.Cells(1, 1).MergeArea
        merged = sheet.range(
            (first.Row, first.Column),
            (first.Row + first.Rows.Count - 1, first.Column + first.Columns.Count - 1),
        )
        bounds = _audit_bounds(merged)
        if region.size != 1 and _audit_bounds(region) != bounds:
            raise BondsLiteError(f"Stage field must name one value cell or merged value area: {name}")
        field_top, field_left, field_bottom, field_right = bounds
        if not (top <= field_top <= field_bottom <= bottom and left <= field_left <= field_right <= right):
            raise BondsLiteError(f"Stage field is outside the named audit panel: {name}")
        if any(
            field_top <= old_bottom and field_bottom >= old_top
            and field_left <= old_right and field_right >= old_left
            for old_top, old_left, old_bottom, old_right in occupied
        ):
            raise BondsLiteError(f"Stage audit fields overlap: {name}")
        occupied.append(bounds)


def _upgrade_stage_panel_names(caller: xw.Book) -> None:
    """Name legacy fields only when the original relative layout is unambiguous."""
    panel = _named_audit_region(caller, "SST_Stage1_Panel")
    if panel.shape != (11, 8):
        raise BondsLiteError("Legacy stage panel was resized; its unnamed fields cannot be inferred safely")
    sheet = panel.sheet
    anchors = {}
    for field, address in STAGE_ONE_FIELDS.items():
        row, column, _, _ = range_bounds(f"{address}:{address}")
        anchor = sheet.range((panel.row + row - 73, panel.column + column - 8))
        label = sheet.range((anchor.row, panel.column))
        if label.value != STAGE_ONE_LABELS[field]:
            raise BondsLiteError("Legacy stage panel layout changed; no names were added")
        anchors[field] = anchor
    for field in ("status", "last_output", "stage_two_input"):
        region = _named_audit_region(caller, STAGE_ONE_FIELD_NAMES[field])
        if region.row != anchors[field].row or region.column != anchors[field].column:
            raise BondsLiteError("Legacy stage field names are inconsistent; no names were added")
    added = []
    try:
        for field in ("last_run", "start", "finish", "elapsed", "user"):
            name = STAGE_ONE_FIELD_NAMES[field]
            caller.names.add(name, f"='I. Control'!{anchors[field].get_address()}")
            added.append(name)
        require_stage_panel(caller)
    except Exception:
        for name in added:
            caller.names[name].delete()
        raise


def initialise_stage_panel(caller: xw.Book) -> None:
    """Install the last-run panel only in a verified free area beside the button."""
    if caller.api.ReadOnly:
        raise BondsLiteError("Draft is read-only; the stage audit cannot be installed")
    sheet = caller.sheets["I. Control"]
    if sheet.api.ProtectContents:
        raise BondsLiteError("I. Control is protected; the stage audit cannot be installed")
    existing = {name.name.split("!")[-1] for name in caller.names}
    present = existing.intersection(STAGE_ONE_NAMES)
    if present:
        legacy = {
            "SST_Stage1_Panel", "SST_Stage1_Status",
            "SST_Stage1_LastSuccessfulOutput", "SST_Stage2_Input",
        }
        if present == legacy:
            _upgrade_stage_panel_names(caller)
            return
        if present != set(STAGE_ONE_NAMES):
            raise BondsLiteError("Stage audit names are incomplete; installation was not changed")
        require_stage_panel(caller)
        return
    area = sheet.range(STAGE_ONE_PANEL)
    if caller.app.api.WorksheetFunction.CountA(area.api) or area.api.MergeCells is not False:
        raise BondsLiteError("Audit area H73:O83 contains cells or merges; nothing was overwritten")
    for shape in sheet.api.Shapes:
        if (
            shape.Left < area.left + area.width
            and shape.Left + shape.Width > area.left
            and shape.Top < area.top + area.height
            and shape.Top + shape.Height > area.top
        ):
            raise BondsLiteError("Audit area H73:O83 overlaps an existing shape or button")

    events = caller.app.api.EnableEvents
    try:
        caller.app.api.EnableEvents = False
        area.color = (255, 230, 230)
        area.font.name = "Calibri"
        area.font.size = 10
        area.api.VerticalAlignment = -4108
        heading = sheet.range("H73:O73")
        heading.merge()
        heading.value = "Stage 1 - Generate Bonds Lite Input"
        heading.color = (255, 127, 127)
        heading.font.color = (255, 255, 255)
        heading.font.bold = True
        for field, address in STAGE_ONE_FIELDS.items():
            row = sheet.range(address).row
            end_row = row + 1 if field in ("last_output", "stage_two_input") else row
            label = sheet.range((row, 8), (end_row, 11))
            value = sheet.range((row, 12), (end_row, 15))
            label.merge()
            value.merge()
            label.value = STAGE_ONE_LABELS[field]
            label.color = (255, 127, 127)
            label.font.color = (255, 255, 255)
            label.api.WrapText = True
            value.api.WrapText = True
            value.font.color = (0, 0, 0)
            value.api.HorizontalAlignment = -4131
        sheet.range(STAGE_ONE_FIELDS["last_run"]).number_format = "dd/mm/yyyy"
        for field in ("start", "finish"):
            sheet.range(STAGE_ONE_FIELDS[field]).number_format = "hh:mm:ss"
        sheet.range(STAGE_ONE_FIELDS["elapsed"]).number_format = "[h]:mm:ss"
        sheet.range(STAGE_ONE_FIELDS["status"]).value = "Not Run"
        for name, address in STAGE_ONE_NAMES.items():
            caller.names.add(name, f"='I. Control'!{sheet.range(address).get_address()}")
    finally:
        caller.app.api.EnableEvents = events


def write_stage_audit(caller: xw.Book, fields: dict[str, object]) -> None:
    """Write the mapped audit fields and persist Draft with events suppressed."""
    require_stage_panel(caller)
    if set(fields) - STAGE_ONE_FIELDS.keys():
        raise BondsLiteError("Unknown Stage 1 audit field")
    cells = _stage_field_cells(caller)
    events = caller.app.api.EnableEvents
    try:
        caller.app.api.EnableEvents = False
        for field, value in fields.items():
            cell = cells[field]
            if field in ("last_output", "stage_two_input"):
                if cell.api.Hyperlinks.Count:
                    cell.api.Hyperlinks.Delete()
                if value:
                    cell.add_hyperlink(str(value), text_to_display=str(value))
                else:
                    cell.value = None
            else:
                cell.value = f"'{value}" if isinstance(value, str) else value
        caller.save()
        if not caller.api.Saved:
            raise BondsLiteError("Draft's stage audit was not saved")
    finally:
        caller.app.api.EnableEvents = events


def _read_stage_fields(caller: xw.Book) -> dict[str, object]:
    """Read only the configured Stage 1 summary and Stage 2 handoff cells."""
    return {field: cell.value for field, cell in _stage_field_cells(caller).items()}


def require_stage_two_panel(caller: xw.Book) -> None:
    """Validate the independently named Stage 2 panel without fixed addresses."""
    _require_named_stage_panel(caller, "SST_Stage2_Panel", STAGE_TWO_FIELD_NAMES)
    first = _audit_bounds(_named_audit_region(caller, "SST_Stage1_Panel"))
    second = _audit_bounds(_named_audit_region(caller, "SST_Stage2_Panel"))
    if first[0] <= second[2] and first[2] >= second[0] and first[1] <= second[3] and first[3] >= second[1]:
        raise BondsLiteError("Stage 1 and Stage 2 audit panels overlap")


def _stage_two_cells(caller: xw.Book) -> dict[str, xw.Range]:
    cells = {}
    for field, name in STAGE_TWO_FIELD_NAMES.items():
        region = caller.names[name].refers_to_range
        merged = region.api.Cells(1, 1).MergeArea
        cells[field] = region.sheet.range((merged.Row, merged.Column))
    return cells


def _read_stage_two_fields(caller: xw.Book) -> dict[str, object]:
    """Read the attempt identity separately from Stage 1's active handoff."""
    return {field: cell.value for field, cell in _stage_two_cells(caller).items()}


def write_stage_two_audit(caller: xw.Book, fields: dict[str, object]) -> None:
    """Persist named Stage 2 fields with relative folder-only hyperlinks."""
    require_stage_two_panel(caller)
    if set(fields) - STAGE_TWO_FIELD_NAMES.keys():
        raise BondsLiteError("Unknown Stage 2 audit field")
    cells = _stage_two_cells(caller)
    events = caller.app.api.EnableEvents
    try:
        caller.app.api.EnableEvents = False
        for field, value in fields.items():
            cell = cells[field]
            if field in ("last_output", "input_run"):
                if cell.api.Hyperlinks.Count:
                    cell.api.Hyperlinks.Delete()
                if value:
                    cell.add_hyperlink(str(value), text_to_display=str(value))
                else:
                    cell.value = None
            else:
                cell.value = f"'{value}" if isinstance(value, str) else value
        caller.save()
        if not caller.api.Saved:
            raise BondsLiteError("Draft's Stage 2 audit was not saved")
    finally:
        caller.app.api.EnableEvents = events


def initialise_stage_two_panel(caller: xw.Book) -> None:
    """Place a new audit below the current Stage 1 panel and button, never over cells."""
    require_stage_panel(caller)
    expected = {"SST_Stage2_Panel", *STAGE_TWO_FIELD_NAMES.values()}
    existing = {name.name.split("!")[-1] for name in caller.names}
    if existing.intersection(expected):
        if not expected.issubset(existing):
            raise BondsLiteError("Stage 2 audit names are incomplete; installation was not changed")
        require_stage_two_panel(caller)
        return
    sheet = caller.sheets["I. Control"]
    previous = _named_audit_region(caller, "SST_Stage1_Panel")
    bottom = previous.top + previous.height
    for shape in sheet.api.Shapes:
        if shape.Name == "btnGenerateBondsLiteInput":
            bottom = max(bottom, shape.Top + shape.Height)
    row, column = previous.row + previous.rows.count + 2, previous.column
    while sheet.range((row, column)).top < bottom + 12:
        row += 1
    area = sheet.range((row, column), (row + 10, column + 7))
    if caller.app.api.WorksheetFunction.CountA(area.api) or area.api.MergeCells is not False:
        raise BondsLiteError("Proposed Stage 2 audit area contains cells or merges; nothing was overwritten")
    for shape in sheet.api.Shapes:
        if shape.Left < area.left + area.width and shape.Left + shape.Width > area.left and shape.Top < area.top + area.height and shape.Top + shape.Height > area.top:
            raise BondsLiteError("Proposed Stage 2 audit area overlaps an existing shape or button")
    events = caller.app.api.EnableEvents
    try:
        caller.app.api.EnableEvents = False
        area.color, area.font.name, area.font.size = (255, 230, 230), "Calibri", 10
        area.api.VerticalAlignment = -4108
        heading = sheet.range((row, column), (row, column + 7))
        heading.merge()
        heading.value, heading.color = "Stage 2 - Run Bonds Lite", (255, 127, 127)
        heading.font.color, heading.font.bold = (255, 255, 255), True
        labels = {**STAGE_ONE_LABELS, "input_run": "Input Run", "last_output": "Last Successful Results"}
        offsets = {"last_run": 1, "start": 2, "finish": 3, "elapsed": 4, "user": 5, "status": 6, "input_run": 7, "last_output": 9}
        caller.names.add("SST_Stage2_Panel", f"='I. Control'!{area.get_address()}")
        for field, offset in offsets.items():
            end = row + offset + (1 if field in ("input_run", "last_output") else 0)
            label = sheet.range((row + offset, column), (end, column + 3))
            value = sheet.range((row + offset, column + 4), (end, column + 7))
            label.merge()
            value.merge()
            label.value, label.color, label.font.color = labels[field], (255, 127, 127), (255, 255, 255)
            label.api.WrapText = value.api.WrapText = True
            value.font.color, value.api.HorizontalAlignment = (0, 0, 0), -4131
            value.number_format = "dd/mm/yyyy" if field == "last_run" else "hh:mm:ss" if field in ("start", "finish") else "[h]:mm:ss" if field == "elapsed" else "General"
            if field == "status":
                value.value = "Not Run"
            caller.names.add(STAGE_TWO_FIELD_NAMES[field], f"='I. Control'!{value[0, 0].get_address()}")
    finally:
        caller.app.api.EnableEvents = events


def resolve_stage_input(caller: xw.Book) -> Path:
    """Resolve the committed run folder, not an arbitrarily selected input file."""
    require_stage_panel(caller)
    fields = _read_stage_fields(caller)
    if not caller.api.Saved:
        raise BondsLiteError("Save Draft's successful Stage 1 audit before running Stage 2")
    value = fields["stage_two_input"]
    if fields["status"] != "Succeeded" or not value or value != fields["last_output"]:
        raise BondsLiteError("Stage 2 requires a successfully saved Stage 1 output")
    if not isinstance(value, str) or re.match(r"^[a-z]+://", value, re.IGNORECASE):
        raise BondsLiteError("Stage 2 input must be a relative output path")
    relative = Path(value)
    if relative.is_absolute() or relative.drive or ".." in relative.parts:
        raise BondsLiteError("Stage 2 input must not be absolute or contain parent traversal")
    caller_path = _caller_path(caller)
    output = caller_path.parent / relative
    portable_output_path(caller_path, output)
    return output.resolve()


def resolve_stage_inputs(caller: xw.Book) -> dict[str, Path]:
    """Return explicitly labelled Bonds and MAP inputs from a committed run."""
    folder = resolve_stage_input(caller)
    return {dataset: folder / filename for dataset, filename in INPUT_FILENAMES.items()}


def read_credit_stresses(caller: xw.Book) -> dict[str, list[list[float]]]:
    """Read and validate the live caller's complete credit-stress blocks."""
    sheet = caller.sheets["I. Credit Stresses"]
    _validate_credit_headers(sheet)
    return validate_stress_values(
        {
            address: sheet.range(address).options(ndim=2, err_to_str=True).value
            for address in STRESS_BLOCKS
        }
    )


def _validate_credit_headers(sheet: xw.Sheet) -> None:
    """Require the fixed rating/year layout before reading or writing stress blocks."""
    ratings = ["GOV", "AAA", "AA", "A", "BBB", "BB", "B"]
    for row in (9, 20, 31, 42):
        for address, expected in (
            (f"B{row}:B{row + 6}", [[rating] for rating in ratings]),
            (f"C{row - 1}:J{row - 1}", [ratings + ["D"]]),
            (f"L{row}:L{row + 6}", [[rating] for rating in ratings]),
            (f"M{row - 1}:DI{row - 1}", [[str(year) for year in range(101)]]),
            (f"DK{row}:DK{row + 6}", [[rating] for rating in ratings]),
            (f"DL{row - 1}:HH{row - 1}", [[str(year) for year in range(101)]]),
        ):
            grid = sheet.range(address).options(ndim=2).value
            labels = [
                [format(value, "g") if isinstance(value, Real) and not isinstance(value, bool) else value for value in line]
                for line in grid
            ]
            if labels != expected:
                raise BondsLiteError(f"Stress rating/year headers differ at {address}")


def populate_credit_stresses(
    settings: StageOneInputs, working: xw.Book,
    stresses: dict[str, list[list[float]]],
) -> None:
    """Populate a validated destination using captured stress values only."""
    destination = working.sheets["I. Sensitivity"]
    _validate_credit_headers(destination)
    for address, grid in stresses.items():
        destination.range(address).value = grid
        if destination.range(address).options(ndim=2).value != grid:
            raise BondsLiteError(f"Credit stress read-back failed at {address}")


@timed_stage("Inspect template formula columns")
def formula_columns(table: object) -> dict[int, str]:
    """Capture uniform calculated columns and reject mixed/legacy array inputs."""
    body = table.DataBodyRange
    if body is None:
        raise BondsLiteError("The template bond table has no formula seed rows")
    if body.HasFormula is False:
        raise BondsLiteError("The template bond table has lost its calculated columns")
    cells = body.SpecialCells(-4123)
    columns = set()
    for area in cells.Areas:
        columns.update(range(area.Column, area.Column + area.Columns.Count))
    result = {}
    for column in sorted(columns):
        region = table.Parent.Range(
            table.Parent.Cells(body.Row, column),
            table.Parent.Cells(body.Row + body.Rows.Count - 1, column),
        )
        if region.HasFormula is not True:
            raise BondsLiteError(f"Mixed formulas and values in {column_name(column)}")
        if region.HasArray is not False:
            raise BondsLiteError(
                f"Legacy array formulas in {column_name(column)} require template review"
            )
        formulas = region.Formula2R1C1
        seed = formulas if isinstance(formulas, str) else formulas[0][0]
        if not isinstance(formulas, str) and any(
            row[0] != seed for row in formulas
        ):
            raise BondsLiteError(f"Inconsistent formulas in {column_name(column)}")
        result[column] = seed
    return result


def native_copy_values(source: xw.Range, destination: xw.Range) -> None:
    """Copy values inside Excel and verify them using scalar Excel-side checks."""
    application = source.sheet.book.app
    if application.pid != destination.sheet.book.app.pid:
        raise BondsLiteError("Native copying requires the same Excel instance")
    if source.shape != destination.shape:
        raise BondsLiteError("Source and destination copy dimensions differ")
    source_address = source.get_address()
    destination_address = destination.get_address(external=True)
    expected = source.size
    try:
        errors = source.sheet.api.Evaluate(
            f"SUMPRODUCT(--ISERROR({source_address}))"
        )
        if not isinstance(errors, Real) or isinstance(errors, bool) or errors != 0:
            raise BondsLiteError(f"Excel error in raw bond inputs at {source_address}")
        source.api.Copy()
        destination.api.PasteSpecial(Paste=-4163, SkipBlanks=False, Transpose=False)
        comparisons = [
            f"EXACT({source_address},{destination_address})",
            *(
                f"{predicate}({source_address})={predicate}({destination_address})"
                for predicate in ("ISTEXT", "ISNUMBER", "ISBLANK")
            ),
        ]
        for comparison in comparisons:
            matches = source.sheet.api.Evaluate(f"SUMPRODUCT(--({comparison}))")
            if (
                not isinstance(matches, Real)
                or isinstance(matches, bool)
                or matches != expected
            ):
                raise BondsLiteError(
                    f"Native copy verification failed at {source_address}"
                )
    finally:
        application.api.CutCopyMode = False


def populate_bond_data(data: xw.Book, working: xw.Book) -> tuple[int, list[int]]:
    """Resize the destination table and copy raw input blocks inside Excel."""
    source_sheet = data.sheets["I. Bonds Data"]
    target_sheet = working.sheets["I. Bonds Data"]
    source = source_sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    target = target_sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    validate_bond_headers(
        list(source.HeaderRowRange.Value2[0]), list(target.HeaderRowRange.Value2[0])
    )
    if source.Range.Column != RAW_FIRST_COLUMN or target.Range.Column != RAW_FIRST_COLUMN:
        raise BondsLiteError("The source and destination bond tables must start in column B")
    if target_sheet.api.ProtectContents:
        raise BondsLiteError("The destination bond sheet is protected")
    count = source.ListRows.Count
    if count == 0:
        raise BondsLiteError("Bond data is empty")
    identifiers = source.ListColumns("ISIN").DataBodyRange.Value2
    validate_identifiers([row[0] for row in identifiers])
    formulas = formula_columns(target)
    header_row = target.HeaderRowRange.Row
    first_column = target.Range.Column
    last_column = first_column + target.ListColumns.Count - 1
    old_last = header_row + target.ListRows.Count
    new_last = header_row + count
    if new_last > 1048576:
        raise BondsLiteError("Bond row count exceeds Excel's worksheet limit")
    if new_last > old_last:
        growth = target_sheet.range((old_last + 1, first_column), (new_last, last_column))
        if working.app.api.WorksheetFunction.CountA(growth.api):
            raise BondsLiteError("Table growth would overwrite content below the table")
    if new_last != old_last:
        with timed_stage("Resize bond table"):
            target.Resize(target_sheet.range((header_row, first_column), (new_last, last_column)).api)
    if new_last < old_last:
        target_sheet.range((new_last + 1, first_column), (old_last, last_column)).clear_contents()
    if any(column not in formulas for column in range(RAW_LAST_COLUMN + 1, last_column + 1)):
        raise BondsLiteError("Destination columns after ATT must contain template formulas")
    raw_columns = [
        column for column in range(first_column, RAW_LAST_COLUMN + 1)
        if column not in formulas
    ]
    intervals = []
    for column in raw_columns:
        if intervals and intervals[-1][1] == column - 1:
            intervals[-1] = (intervals[-1][0], column)
        else:
            intervals.append((column, column))
    for first, last in intervals:
        source_first = source.Range.Column + first - first_column
        source_last = source.Range.Column + last - first_column
        origin = source_sheet.range(
            (source.HeaderRowRange.Row + 1, source_first),
            (source.HeaderRowRange.Row + count, source_last),
        )
        destination = target_sheet.range(
            (header_row + 1, first), (new_last, last)
        )
        with timed_stage(f"Native value copy: {column_name(first)}:{column_name(last)}"):
            native_copy_values(origin, destination)
    if new_last > old_last:
        with timed_stage("Fill formulas in newly added bond rows"):
            for column, formula in formulas.items():
                region = target_sheet.range((old_last + 1, column), (new_last, column))
                region.api.Formula2R1C1 = formula
                if region.api.HasFormula is not True:
                    raise BondsLiteError(f"Formula fill failed in {column_name(column)}")
    return count, list(formulas)


@contextmanager
def excel_state(application: xw.App) -> Iterator[None]:
    """Disable prompts and macros temporarily, restoring all changed settings."""
    settings = {
        "AutomationSecurity": 3,
        "EnableEvents": False,
        "DisplayAlerts": False,
        "ScreenUpdating": False,
        "Calculation": -4135,
        "CalculateBeforeSave": False,
    }
    with timed_stage("Read Excel settings"):
        originals = {name: getattr(application.api, name) for name in settings}
    try:
        with timed_stage("Configure Excel settings"):
            for name, value in settings.items():
                setattr(application.api, name, value)
        yield
    finally:
        with timed_stage("Restore Excel settings"):
            failures = []
            for name, value in reversed(list(originals.items())):
                try:
                    setattr(application.api, name, value)
                except Exception as error:
                    failures.append(f"{name}: {error}")
            if failures:
                raise BondsLiteError("Unable to restore Excel settings: " + "; ".join(failures))


@contextmanager
def owned_book(application: xw.App, path: Path, read_only: bool = True) -> Iterator[xw.Book]:
    """Open a private workbook without macros or automatic external-link updates."""
    if any(book.name.casefold() == path.name.casefold() for book in application.books):
        raise BondsLiteError(f"Close the input workbook before running: {path.name}")
    with timed_stage(f"Open workbook: {path.name}"):
        book = application.books.open(
            str(path), update_links=False, read_only=read_only,
            ignore_read_only_recommended=True, notify=False,
        )
    try:
        if not read_only and book.api.ReadOnly:
            raise BondsLiteError(f"Working workbook is read-only: {path}")
        yield book
    finally:
        with timed_stage(f"Close workbook: {path.name}"):
            book.close()


@contextmanager
def _stage_source_locally(path: Path) -> Iterator[Path]:
    """Give Excel a verified local copy of a read-only Stage 1 source workbook."""
    source = path.resolve(strict=True)
    local_root = Path(os.environ.get("LOCALAPPDATA", tempfile.gettempdir())) / "Temp"
    if not local_root.is_dir():
        local_root = Path(tempfile.gettempdir())
    with tempfile.TemporaryDirectory(prefix="bonds-lite-source-", dir=local_root) as temporary:
        local_copy = Path(temporary) / source.name
        shutil.copy2(source, local_copy)
        if _file_hash(source) != _file_hash(local_copy):
            raise BondsLiteError(f"Local source copy verification failed: {source.name}")
        yield local_copy


def validate_entity(entity: str) -> str:
    """Require entity text that can be used unchanged in Windows folder names."""
    if (
        not isinstance(entity, str)
        or not entity
        or entity != entity.strip()
        or re.search(r'[<>:"/\\|?*\x00-\x1f\x7f]', entity)
    ):
        raise BondsLiteError("I. Control!D9 must contain a Windows-safe entity without surrounding whitespace")
    return entity


def validate_inputs(caller: xw.Book) -> StageOneInputs:
    """Validate file controls and live stress inputs before creating working files."""
    control = caller.sheets["I. Control"]
    paths = resolve_paths(
        caller.fullname, control.range("D39").value or "",
        control.range("D40").value or "", control.range("D42").value or "",
    )
    date = control.range("D11").value
    entity = validate_entity(control.range("D9").value)
    if not isinstance(date, datetime):
        raise BondsLiteError("I. Control!D11 must contain a valid Excel valuation date")
    return StageOneInputs(*paths, date, entity, read_credit_stresses(caller))


def synchronise_controls(settings: StageOneInputs, working: xw.Book) -> None:
    """Synchronise only the approved valuation date and entity controls."""
    destination = working.sheets["I. Control"]
    destination.range("D10").value = settings.valuation_date
    destination.range("D13").value = settings.entity


def _model_output_cell(working: xw.Book) -> xw.Range:
    """Resolve the template's output-folder name without a fixed cell address."""
    try:
        region = working.names["output_folder_path"].refers_to_range
        if (
            region.sheet.name != "I. Control"
            or os.path.normcase(region.sheet.book.fullname) != os.path.normcase(working.fullname)
            or region.api.Areas.Count != 1
        ):
            raise BondsLiteError("output_folder_path must refer to this input's I. Control")
        merged = region.api.Cells(1, 1).MergeArea
        if region.size != 1 and _audit_bounds(region) != (
            merged.Row, merged.Column,
            merged.Row + merged.Rows.Count - 1,
            merged.Column + merged.Columns.Count - 1,
        ):
            raise BondsLiteError("output_folder_path must name one value cell")
        return region.sheet.range((merged.Row, merged.Column))
    except BondsLiteError:
        raise
    except Exception as error:
        raise BondsLiteError("Input template output_folder_path name is missing or invalid") from error


def set_model_output_folder(working: xw.Book, folder: Path) -> None:
    """Set this generated input's existing final model-output folder."""
    resolved = folder.resolve()
    if not resolved.is_dir():
        raise BondsLiteError(f"Model output folder does not exist: {folder}")
    _model_output_cell(working).value = str(resolved)


def required_links(working: xw.Book, columns: list[int]) -> list[str]:
    """Find external dependencies in bond formulas and referenced local tables."""
    sheet = working.sheets["I. Bonds Data"]
    formulas = [sheet.range((9, column)).api.Formula2 for column in columns]
    seen = set()
    pending = list(formulas)
    tables = {
        table.Name: table
        for worksheet in working.sheets
        for table in worksheet.api.ListObjects
    }
    while pending:
        formula = pending.pop()
        for name in re.findall(r"\b([A-Za-z_][A-Za-z0-9_.]*)\[", formula):
            if name in seen or name not in tables:
                continue
            seen.add(name)
            body = tables[name].DataBodyRange
            if body is not None and body.HasFormula is not False:
                for area in body.SpecialCells(-4123).Areas:
                    values = area.Formula2
                    nested = [[values]] if isinstance(values, str) else values
                    for row in nested:
                        for value in row:
                            if isinstance(value, str) and value.startswith("="):
                                formulas.append(value)
                                pending.append(value)
    text = "\n".join(formulas).casefold()
    return [link for link in (working.api.LinkSources(1) or ()) if f"[{Path(link).name}]".casefold() in text]


def refresh_and_validate(working: xw.Book, columns: list[int], timeout: float = 180.0) -> None:
    """Refresh required sources and reject errors in recalculated bond columns."""
    import pythoncom
    from pywintypes import com_error

    started = time.monotonic()
    with timed_stage("Discover required formula dependencies"):
        links = required_links(working, columns)
    for link in links:
        path = Path(link)
        if not path.is_file():
            raise BondsLiteError(f"Required formula dependency is unavailable: {path}")
        if any(book.name.casefold() == path.name.casefold() for book in working.app.books):
            raise BondsLiteError(f"Close the required dependency before running: {path}")
        try:
            with timed_stage(f"Refresh dependency: {path.name}"):
                working.api.UpdateLink(Name=link, Type=1)
        except com_error as error:
            raise BondsLiteError(
                f"Required formula dependency could not be refreshed: {link}\n"
                "Check the dependency workbook and its referenced sheets in Excel"
            ) from error
    for worksheet in working.sheets:
        with timed_stage(f"Calculate worksheet: {worksheet.name}"):
            worksheet.api.Calculate()
    with timed_stage("Wait for Excel calculation"):
        while working.app.api.CalculationState == 1:
            if time.monotonic() - started > timeout:
                raise BondsLiteError("Excel calculation exceeded the allowed runtime")
            pythoncom.PumpWaitingMessages()
            time.sleep(0.05)
    if time.monotonic() - started > timeout:
        raise BondsLiteError("Dependency refresh/calculation exceeded the allowed runtime")
    with timed_stage("Validate refreshed dependency status"):
        for link in links:
            status = working.api.LinkInfo(link, 3, 1)
            if status == 3:
                logger.warning(
                    f"Warning: Excel reports {Path(link).name} as old (status 3) "
                    "after refresh; checking calculated formulas. Keep source "
                    "workbook calculations current upstream"
                )
            elif status != 0:
                raise BondsLiteError(
                    f"Required link did not refresh successfully: {Path(link).name} "
                    f"(status {status})"
                )
    with timed_stage("Validate model base curve values"):
        curve_sheet = working.sheets["P. Base Curves"]
        curve_table = curve_sheet.api.ListObjects("P_tbl_BaseCurves")
        curve_range = curve_table.ListColumns("GBP SONIA").DataBodyRange
        curve_address = curve_range.Address
        curve_errors = curve_sheet.api.Evaluate(
            f"SUMPRODUCT(--ISERROR({curve_address}))"
        )
        curve_numbers = curve_sheet.api.Evaluate(
            f"SUMPRODUCT(--ISNUMBER({curve_address}))"
        )
        if (
            not isinstance(curve_errors, Real)
            or curve_errors != 0
            or not isinstance(curve_numbers, Real)
            or curve_numbers != curve_table.ListRows.Count
        ):
            raise BondsLiteError(
                "P. Base Curves[GBP SONIA] must contain numeric values after "
                "calculation; ensure the JRC_SpotAnnual add-in is loaded in Excel"
            )
    sheet = working.sheets["I. Bonds Data"]
    table = sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    last = table.Range.Row + table.ListRows.Count
    failures = []
    with timed_stage("Validate recalculated bond formulas"):
        for column in columns:
            address = f"{column_name(column)}9:{column_name(column)}{last}"
            count = sheet.api.Evaluate(f"SUMPRODUCT(--ISERROR({address}))")
            if not isinstance(count, (int, float)) or count != 0:
                failures.append(address)
        if failures:
            raise BondsLiteError("Excel formula errors after recalculation: " + ", ".join(failures))


def verify_saved_copy(
    book: xw.Book, settings: StageOneInputs, rows: int, columns: list[int],
    stresses: dict[str, list[list[float]]],
    model_output_folder: Path | None = None,
) -> None:
    """Verify the generated workbook after saving, without reopening it in Excel."""
    sheet = book.sheets["I. Bonds Data"]
    table = sheet.api.ListObjects("I_tbl_Monthly_Bonds")
    if table.ListRows.Count != rows:
        raise BondsLiteError("Saved output has an incorrect bond row count")
    for column in columns:
        if sheet.range((9, column), (8 + rows, column)).api.HasFormula is not True:
            raise BondsLiteError(f"Saved output lost formulas in {column_name(column)}")
    for address, grid in stresses.items():
        if book.sheets["I. Sensitivity"].range(address).options(ndim=2).value != grid:
            raise BondsLiteError(f"Saved output stress values differ at {address}")
    target = book.sheets["I. Control"]
    if target.range("D10").value != settings.valuation_date or target.range("D13").value != settings.entity:
        raise BondsLiteError("Saved output control values differ")
    if model_output_folder is not None:
        value = _model_output_cell(book).value
        if not isinstance(value, str) or Path(value).resolve() != model_output_folder.resolve():
            raise BondsLiteError("Saved output_folder_path differs from its dataset output folder")


def generate_bonds_lite_input(caller: xw.Book) -> Path:
    """Generate a paired run under the guard without saving caller audit data."""
    with workbook_run_guard(_caller_path(caller)):
        return _generate_bonds_lite_input(caller)


def run_stage_one(caller: xw.Book) -> Path:
    """Persist Stage 1's latest attempt and expose only a committed output handoff."""
    caller_path = _caller_path(caller)
    with workbook_run_guard(caller_path):
        require_stage_panel(caller)
        previous_output = _read_stage_fields(caller).get("last_output")
        started = datetime.now()
        clock = time.perf_counter()
        with timed_stage("Record Stage 1 audit start"):
            write_stage_audit(caller, {
                "last_run": started.replace(hour=0, minute=0, second=0, microsecond=0),
                "start": started, "finish": None, "elapsed": None,
                "user": getpass.getuser(), "status": "Running",
                "stage_two_input": None,
            })
        try:
            output = _generate_bonds_lite_input(caller)
            relative = portable_output_path(caller_path, output)
        except BaseException as error:
            terminal = {
                "finish": datetime.now(),
                "elapsed": (time.perf_counter() - clock) / 86400,
                "status": "Interrupted" if isinstance(error, KeyboardInterrupt) else "Failed",
                "last_output": previous_output, "stage_two_input": None,
            }
            try:
                with timed_stage("Record unsuccessful Stage 1 audit"):
                    write_stage_audit(caller, terminal)
            except Exception as audit_error:
                logger.error(f"Stage 1 failed and its audit could not be saved: {audit_error}")
            raise

        finished = datetime.now()
        elapsed = (time.perf_counter() - clock) / 86400
        try:
            with timed_stage("Commit successful Stage 1 audit and handoff"):
                write_stage_audit(caller, {
                    "finish": finished, "elapsed": elapsed, "status": "Succeeded",
                    "last_output": relative, "stage_two_input": relative,
                })
        except BaseException as audit_error:
            try:
                write_stage_audit(caller, {
                    "finish": finished, "elapsed": elapsed, "status": "Failed",
                    "last_output": previous_output, "stage_two_input": None,
                })
            except Exception as recovery_error:
                logger.error(f"Stage 1 audit recovery could not be saved: {recovery_error}")
            raise BondsLiteError(
                f"Generated output retained at {output}; Stage 1 audit/handoff "
                f"could not be saved: {audit_error}"
            ) from audit_error
        return output


@timed_stage("Generate Bonds and MAP inputs")
def _generate_bonds_lite_input(caller: xw.Book) -> Path:
    """Generate one complete Bonds/MAP run while the caller owns the run guard."""
    with timed_stage("Validate inputs and credit stresses"):
        settings = validate_inputs(caller)
    run_folder = allocate_run_folder(
        settings.caller_path.parent / "Bonds Lite output", datetime.now()
    )
    try:
        for dataset in INPUT_FILENAMES:
            (run_folder / f"{settings.entity} {dataset}").mkdir()
        with excel_state(caller.app):
            for dataset, data_path in (
                ("Bonds", settings.bonds_path), ("MAP", settings.map_path),
            ):
                with timed_stage(f"Generate {dataset} input"):
                    _generate_dataset_input(caller.app, dataset, data_path, settings, run_folder)
        with timed_stage("Validate complete Bonds and MAP run"):
            portable_output_path(settings.caller_path, run_folder)
        return run_folder
    except BaseException:
        try:
            with timed_stage("Remove incomplete Stage 1 attempt"):
                shutil.rmtree(run_folder)
        except Exception as cleanup_error:
            logger.error(
                f"Incomplete run folder could not be removed: {run_folder}; "
                f"{cleanup_error}. This attempt is not a successful Stage 1 output"
            )
        raise


def _generate_dataset_input(
    application: xw.App, dataset: str, data_path: Path,
    settings: StageOneInputs, run_folder: Path,
) -> Path:
    """Prepare one input privately and publish it only after save/verification."""
    model_output_folder = run_folder / f"{settings.entity} {dataset}"
    with tempfile.TemporaryDirectory(prefix=".bonds-lite-", dir=run_folder) as temporary:
        private_folder = Path(temporary)
        working_path = private_folder / f"{dataset}-{private_folder.name}.xlsx"
        with _stage_source_locally(data_path) as local_data_path:
            with owned_book(application, local_data_path) as data:
                with timed_stage(f"Copy {dataset} input template"):
                    shutil.copy2(settings.template_path, working_path)
                with owned_book(application, working_path, read_only=False) as working:
                    with timed_stage(f"Import {dataset} data and resize table"):
                        rows, columns = populate_bond_data(data, working)
                        logger.info(f"Imported {rows} {dataset} rows")
                    with timed_stage(f"Populate {dataset} credit sensitivities"):
                        populate_credit_stresses(settings, working, settings.stresses)
                    with timed_stage(f"Set {dataset} controls and model output folder"):
                        synchronise_controls(settings, working)
                        set_model_output_folder(working, model_output_folder)
                    with timed_stage(f"Refresh {dataset} formulas"):
                        refresh_and_validate(working, columns)
                    with timed_stage(f"Save {dataset} input"):
                        working.save()
                        if not working.api.Saved:
                            raise BondsLiteError(f"Excel did not finish saving the {dataset} input")
                    with timed_stage(f"Verify {dataset} input"):
                        verify_saved_copy(
                            working, settings, rows, columns, settings.stresses,
                            model_output_folder=model_output_folder,
                        )
        if not working_path.is_file() or working_path.stat().st_size == 0:
            raise BondsLiteError(f"The {dataset} input was not saved to disk")
        final = run_folder / INPUT_FILENAMES[dataset]
        with timed_stage(f"Publish {dataset} input"):
            working_path.rename(final)
        return final


def _require_model_inputs_closed(caller: xw.Book, inputs: dict[str, ModelInput]) -> None:
    if not hasattr(caller, "app"):
        return
    import xlwings as xw

    names = {item.input_path.name.casefold() for item in inputs.values()}
    for application in xw.apps:
        for book in application.books:
            if book.name.casefold() in names:
                raise BondsLiteError(f"Close the Stage 2 input workbook in Excel before running: {book.name}")


def run_stage_two(caller: xw.Book) -> Path | None:
    """Run the committed folder once and commit results only with their saved audit."""
    caller_path = _caller_path(caller)
    with workbook_run_guard(caller_path):
        with timed_stage("Validate Stage 2 handoff and model environment"):
            require_stage_two_panel(caller)
            run_folder = resolve_stage_input(caller)
            config = load_model_config(caller_path)
            inputs = read_model_inputs(run_folder)
            _require_model_inputs_closed(caller, inputs)
            preflight_model_environment(config)
            previous = _read_stage_two_fields(caller)["last_output"]
            relative = portable_output_path(caller_path, run_folder)
        if not confirm_model_overwrite(inputs, run_folder):
            logger.info("Stage 2 cancelled; existing results and audit were not changed")
            return None
        started, clock = datetime.now(), time.perf_counter()
        write_stage_two_audit(caller, {
            "last_run": started.replace(hour=0, minute=0, second=0, microsecond=0),
            "start": started, "finish": None, "elapsed": None,
            "user": getpass.getuser(), "status": "Running",
            "input_run": relative, "last_output": previous,
        })
        try:
            with model_results_transaction(run_folder, inputs) as transaction:
                with timed_stage("Run Bonds Lite batch"):
                    run_model_batch(run_folder, config, getattr(caller, "app", None))
                with timed_stage("Validate both model results"):
                    validate_model_results(run_folder, inputs)
                with timed_stage("Commit successful Stage 2 audit"):
                    write_stage_two_audit(caller, {
                        "finish": datetime.now(), "elapsed": (time.perf_counter() - clock) / 86400,
                        "status": "Succeeded", "input_run": relative, "last_output": relative,
                    })
                transaction.commit()
        except BaseException as error:
            status = "Interrupted" if isinstance(error, KeyboardInterrupt) else "Timed Out" if isinstance(error, TimeoutError) else "Failed"
            try:
                write_stage_two_audit(caller, {
                    "finish": datetime.now(), "elapsed": (time.perf_counter() - clock) / 86400,
                    "status": status, "input_run": relative, "last_output": previous,
                })
            except Exception as audit_error:
                logger.error("Stage 2 failed and its audit could not be saved: %s", audit_error)
            raise
        return run_folder


def resolve_stage_two_results(caller: xw.Book) -> Path:
    """Permit downstream use only when both saved stages identify the same run."""
    require_stage_two_panel(caller)
    folder = resolve_stage_input(caller)
    fields = _read_stage_two_fields(caller)
    relative = portable_output_path(_caller_path(caller), folder)
    if fields["status"] != "Succeeded" or fields["input_run"] != relative or fields["last_output"] != relative:
        raise BondsLiteError("Stage 2 has not succeeded for the current committed Stage 1 run")
    return folder


def workbook_mutex_name(caller: Path) -> str:
    """Return a stable session-local mutex identity without exposing file paths."""
    identity = os.path.normcase(str(caller.resolve())).encode("utf-8")
    return "Local\\BondsLite_" + hashlib.sha256(identity).hexdigest()


@contextmanager
def workbook_run_guard(caller: Path) -> Iterator[None]:
    """Prevent concurrent runs for the same caller within a Windows session."""
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p]
    kernel.CreateMutexW.restype = ctypes.c_void_p
    kernel.ReleaseMutex.argtypes = [ctypes.c_void_p]
    kernel.CloseHandle.argtypes = [ctypes.c_void_p]
    handle = kernel.CreateMutexW(None, True, workbook_mutex_name(caller))
    if not handle:
        raise ctypes.WinError(ctypes.get_last_error())
    if ctypes.get_last_error() == 183:
        kernel.CloseHandle(handle)
        raise BondsLiteError("A Bonds Lite run is already in progress for this workbook")
    try:
        yield
    finally:
        with timed_stage("Release workbook run guard"):
            kernel.ReleaseMutex(handle)
            kernel.CloseHandle(handle)


def _install_stage_button(book: xw.Book, name: str, caption: str, macro: str, anchor: xw.Range) -> None:
    sheet = anchor.sheet
    if name in [shape.Name for shape in sheet.api.Shapes]:
        button = sheet.api.Buttons().Item(name)
    else:
        width = max(anchor.width, 205)
        end_column = anchor.column
        while sheet.range((anchor.row, end_column)).left < anchor.left + width:
            end_column += 1
        covered = sheet.range((anchor.row, anchor.column), (anchor.row + anchor.rows.count - 1, end_column - 1))
        if book.app.api.WorksheetFunction.CountA(covered.api) or covered.api.MergeCells is not False:
            raise BondsLiteError(f"Proposed {caption} button covers occupied cells or merges")
        for shape in sheet.api.Shapes:
            if shape.Left < anchor.left + width and shape.Left + shape.Width > anchor.left and shape.Top < anchor.top + anchor.height and shape.Top + shape.Height > anchor.top:
                raise BondsLiteError(f"Proposed {caption} button overlaps an existing shape")
        button = sheet.api.Buttons().Add(anchor.left, anchor.top, width, anchor.height)
        button.Name = name
    button.Caption = caption
    button.OnAction = "'" + book.name.replace("'", "''") + "'!" + macro


@contextmanager
def _installer_excel_application() -> Iterator[xw.App]:
    import xlwings as xw

    application = xw.App(visible=False, add_book=True)
    try:
        yield application
    finally:
        application.api.DisplayAlerts = False
        application.api.Quit()


def install_button(caller: Path) -> Path:
    """Install the launcher into a closed Draft workbook, preserving a backup."""
    import xlwings as xw

    caller = caller.resolve(strict=True)
    for application in xw.apps:
        for book in application.books:
            if book.name.casefold() == caller.name.casefold():
                raise BondsLiteError(f"Close {caller.name} in Excel before installation")
    backup_folder = caller.parent / "references" / "backups"
    backup_folder.mkdir(parents=True, exist_ok=True)
    backup = backup_folder / f"{caller.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsm"
    shutil.copy2(caller, backup)
    with _installer_excel_application() as application:
        application.api.AutomationSecurity = 3
        application.api.EnableEvents = False
        application.api.Calculation = -4135
        application.display_alerts = False
        book = application.books.open(str(caller), update_links=False)
        components = module = sheet = anchor = button = None
        try:
            if book.api.ReadOnly:
                raise BondsLiteError("Draft workbook is read-only or locked")
            initialise_stage_panel(book)
            components = book.api.VBProject.VBComponents
            names = [components.Item(index).Name for index in range(1, components.Count + 1)]
            if "modBondsLiteLauncher" in names:
                module = components.Item("modBondsLiteLauncher")
                if module.CodeModule.CountOfLines:
                    module.CodeModule.DeleteLines(1, module.CodeModule.CountOfLines)
            else:
                module = components.Add(1)
                module.Name = "modBondsLiteLauncher"
            module.CodeModule.AddFromString(VBA_LAUNCHER)
            sheet = book.sheets["I. Control"]
            anchor = sheet.range("D74:F76")
            _install_stage_button(
                book, "btnGenerateBondsLiteInput", "Generate Bonds Lite Input",
                "GenerateBondsLiteInput", anchor,
            )
            initialise_stage_two_panel(book)
            panel = book.names["SST_Stage2_Panel"].refers_to_range
            button_row = panel.row
            if sheet.range((button_row, 4)).left + 205 > panel.left:
                button_row += panel.rows.count + 2
            _install_stage_button(
                book, "btnRunBondsLite", "Run Bonds Lite", "RunBondsLite",
                sheet.range((button_row, 4), (button_row + 2, 6)),
            )
            if sheet.range("D39").value is None:
                sheet.range("D39").value = "Bonds Lite data template.xlsx"
            if sheet.range("D42").value is None:
                sheet.range("D42").value = "Bonds Lite input template.xlsx"
            book.save()
        finally:
            components = module = sheet = anchor = button = None
            book.close()
    return backup


def main() -> int:
    """Run the installed command-line entry point or install the Excel button."""
    configure_console_logging()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--install-button", type=Path)
    parser.add_argument("--caller")
    parser.add_argument("--excel-hwnd", type=int)
    parser.add_argument("--pause", action="store_true", help="Wait for Enter after a run")
    parser.add_argument("--stage", choices=("1", "2"), default="1")
    arguments = parser.parse_args()
    if arguments.install_button:
        backup = install_button(arguments.install_button)
        logger.info(f"Button installed; original backup: {backup}")
        return 0
    if not arguments.caller or arguments.excel_hwnd is None:
        parser.error("Specify --caller and --excel-hwnd, or --install-button")
    result = 0
    with run_progress() as progress:
        logger.info(
            "Bonds Lite input started | timings exclude Python startup and the Enter wait"
        )
        try:
            with timed_stage("Locate calling Excel workbook"):
                import xlwings as xw

                applications = [application for application in xw.apps if int(application.api.Hwnd) == arguments.excel_hwnd]
                if len(applications) != 1:
                    raise BondsLiteError("The calling Excel instance is no longer available")
                matches = [
                    book for book in applications[0].books
                    if os.path.normcase(book.api.FullName) == os.path.normcase(arguments.caller)
                ]
                if len(matches) != 1:
                    raise BondsLiteError("The exact calling workbook is no longer open")
            if arguments.stage == "2":
                output = run_stage_two(matches[0])
                if output is None:
                    logger.info("Stage 2 cancelled; no model batch was run")
                else:
                    logger.info(f"Bonds Lite model succeeded in run folder:\n{output}")
            else:
                output = run_stage_one(matches[0])
                logger.info(f"Saved Bonds Lite inputs in run folder:\n{output}")
        except KeyboardInterrupt:
            logger.error("Run interrupted by the user; managed cleanup has finished")
            result = 1
        except Exception as error:
            logger.error(f"Bonds Lite input failed: {error}")
            result = 1
    logger.info(f"Total processing time: {time.perf_counter() - progress.started:.1f}s")
    if arguments.pause:
        try:
            input("Press Enter to close this console...")
        except (EOFError, KeyboardInterrupt):
            pass
    return result


def notify_user(message: str, owner: int = 0, failure: bool = False) -> None:
    """Display one concise result without exposing bond data in a log file."""
    if len(message) > 2500:
        message = message[:2500] + "\nFurther validation errors omitted; complete all required blocks"
    user = ctypes.WinDLL("user32", use_last_error=True)
    user.MessageBoxW.argtypes = [ctypes.c_void_p, ctypes.c_wchar_p, ctypes.c_wchar_p, ctypes.c_uint]
    user.MessageBoxW(owner, message, "Bonds Lite input", 0x10 if failure else 0x40)


if __name__ == "__main__":
    sys.exit(main())